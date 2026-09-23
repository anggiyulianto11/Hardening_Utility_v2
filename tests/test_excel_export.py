from openpyxl import load_workbook

from controls.registry import ControlRegistry
from reporting.excel_export import AssessmentExcelExporter


def test_export_contains_expected_sheets(tmp_path):
    output = tmp_path / "assessment.xlsx"
    AssessmentExcelExporter(ControlRegistry()).export(output)
    wb = load_workbook(output, read_only=True, data_only=False)
    assert wb.sheetnames == [
        "Ringkasan",
        "Detail Kontrol",
        "Perlu Bukti",
        "Peringatan",
        "Kesiapan Implementasi",
        "Bukti Teknis",
        "Informasi Pemeriksaan",
    ]


def test_detail_contains_118_controls(tmp_path):
    output = tmp_path / "assessment.xlsx"
    AssessmentExcelExporter(ControlRegistry()).export(output)
    wb = load_workbook(output, read_only=True)
    ws = wb["Detail Kontrol"]
    assert ws.max_row == 119


def test_no_secret_in_export(tmp_path):
    output = tmp_path / "assessment.xlsx"
    AssessmentExcelExporter(ControlRegistry()).export(output)
    wb = load_workbook(output, read_only=True, data_only=False)
    all_values = " ".join(
        str(cell.value)
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "token=secret" not in all_values
