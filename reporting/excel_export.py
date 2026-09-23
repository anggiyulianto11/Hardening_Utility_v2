from collections import Counter
from datetime import datetime
from pathlib import Path
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from analysis.evidence_sanitizer import sanitize_evidence


STATUS_COLORS = {
    "SESUAI": "C6EFCE",
    "TIDAK SESUAI": "FFC7CE",
    "PERINGATAN": "FCE4D6",
    "PERLU TINJAUAN": "FFF2CC",
    "BELUM DIPERIKSA": "E7E6E6",
    "TIDAK BERLAKU": "D9EAD3",
    "TIDAK DAPAT DINILAI": "D9E1F2",
    "ERROR": "F4CCCC",
}


class AssessmentExcelExporter:
    def __init__(self, control_registry):
        self.control_registry = control_registry

    def export(
        self,
        path,
        assessment_results=None,
        semi_workspace=None,
        manual_workspace=None,
        assessment_run=None,
        organization_name="",
        assessor_name="",
    ):
        assessment_results = list(assessment_results or [])
        result_map = self._result_map(assessment_results)
        semi_map = self._workspace_map(semi_workspace)
        manual_map = self._workspace_map(manual_workspace)

        wb = Workbook()
        wb.remove(wb.active)
        self._summary_sheet(
            wb, result_map, semi_map, manual_map,
            organization_name, assessor_name, assessment_run,
        )
        self._detail_sheet(wb, result_map, semi_map, manual_map)
        self._needs_evidence_sheet(wb, result_map, semi_map, manual_map)
        self._warnings_sheet(wb, result_map, semi_map, manual_map)
        self._readiness_sheet(wb)
        self._technical_evidence_sheet(wb, assessment_results)
        self._information_sheet(
            wb, organization_name, assessor_name, assessment_run,
        )

        for ws in wb.worksheets:
            self._finish_sheet(ws)
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        destination = Path(path)
        wb.save(destination)
        return destination

    def _result_map(self, results):
        grouped = {}
        for result in results:
            grouped.setdefault(result.control_id, []).append(result)
        return grouped

    @staticmethod
    def _workspace_map(workspace):
        if workspace is None:
            return {}
        return {item.control_id: item for item in workspace.all()}

    def _effective(self, control, result_map, semi_map, manual_map):
        if control.control_id in manual_map:
            record = manual_map[control.control_id]
            return {
                "status": record.status,
                "current": record.reviewer_notes or "-",
                "evidence": self._evidence_refs(record.evidence_references),
                "notes": record.applicability_reason or record.business_exception,
                "reviewer": record.reviewer_name,
            }
        if control.control_id in semi_map:
            record = semi_map[control.control_id]
            return {
                "status": record.status,
                "current": json.dumps(
                    sanitize_evidence(record.technical_summary),
                    ensure_ascii=False,
                ) if record.technical_summary else "-",
                "evidence": self._evidence_refs(record.evidence_references),
                "notes": record.reviewer_notes,
                "reviewer": "",
            }
        results = result_map.get(control.control_id, [])
        if not results:
            return {
                "status": "BELUM DIPERIKSA",
                "current": "-",
                "evidence": "-",
                "notes": "Analyzer/bukti belum tersedia.",
                "reviewer": "",
            }
        statuses = [item.status.value for item in results]
        priority = [
            "ERROR", "TIDAK SESUAI", "PERINGATAN", "PERLU TINJAUAN",
            "TIDAK DAPAT DINILAI", "SESUAI",
        ]
        status = next((value for value in priority if value in statuses), statuses[0])
        current = "\n".join(
            f"{item.target_name}: {item.current_condition or '-'}"
            for item in results
        )
        evidence = "\n".join(
            f"{item.target_name}: {json.dumps(sanitize_evidence(item.evidence), ensure_ascii=False)}"
            for item in results if item.evidence
        ) or "-"
        errors = "; ".join(item.error for item in results if item.error)
        return {
            "status": status,
            "current": current,
            "evidence": evidence,
            "notes": errors,
            "reviewer": "",
        }

    @staticmethod
    def _evidence_refs(references):
        if not references:
            return "-"
        return "\n".join(
            f"[{item.evidence_type}] {item.label}: {item.location}"
            for item in references
        )

    def _summary_sheet(
        self, wb, result_map, semi_map, manual_map,
        organization_name, assessor_name, assessment_run,
    ):
        ws = wb.create_sheet("Ringkasan")
        ws["A1"] = "ArcGIS Enterprise Hardening Assessment"
        ws["A2"] = "Ringkasan Assessment"
        ws["A4"] = "Organisasi"
        ws["B4"] = organization_name or "-"
        ws["A5"] = "Assessor"
        ws["B5"] = assessor_name or "-"
        ws["A6"] = "Tanggal Ekspor"
        ws["B6"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws["A7"] = "Catalog Version"
        ws["B7"] = self.control_registry.metadata.get("catalog_version", "-")
        ws["A8"] = "Assessment Run ID"
        ws["B8"] = getattr(assessment_run, "run_id", "-") if assessment_run else "-"

        statuses = []
        profiles = Counter()
        for control in self.control_registry.get_all_controls():
            effective = self._effective(control, result_map, semi_map, manual_map)
            statuses.append(effective["status"])
            profiles[control.profile] += 1
        counts = Counter(statuses)
        decided = len(statuses) - counts.get("BELUM DIPERIKSA", 0)
        ws["A10"] = "Total Kontrol"
        ws["B10"] = len(statuses)
        ws["A11"] = "Basic"
        ws["B11"] = profiles["Basic"]
        ws["A12"] = "Advanced"
        ws["B12"] = profiles["Advanced"]
        ws["A13"] = "Kontrol Tuntas"
        ws["B13"] = decided
        ws["A14"] = "Completion"
        ws["B14"] = decided / len(statuses) if statuses else 1
        ws["B14"].number_format = "0.0%"

        ws["D4"] = "Status"
        ws["E4"] = "Jumlah"
        for row, status in enumerate(
            ["SESUAI", "TIDAK SESUAI", "PERINGATAN", "PERLU TINJAUAN",
             "BELUM DIPERIKSA", "TIDAK BERLAKU", "TIDAK DAPAT DINILAI", "ERROR"],
            start=5,
        ):
            ws.cell(row, 4, status)
            ws.cell(row, 5, counts.get(status, 0))
            ws.cell(row, 4).fill = PatternFill("solid", fgColor=STATUS_COLORS[status])

        eligible = decided - counts.get("TIDAK BERLAKU", 0) - counts.get("TIDAK DAPAT DINILAI", 0) - counts.get("ERROR", 0)
        score_numerator = counts.get("SESUAI", 0) + 0.5 * counts.get("PERINGATAN", 0)
        ws["A16"] = "Skor Sementara"
        ws["B16"] = score_numerator / eligible if eligible > 0 else 0
        ws["B16"].number_format = "0%"
        ws["A17"] = "Catatan"
        ws["B17"] = "Skor hanya menghitung kontrol yang telah diputuskan dan dapat dinilai."

    def _detail_sheet(self, wb, result_map, semi_map, manual_map):
        ws = wb.create_sheet("Detail Kontrol")
        headers = [
            "Urutan", "ID Kontrol", "Bagian Panduan", "Profil", "Area",
            "Nama Kontrol", "Cara Pemeriksaan", "Esri / 3rd Party",
            "Risiko Keamanan", "Risiko Privasi", "Status", "Kondisi Saat Ini",
            "Kondisi yang Diharapkan", "Rekomendasi", "Bukti / Referensi",
            "Catatan Pemeriksaan", "Reviewer",
        ]
        ws.append(headers)
        for control in self.control_registry.get_all_controls():
            effective = self._effective(control, result_map, semi_map, manual_map)
            ws.append([
                control.sequence, control.control_id, control.guide_section,
                control.profile, control.area, control.control_name,
                control.verification_method, control.implementation_side,
                control.security_risk, control.privacy_risk,
                effective["status"], effective["current"],
                control.expected_condition, control.recommendation,
                effective["evidence"], effective["notes"], effective["reviewer"],
            ])
        self._color_status_column(ws, 11)

    def _filtered_sheet(self, wb, title, accepted, result_map, semi_map, manual_map):
        ws = wb.create_sheet(title)
        ws.append(["ID Kontrol", "Profil", "Area", "Nama Kontrol", "Status", "Kondisi Saat Ini", "Rekomendasi", "Bukti / Catatan"])
        for control in self.control_registry.get_all_controls():
            effective = self._effective(control, result_map, semi_map, manual_map)
            if effective["status"] not in accepted:
                continue
            ws.append([
                control.control_id, control.profile, control.area,
                control.control_name, effective["status"], effective["current"],
                control.recommendation,
                "\n".join(x for x in [effective["evidence"], effective["notes"]] if x and x != "-"),
            ])
        self._color_status_column(ws, 5)

    def _needs_evidence_sheet(self, wb, result_map, semi_map, manual_map):
        self._filtered_sheet(
            wb, "Perlu Bukti",
            {"PERLU TINJAUAN", "BELUM DIPERIKSA", "TIDAK DAPAT DINILAI", "ERROR"},
            result_map, semi_map, manual_map,
        )

    def _warnings_sheet(self, wb, result_map, semi_map, manual_map):
        self._filtered_sheet(
            wb, "Peringatan", {"PERINGATAN", "TIDAK SESUAI"},
            result_map, semi_map, manual_map,
        )

    def _readiness_sheet(self, wb):
        ws = wb.create_sheet("Kesiapan Implementasi")
        ws.append([
            "ID Kontrol", "Nama Kontrol", "Dapat Diubah", "Kesiapan",
            "Target", "Parameter / Keputusan", "Perubahan yang Direncanakan",
            "Dampak Operasional", "Rencana Pemulihan",
        ])
        for control in self.control_registry.get_all_controls():
            ws.append([
                control.control_id, control.control_name,
                "Ya" if control.changeable else "Tidak",
                control.implementation_readiness or "-",
                control.implementation_target or "-",
                control.implementation_parameter_decision or "-",
                control.planned_change or "-",
                control.operational_impact or "-",
                control.recovery_plan or "-",
            ])

    def _technical_evidence_sheet(self, wb, results):
        ws = wb.create_sheet("Bukti Teknis")
        ws.append([
            "Target", "Component", "Control ID", "Status",
            "Current Condition", "Evidence Sanitized", "Error",
        ])
        for result in results:
            ws.append([
                result.target_name, result.component_type, result.control_id,
                result.status.value, result.current_condition,
                json.dumps(sanitize_evidence(result.evidence), ensure_ascii=False),
                result.error or "-",
            ])
        self._color_status_column(ws, 4)

    def _information_sheet(self, wb, organization_name, assessor_name, assessment_run):
        ws = wb.create_sheet("Informasi Pemeriksaan")
        rows = [
            ("Organisasi", organization_name or "-"),
            ("Assessor", assessor_name or "-"),
            ("Tanggal Ekspor", datetime.now().isoformat(timespec="seconds")),
            ("Catalog Version", self.control_registry.metadata.get("catalog_version", "-")),
            ("Source Catalog", self.control_registry.metadata.get("source_file", "-")),
            ("Assessment Run ID", getattr(assessment_run, "run_id", "-") if assessment_run else "-"),
            ("Catatan Keamanan", "Token, password, shared key, secret, cookie, credential, dan authorization data disanitasi sebelum ekspor."),
            ("Interpretasi Skor", "Skor bersifat sementara sampai seluruh kontrol yang berlaku telah diputuskan."),
        ]
        for key, value in rows:
            ws.append([key, value])

    def _finish_sheet(self, ws):
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2" if ws.max_row > 2 else None
        header_fill = PatternFill("solid", fgColor="1F4E78")
        header_font = Font(color="FFFFFF", bold=True)
        thin = Side(style="thin", color="D9E2F3")
        if ws.max_row >= 1:
            for cell in ws[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = Border(bottom=thin)
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in range(1, ws.max_column + 1):
            max_len = 0
            for cell in ws.iter_cols(min_col=column, max_col=column):
                for item in cell:
                    value = "" if item.value is None else str(item.value)
                    max_len = max(max_len, min(max(len(line) for line in value.split("\n")), 60))
            ws.column_dimensions[get_column_letter(column)].width = max(12, min(max_len + 2, 60))
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions

    @staticmethod
    def _color_status_column(ws, column):
        for row in range(2, ws.max_row + 1):
            status = ws.cell(row, column).value
            if status in STATUS_COLORS:
                ws.cell(row, column).fill = PatternFill("solid", fgColor=STATUS_COLORS[status])
