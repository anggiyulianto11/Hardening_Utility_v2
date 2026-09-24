import sys
import time
import json
import html
import ssl
import socket
import hashlib
from datetime import datetime, timezone
from email.utils import parsedate_to_datetime
from pathlib import Path
from urllib.parse import urlparse, urlunparse, quote

import requests
import urllib3
from concurrent.futures import ThreadPoolExecutor, as_completed
from PySide6.QtCore import QObject, QThread, Signal, Qt, QTimer, QSettings
from PySide6.QtGui import QColor, QFont
from PySide6.QtWidgets import (
    QApplication, QButtonGroup, QCheckBox, QComboBox, QDialog, QDialogButtonBox, QFormLayout, QFrame,
    QFileDialog, QHBoxLayout, QHeaderView, QLabel, QLineEdit, QListWidget, QMainWindow,
    QMessageBox, QPlainTextEdit, QPushButton, QRadioButton, QScrollArea, QStatusBar, QTabWidget,
    QTextBrowser,
    QTableWidget, QTableWidgetItem, QVBoxLayout, QWidget,
)

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)


def normalize_url(value):
    return value.strip().rstrip("/")


def derive_portal_base_url(portal_admin_url):
    value = normalize_url(portal_admin_url)
    parsed = urlparse(value)
    path = parsed.path.rstrip("/")
    if not path.lower().endswith("/portaladmin"):
        raise ValueError("Portal Admin URL harus berakhir dengan /portaladmin.")
    base_path = path[:-len("/portaladmin")]
    return urlunparse((parsed.scheme, parsed.netloc, base_path, "", "", "")).rstrip("/")


def find_value(data, *keys, default="Tidak tersedia"):
    wanted = {key.lower() for key in keys}
    if isinstance(data, dict):
        for key, value in data.items():
            if str(key).lower() in wanted:
                return value
        for value in data.values():
            found = find_value(value, *keys, default=None)
            if found is not None:
                return found
    elif isinstance(data, list):
        for value in data:
            found = find_value(value, *keys, default=None)
            if found is not None:
                return found
    return default


def bool_text(value):
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)
def public_probe_path(component, registered_url):
    base = normalize_url(registered_url)
    if component == "portal":
        return f"{base}/home"
    return f"{base}/rest/services"
def http_variant(url):
    parsed = urlparse(url)
    return urlunparse(("http", parsed.netloc, parsed.path, "", parsed.query, ""))
def https_variant(url):
    parsed = urlparse(url)
    return urlunparse(("https", parsed.netloc, parsed.path, "", parsed.query, ""))



def tls_certificate_evidence(url, timeout=6):
    """Inspect the certificate actually served by an HTTPS hostname and port."""
    parsed = urlparse(normalize_url(url))
    host = parsed.hostname
    port = parsed.port or 443
    if not host:
        return {"status": "UNKNOWN", "error": "Invalid hostname", "url": url}
    result = {"url": url, "hostname": host, "port": port, "path": parsed.path or "/"}
    try:
        context = ssl.create_default_context()
        with socket.create_connection((host, port), timeout=timeout) as raw:
            with context.wrap_socket(raw, server_hostname=host) as wrapped:
                cert = wrapped.getpeercert()
                der = wrapped.getpeercert(binary_form=True)
        subject = dict(item[0] for item in cert.get("subject", ()))
        issuer = dict(item[0] for item in cert.get("issuer", ()))
        sans = [value for kind, value in cert.get("subjectAltName", ()) if kind == "DNS"]
        not_before = cert.get("notBefore")
        not_after = cert.get("notAfter")
        expires_epoch = ssl.cert_time_to_seconds(not_after) if not_after else None
        days_remaining = int((expires_epoch - time.time()) // 86400) if expires_epoch else None
        self_signed = bool(subject and issuer and subject == issuer)
        result.update({
            "status": "PASS", "trusted": True, "hostname_match": True,
            "subject_cn": subject.get("commonName", "Tidak tersedia"),
            "issuer_cn": issuer.get("commonName", "Tidak tersedia"),
            "sans": sans, "not_before": not_before, "not_after": not_after,
            "days_remaining": days_remaining, "self_signed": self_signed,
            "sha256": hashlib.sha256(der).hexdigest().upper(),
        })
        if self_signed or (days_remaining is not None and days_remaining < 0):
            result["status"] = "FAIL"
        elif days_remaining is not None and days_remaining < 60:
            result["status"] = "PASS WITH NOTE"
    except ssl.SSLCertVerificationError as error:
        result.update({"status": "FAIL", "trusted": False, "error": str(error)})
    except (socket.timeout, TimeoutError) as error:
        result.update({"status": "UNKNOWN", "error": type(error).__name__})
    except Exception as error:
        result.update({"status": "UNKNOWN", "error": f"{type(error).__name__}: {error}"})
    return result

class ArcGISRequestError(Exception):
    pass


class ConnectionWorker(QObject):
    progress = Signal(str)
    completed = Signal(dict)
    finished = Signal()

    def __init__(
        self, mode, purpose, portal_admin_url, portal_username, portal_password,
        server_admin_url, server_username, server_password, verify_ssl=False,
        portal_token=None, portal_token_expires=0,
        server_token=None, server_token_expires=0,
        run_network_probes=False,
    ):
        super().__init__()
        self.mode = mode
        self.purpose = purpose
        self.portal_admin_url = normalize_url(portal_admin_url)
        self.portal_base_url = (
            derive_portal_base_url(self.portal_admin_url)
            if self.portal_admin_url else ""
        )
        self.portal_username = portal_username.strip()
        self.portal_password = portal_password
        self.server_admin_url = normalize_url(server_admin_url)
        self.server_username = server_username.strip()
        self.server_password = server_password
        self.verify_ssl = verify_ssl
        self.portal_token = portal_token
        self.portal_token_expires = float(portal_token_expires or 0)
        self.server_token = server_token
        self.server_token_expires = float(server_token_expires or 0)
        self.run_network_probes = bool(run_network_probes)
        self.web_tier_tls_cache = {}

    def session(self):
        session = requests.Session()
        session.verify = self.verify_ssl
        session.headers.update({"User-Agent": "ArcGIS-Hardening-Utility/0.6"})
        if self.portal_base_url:
            session.headers.update({"Referer": f"{self.portal_base_url}/home/"})
        return session

    def request_json(self, session, method, url, **kwargs):
        try:
            response = session.request(method, url, timeout=60, **kwargs)
        except requests.exceptions.SSLError as error:
            raise ArcGISRequestError(
                "Koneksi HTTPS tidak dapat diverifikasi. Matikan verifikasi TLS untuk pengujian."
            ) from error
        except requests.exceptions.Timeout as error:
            raise ArcGISRequestError(f"Koneksi timeout.\nEndpoint: {url}") from error
        except requests.exceptions.ConnectionError as error:
            raise ArcGISRequestError(
                f"Tidak dapat terhubung. Periksa DNS, VPN, firewall, dan port.\nEndpoint: {url}"
            ) from error
        except requests.exceptions.RequestException as error:
            raise ArcGISRequestError(f"HTTP request gagal.\nEndpoint: {url}\nDetail: {error}") from error

        if response.status_code >= 400:
            labels = {
                401: "Autentikasi ditolak.", 403: "Akses endpoint ditolak.",
                404: "Endpoint tidak ditemukan.", 405: "Metode HTTP ditolak.",
            }
            raise ArcGISRequestError(
                f"{labels.get(response.status_code, 'Request gagal.')}\n"
                f"HTTP {response.status_code}\nEndpoint: {url}"
            )
        try:
            result = response.json()
        except ValueError as error:
            raise ArcGISRequestError(
                f"Endpoint tidak mengembalikan JSON.\nEndpoint: {url}\n"
                f"Content-Type: {response.headers.get('Content-Type', 'Tidak diketahui')}"
            ) from error
        if isinstance(result, dict) and result.get("error"):
            err = result["error"]
            raise ArcGISRequestError(
                f"ArcGIS API menolak request.\nCode: {err.get('code', '-')}\n"
                f"Message: {err.get('message', '-')}\nDetails: {err.get('details', '-')}"
            )
        return result

    def token_is_usable(self, token, expires_at):
        # Refresh proaktif jika sisa masa berlaku kurang dari dua menit.
        return bool(token) and float(expires_at or 0) > (time.time() + 120)

    def generate_portal_token(self, session):
        self.progress.emit("Portal: token tidak tersedia/kedaluwarsa; membuat token baru...")
        result = self.request_json(
            session, "POST", f"{self.portal_base_url}/sharing/rest/generateToken",
            data={
                "username": self.portal_username,
                "password": self.portal_password,
                "client": "referer",
                "referer": f"{self.portal_base_url}/home/",
                "expiration": "30",
                "f": "json",
            },
        )
        token = result.get("token")
        if not token:
            raise ArcGISRequestError(
                "Autentikasi Portal gagal. Periksa username, password, dan hak administrator."
            )
        expires_ms = result.get("expires", 0)
        self.portal_token = token
        self.portal_token_expires = (
            float(expires_ms) / 1000 if expires_ms else time.time() + 25 * 60
        )
        return token

    def ensure_portal_token(self, session):
        if self.token_is_usable(self.portal_token, self.portal_token_expires):
            self.progress.emit("Portal: menggunakan token yang masih valid dari memory...")
            return self.portal_token
        return self.generate_portal_token(session)

    def generate_server_token(self, session):
        self.progress.emit("Server: token tidak tersedia/kedaluwarsa; membuat token baru...")
        result = self.request_json(
            session, "POST", f"{self.server_admin_url}/generateToken",
            data={
                "username": self.server_username,
                "password": self.server_password,
                "client": "requestip",
                "expiration": "30",
                "f": "json",
            },
        )
        token = result.get("token")
        if not token:
            raise ArcGISRequestError(
                "Autentikasi ArcGIS Server gagal. Periksa Primary Site Administrator credentials."
            )
        expires_ms = result.get("expires", 0)
        self.server_token = token
        self.server_token_expires = (
            float(expires_ms) / 1000 if expires_ms else time.time() + 25 * 60
        )
        return token

    def ensure_server_token(self, session):
        if self.token_is_usable(self.server_token, self.server_token_expires):
            self.progress.emit("Server: menggunakan token yang masih valid dari memory...")
            return self.server_token
        return self.generate_server_token(session)
    @staticmethod
    def web_adaptor_ids(collection):
        if not isinstance(collection, dict):
            return []
        candidates = collection.get("webAdaptors", collection.get("items", []))
        if isinstance(candidates, dict):
            candidates = list(candidates.values())
        ids = []
        for item in candidates if isinstance(candidates, list) else []:
            if isinstance(item, str):
                ids.append(item)
            elif isinstance(item, dict):
                value = item.get("id", item.get("webAdaptorId", item.get("machineName")))
                if value:
                    ids.append(str(value))
        return list(dict.fromkeys(ids))
    def discover_web_adaptors(self, session, admin_url, token, component):
        started = time.perf_counter()
        self.progress.emit(f"{component.title()}: menemukan registered Web Adaptors...")
        collection = self.request_json(
            session, "GET", f"{admin_url}/system/webadaptors",
            params={"token": token, "f": "json", "_ts": str(int(time.time() * 1000))},
            headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
        )
        raw_items = collection.get("webAdaptors", []) if isinstance(collection, dict) else []
        collection_by_id = {}
        if isinstance(raw_items, list):
            for item in raw_items:
                if isinstance(item, dict):
                    item_id = find_value(item, "id", "webAdaptorId", default=None)
                    if item_id:
                        collection_by_id[str(item_id)] = dict(item)
        ids = self.web_adaptor_ids(collection)
        details = []
        if not ids and isinstance(raw_items, list):
            details.extend(dict(item) for item in raw_items if isinstance(item, dict))
        for adaptor_id in ids:
            detail = self.request_json(
                session, "GET", f"{admin_url}/system/webadaptors/{adaptor_id}",
                params={"token": token, "f": "json", "_ts": str(int(time.time() * 1000))},
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
            merged = dict(collection_by_id.get(str(adaptor_id), {}))
            if isinstance(detail, dict):
                merged.update(detail)
            merged.setdefault("id", adaptor_id)
            details.append(merged)
        normalized = []
        for item in details:
            registered_url = find_value(
                item, "url", "webAdaptorURL", "webAdaptorUrl", default=None
            )
            normalized.append({
                "id": find_value(item, "id", "webAdaptorId", default="Tidak tersedia"),
                "url": registered_url,
                "machine_name": find_value(item, "machineName", default="Tidak tersedia"),
                "machine_ip": find_value(item, "machineIP", "machineIp", default="Tidak tersedia"),
                "http_port": find_value(item, "httpPort", default="Tidak tersedia"),
                "https_port": find_value(item, "httpsPort", default="Tidak tersedia"),
                "version": find_value(item, "currentVersion", "version", default="Tidak tersedia"),
                "web_server": find_value(item, "webServer", default="Tidak tersedia"),
                "os": find_value(item, "os", "operatingSystem", default="Tidak tersedia"),
                "probe": {"http_result": "NOT_RUN", "https_result": "NOT_RUN"},
            })
        inventory_elapsed = time.perf_counter() - started
        self.progress.emit(
            f"{component.title()}: Web Adaptor metadata selesai dalam {inventory_elapsed:.2f}s "
            f"({len(normalized)} ditemukan)."
        )
        if not self.run_network_probes:
            self.progress.emit(
                f"{component.title()}: network probe dilewati pada Test Connection; "
                "probe dijalankan saat Analyze Hardening."
            )
            return normalized
        jobs = []
        for index, record in enumerate(normalized):
            if record.get("url"):
                jobs.append((index, public_probe_path(component, record["url"])))
            else:
                record["probe"] = {
                    "http_result": "UNKNOWN", "https_result": "UNKNOWN",
                    "detail": "Registered URL unavailable",
                }
        if not jobs:
            return normalized
        probe_started = time.perf_counter()
        self.progress.emit(
            f"{component.title()}: memulai {len(jobs)} HTTP/HTTPS probe paralel..."
        )
        with ThreadPoolExecutor(max_workers=min(8, len(jobs))) as executor:
            futures = {
                executor.submit(self.probe_https_endpoint, secure_url, component, index + 1): index
                for index, secure_url in jobs
            }
            for future in as_completed(futures):
                index = futures[future]
                try:
                    normalized[index]["probe"] = future.result()
                except Exception as error:
                    normalized[index]["probe"] = {
                        "http_result": "UNKNOWN", "https_result": "UNKNOWN",
                        "detail": type(error).__name__,
                    }
        elapsed = time.perf_counter() - probe_started
        self.progress.emit(
            f"{component.title()}: seluruh network probe selesai dalam {elapsed:.2f}s."
        )
        return normalized
    def probe_one_url(self, url, scheme, label):
        started = time.perf_counter()
        headers = {"User-Agent": "ArcGIS-Hardening-Utility/0.8", "Cache-Control": "no-cache"}
        try:
            response = requests.get(
                url, headers=headers, timeout=(4, 8),
                allow_redirects=(scheme == "https"), verify=self.verify_ssl,
            )
            return {
                "status": response.status_code,
                "location": response.headers.get("Location"),
                "final_url": response.url,
                "hsts": bool(response.headers.get("Strict-Transport-Security")),
                "elapsed": time.perf_counter() - started,
            }
        except requests.exceptions.SSLError:
            return {"error": "TLS_VERIFICATION_FAILED", "elapsed": time.perf_counter() - started}
        except requests.exceptions.ConnectTimeout:
            return {"error": "CONNECT_TIMEOUT", "elapsed": time.perf_counter() - started}
        except requests.exceptions.ReadTimeout:
            return {"error": "READ_TIMEOUT", "elapsed": time.perf_counter() - started}
        except requests.exceptions.ConnectionError:
            return {"error": "NOT_EXPOSED_FROM_TEST_LOCATION", "elapsed": time.perf_counter() - started}
        except requests.exceptions.RequestException as error:
            return {"error": type(error).__name__.upper(), "elapsed": time.perf_counter() - started}
    def probe_https_endpoint(self, secure_url, component, number):
        http_url = http_variant(secure_url)
        https_url = https_variant(secure_url)
        self.progress.emit(f"{component.title()} Web Adaptor {number}: menguji HTTP dan HTTPS secara paralel...")
        with ThreadPoolExecutor(max_workers=2) as executor:
            http_future = executor.submit(self.probe_one_url, http_url, "http", "HTTP")
            https_future = executor.submit(self.probe_one_url, https_url, "https", "HTTPS")
            http = http_future.result()
            https = https_future.result()
        result = {
            "secure_url": secure_url, "http_url": http_url,
            "http_status": http.get("status"), "http_location": None,
            "http_result": "UNKNOWN", "https_status": https.get("status"),
            "https_result": "UNKNOWN", "hsts": bool(https.get("hsts")),
            "http_elapsed": http.get("elapsed", 0), "https_elapsed": https.get("elapsed", 0),
            "final_https_url": https.get("final_url"),
        }
        if http.get("error"):
            result["http_result"] = http["error"]
        else:
            location = http.get("location")
            result["http_location"] = requests.compat.urljoin(http_url, location) if location else None
            if http.get("status") in (301, 302, 303, 307, 308) and location:
                result["http_result"] = (
                    "REDIRECTS_TO_HTTPS"
                    if urlparse(result["http_location"]).scheme.lower() == "https"
                    else "REDIRECTS_TO_NON_HTTPS"
                )
            elif (http.get("status") or 999) < 400:
                result["http_result"] = "PLAINTEXT_CONTENT_AVAILABLE"
            else:
                result["http_result"] = f"HTTP_{http.get('status')}"
        result["https_result"] = https.get("error") or (
            "REACHABLE" if (https.get("status") or 999) < 500
            else f"HTTP_{https.get('status')}"
        )
        total = max(result["http_elapsed"], result["https_elapsed"])
        self.progress.emit(
            f"{component.title()} Web Adaptor {number}: HTTP={result['http_result']} "
            f"({result['http_elapsed']:.2f}s), HTTPS={result['https_result']} "
            f"({result['https_elapsed']:.2f}s), selesai {total:.2f}s."
        )
        return result
    @staticmethod
    def certificate_aliases(collection):
        if not isinstance(collection, dict):
            return []
        candidates = find_value(
            collection, "sslCertificates", "certificates", "aliases", "items", default=[]
        )
        if isinstance(candidates, dict):
            candidates = list(candidates.values())
        aliases = []
        for item in candidates if isinstance(candidates, list) else []:
            if isinstance(item, str):
                aliases.append(item)
            elif isinstance(item, dict):
                alias = find_value(item, "aliasName", "alias", "name", default=None)
                if alias:
                    aliases.append(str(alias))
        return list(dict.fromkeys(aliases))
    @staticmethod
    def certificate_date(value):
        if value in (None, "", "Tidak tersedia"):
            return None
        text = str(value).strip()
        for zone in ("WIB", "WITA", "WIT", "UTC"):
            text = text.replace(f" {zone} ", " +0700 " if zone == "WIB" else " +0800 " if zone == "WITA" else " +0900 " if zone == "WIT" else " +0000 ")
        try:
            dt = parsedate_to_datetime(text)
            return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)
        except Exception:
            pass
        for fmt in ("%a %b %d %H:%M:%S %z %Y", "%Y-%m-%dT%H:%M:%S.%f%z", "%Y-%m-%dT%H:%M:%S%z"):
            try:
                return datetime.strptime(text, fmt)
            except ValueError:
                continue
        return None
    @staticmethod
    def normalized_dn(value):
        return ",".join(part.strip().lower() for part in str(value or "").split(",") if part.strip())
    @classmethod
    def normalize_certificate_detail(cls, detail, fallback_alias=None):
        detail = detail if isinstance(detail, dict) else {}
        sans = find_value(detail, "subjectAlternativeNames", "subjectAltNames", "sans", default=[])
        if isinstance(sans, dict):
            sans = list(sans.values())
        if isinstance(sans, str):
            sans = [x.strip() for x in sans.split(",") if x.strip()]
        normalized_sans = []
        for value in sans if isinstance(sans, list) else []:
            if isinstance(value, dict):
                value = find_value(value, "DNSName", "dnsName", "value", default=None)
            value = str(value or "").strip()
            if value.lower().startswith("dnsname:"):
                value = value.split(":", 1)[1].strip()
            if value:
                normalized_sans.append(value)
        valid_from_raw = find_value(detail, "validFrom", "notBefore", default=None)
        valid_until_raw = find_value(detail, "validUntil", "notAfter", default=None)
        valid_from = cls.certificate_date(valid_from_raw)
        valid_until = cls.certificate_date(valid_until_raw)
        now = datetime.now(timezone.utc)
        days = int((valid_until.astimezone(timezone.utc) - now).total_seconds() // 86400) if valid_until else None
        issuer = find_value(detail, "issuer", default=None)
        subject = find_value(detail, "subject", default=None)
        entry_type = find_value(detail, "entryType", default=None)
        key_size = find_value(detail, "keySize", default=None)
        signature = find_value(detail, "signatureAlgorithm", default=None)
        key_usage = find_value(detail, "keyUsage", default=[])
        if isinstance(key_usage, str):
            key_usage = [x.strip() for x in key_usage.split(",") if x.strip()]
        return {
            "alias": find_value(detail, "aliasName", "alias", "name", default=fallback_alias),
            "entry_type": entry_type, "issuer": issuer, "subject": subject,
            "sans": normalized_sans, "valid_from": valid_from_raw, "valid_until": valid_until_raw,
            "days_remaining": days, "key_algorithm": find_value(detail, "keyAlgorithm", default=None),
            "key_size": key_size, "serial_number": find_value(detail, "serialNumber", default=None),
            "signature_algorithm": signature, "key_usage": key_usage,
            "sha256": find_value(detail, "sha256Fingerprint", "SHA256Fingerprint", default=None),
            "issuer_dn": cls.normalized_dn(issuer), "subject_dn": cls.normalized_dn(subject),
            "not_yet_valid": bool(valid_from and valid_from.astimezone(timezone.utc) > now),
            "expired": bool(valid_until and valid_until.astimezone(timezone.utc) < now),
        }
    @staticmethod
    def hostname_matches_pattern(hostname, pattern):
        host = str(hostname or "").strip().lower().rstrip(".")
        pat = str(pattern or "").strip().lower().rstrip(".")
        if not host or not pat:
            return False
        if pat.startswith("*."):
            suffix = pat[1:]
            return host.endswith(suffix) and host.count(".") == pat.count(".")
        return host == pat
    @staticmethod
    def san_target(machine_name=None, admin_url=None):
        admin_hostname = urlparse(str(admin_url or "")).hostname
        if admin_hostname and "." in admin_hostname:
            return admin_hostname.lower(), "Native Admin URL", "FQDN"
        registered = str(machine_name or "").strip().rstrip(".")
        if registered and "." in registered:
            return registered.lower(), "Registered machine name", "FQDN"
        if admin_hostname:
            return admin_hostname.lower(), "Native Admin URL", "SHORT_HOSTNAME"
        if registered:
            return registered.lower(), "Registered machine name", "SHORT_HOSTNAME"
        return None, "Unavailable", "UNAVAILABLE"
    @classmethod
    def assess_leaf_certificate(cls, cert, hostname=None, hostname_kind="FQDN"):
        if not cert or not cert.get("alias"):
            return "UNKNOWN", ["Active certificate detail unavailable"], "UNKNOWN"
        reasons = []
        issuer = cert.get("issuer_dn", "")
        subject = cert.get("subject_dn", "")
        self_signed = bool(subject and issuer and subject == issuer) or "selfsignedcertificate" in issuer.replace(" ", "") or "self signed certificate" in issuer
        if str(cert.get("entry_type", "")).lower() != "privatekeyentry":
            reasons.append("Active certificate is not PrivateKeyEntry")
        if self_signed:
            reasons.append("Active leaf certificate is self-signed")
        if cert.get("expired"):
            reasons.append("Certificate has expired")
        if cert.get("not_yet_valid"):
            reasons.append("Certificate is not valid yet")
        try:
            if str(cert.get("key_algorithm", "")).upper() == "RSA" and int(cert.get("key_size") or 0) < 2048:
                reasons.append("RSA key size is below 2048 bits")
        except (TypeError, ValueError):
            pass
        signature = str(cert.get("signature_algorithm", "")).lower()
        if signature and ("md5" in signature or "sha1" in signature):
            reasons.append("Weak certificate signature algorithm")
        sans = cert.get("sans", [])
        if hostname_kind != "FQDN" or not hostname:
            san_status = "UNKNOWN"
        elif not sans:
            san_status = "UNKNOWN"
        elif any(cls.hostname_matches_pattern(hostname, san) for san in sans):
            san_status = "PASS"
        else:
            san_status = "FAIL"
            reasons.append(f"SAN does not cover native hostname/FQDN {hostname}")
        if reasons:
            return "FAIL", reasons, san_status
        if cert.get("days_remaining") is None:
            return "UNKNOWN", ["Certificate validity date could not be parsed"], san_status
        if san_status == "UNKNOWN":
            return "UNKNOWN", [
                "Native Admin URL only uses a short hostname or the agreed native FQDN cannot be determined automatically; DNS, hosts file, DNS alias, and installation-time machine naming are outside tool visibility"
            ], san_status
        if cert.get("days_remaining") < 60:
            return "PASS WITH NOTE", ["Certificate expires in less than 60 days"], san_status
        return "PASS", [], san_status
    @classmethod
    def resolve_certificate_chain(cls, active_cert, certificate_details):
        details = [x for x in certificate_details if isinstance(x, dict)]
        intermediate = next((x for x in details if x.get("subject_dn") and x.get("subject_dn") == active_cert.get("issuer_dn") and x.get("alias") != active_cert.get("alias")), None)
        root = None
        if intermediate:
            root = next((x for x in details if x.get("subject_dn") and x.get("subject_dn") == intermediate.get("issuer_dn") and x.get("alias") != intermediate.get("alias")), None)
        chain_status = "PASS" if intermediate and root else "UNKNOWN"
        if intermediate and (intermediate.get("expired") or intermediate.get("not_yet_valid")):
            chain_status = "FAIL"
        if root and (root.get("expired") or root.get("not_yet_valid")):
            chain_status = "FAIL"
        return {"intermediate": intermediate, "root": root, "status": chain_status}
    def collect_web_tier_certificates(self, adaptors, component):
        records = []
        for adaptor in adaptors or []:
            url = adaptor.get("url") if isinstance(adaptor, dict) else None
            if not url:
                continue
            parsed = urlparse(url)
            key = (str(parsed.hostname or "").lower(), parsed.port or 443)
            if key not in self.web_tier_tls_cache:
                self.progress.emit(f"Web tier certificate: TLS handshake {key[0]}:{key[1]}...")
                self.web_tier_tls_cache[key] = tls_certificate_evidence(url)
            record = dict(self.web_tier_tls_cache[key])
            record.update({"listener_key": key, "registered_url": url, "component": component,
                           "context_path": parsed.path or "/"})
            records.append(record)
        return records
    def portal_certificate_inventory(self, session, token):
        result = {"status": "UNKNOWN", "error": None, "aliases": [], "active_alias": None,
                  "active_certificate": None, "certificate_details": [], "chain": {"status": "UNKNOWN"},
                  "machines": []}
        try:
            self.progress.emit("Portal certificate: membaca registered Portal machines...")
            machine_collection = self.request_json(
                session, "GET", f"{self.portal_admin_url}/machines",
                params={"token": token, "f": "json", "_ts": str(int(time.time() * 1000))},
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
            raw_machines = find_value(machine_collection, "machines", "items", default=[])
            if isinstance(raw_machines, dict):
                raw_machines = list(raw_machines.values())
            names = []
            for item in raw_machines if isinstance(raw_machines, list) else []:
                name = item if isinstance(item, str) else find_value(item, "machineName", "name", default=None)
                if name:
                    names.append(str(name))
            for name in list(dict.fromkeys(names)):
                machine = {"machine_name": name, "admin_url": None, "platform": None, "role": None, "error": None}
                try:
                    detail = self.request_json(
                        session, "GET", f"{self.portal_admin_url}/machines/{quote(name, safe='')}",
                        params={"token": token, "f": "json", "_ts": str(int(time.time() * 1000))},
                        headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
                    )
                    machine["admin_url"] = find_value(detail, "adminURL", "adminUrl", default=None)
                    machine["platform"] = find_value(detail, "platform", default=None)
                    machine["role"] = find_value(detail, "role", default=None)
                except Exception as error:
                    machine["error"] = f"{type(error).__name__}: {error}"
                target, source, kind = self.san_target(name, machine.get("admin_url"))
                machine.update({"san_target": target, "san_target_source": source, "san_target_kind": kind})
                result["machines"].append(machine)
            self.progress.emit("Portal certificate: membaca SSL certificate inventory dan active alias...")
            collection = self.request_json(session, "GET", f"{self.portal_admin_url}/security/sslCertificates",
                params={"token": token, "f": "json", "_ts": str(int(time.time()*1000))}, headers={"Cache-Control":"no-cache","Pragma":"no-cache"})
            result["aliases"] = self.certificate_aliases(collection)
            result["active_alias"] = find_value(collection, "webServerCertificate", "webServerCertificateAlias", default=None)
            result["protocols"] = find_value(collection, "sslProtocols", "webServerSSLProtocols", default=None)
            result["hsts"] = find_value(collection, "HSTSEnabled", "hstsEnabled", default=None)
            for alias in result["aliases"]:
                try:
                    raw = self.request_json(session, "GET", f"{self.portal_admin_url}/security/sslCertificates/{quote(str(alias), safe='')}",
                        params={"token":token,"f":"json","_ts":str(int(time.time()*1000))}, headers={"Cache-Control":"no-cache","Pragma":"no-cache"})
                    result["certificate_details"].append(self.normalize_certificate_detail(raw, alias))
                except ArcGISRequestError:
                    continue
            active = next((x for x in result["certificate_details"] if str(x.get("alias","")).lower() == str(result["active_alias"] or "").lower()), None)
            if not active and result["active_alias"]:
                raw = self.request_json(session, "GET", f"{self.portal_admin_url}/security/sslCertificates/{quote(str(result['active_alias']), safe='')}", params={"token":token,"f":"json"})
                active = self.normalize_certificate_detail(raw, result["active_alias"])
                result["certificate_details"].append(active)
            result["active_certificate"] = active
            machine_states = []
            for machine in result["machines"]:
                status, reasons, san_status = self.assess_leaf_certificate(
                    active, machine.get("san_target"), machine.get("san_target_kind")
                )
                machine.update({"status": status, "reasons": reasons, "san_status": san_status})
                machine_states.append(status)
            if not machine_states:
                result["status"] = "UNKNOWN"
                result["reasons"] = ["Portal machine inventory is unavailable"]
            elif "FAIL" in machine_states:
                result["status"] = "FAIL"
            elif "UNKNOWN" in machine_states:
                result["status"] = "UNKNOWN"
            elif "PASS WITH NOTE" in machine_states:
                result["status"] = "PASS WITH NOTE"
            else:
                result["status"] = "PASS"
            result["reasons"] = list(dict.fromkeys(
                reason for machine in result["machines"] for reason in machine.get("reasons", [])
            ))
            result["chain"] = self.resolve_certificate_chain(active or {}, result["certificate_details"])
            if result["status"] in ("PASS", "PASS WITH NOTE") and result["chain"]["status"] == "FAIL":
                result["status"] = "FAIL"
        except Exception as error:
            result["error"] = f"{type(error).__name__}: {error}"
        return result
    def server_machine_certificates(self, session, token):
        output = {"status": "UNKNOWN", "error": None, "machines": []}
        try:
            self.progress.emit("Server certificate: membaca seluruh machine dalam site...")
            collection = self.request_json(session, "GET", f"{self.server_admin_url}/machines",
                params={"token":token,"f":"json","_ts":str(int(time.time()*1000))}, headers={"Cache-Control":"no-cache","Pragma":"no-cache"})
            raw_machines = find_value(collection, "machines", "items", default=[])
            if isinstance(raw_machines, dict): raw_machines = list(raw_machines.values())
            names=[]
            for item in raw_machines if isinstance(raw_machines,list) else []:
                name = item if isinstance(item,str) else find_value(item,"machineName","name",default=None)
                if name: names.append(str(name))
            for name in list(dict.fromkeys(names)):
                machine={"machine_name":name,"status":"UNKNOWN","error":None,"certificate_details":[],"chain":{"status":"UNKNOWN"}}
                try:
                    base=f"{self.server_admin_url}/machines/{quote(name, safe='')}"
                    props=self.request_json(session,"GET",base,params={"token":token,"f":"json"})
                    machine["admin_url"]=find_value(props,"adminURL","adminUrl",default=None)
                    machine["ssl_enabled"]=find_value(props,"webServerSSLEnabled","sslEnabled",default=None)
                    machine["active_alias"]=find_value(props,"webServerCertificate","webServerCertificateAlias",default=None)
                    certs=self.request_json(session,"GET",f"{base}/sslcertificates",params={"token":token,"f":"json"})
                    machine["aliases"]=self.certificate_aliases(certs)
                    for alias in machine["aliases"]:
                        try:
                            raw=self.request_json(session,"GET",f"{base}/sslcertificates/{quote(str(alias),safe='')}",params={"token":token,"f":"json"})
                            machine["certificate_details"].append(self.normalize_certificate_detail(raw,alias))
                        except ArcGISRequestError:
                            continue
                    active=next((x for x in machine["certificate_details"] if str(x.get("alias","")).lower()==str(machine["active_alias"] or "").lower()),None)
                    if not active and machine["active_alias"]:
                        raw=self.request_json(session,"GET",f"{base}/sslcertificates/{quote(str(machine['active_alias']),safe='')}",params={"token":token,"f":"json"})
                        active=self.normalize_certificate_detail(raw,machine["active_alias"]); machine["certificate_details"].append(active)
                    machine["active_certificate"]=active
                    target, source, kind = self.san_target(name, machine.get("admin_url"))
                    machine.update({"san_target": target, "san_target_source": source, "san_target_kind": kind})
                    machine["status"], machine["reasons"], machine["san_status"] = self.assess_leaf_certificate(active, target, kind)
                    if str(machine.get("ssl_enabled")).lower()=="false":
                        machine["status"]="FAIL"; machine.setdefault("reasons",[]).append("Native web server SSL is disabled")
                    machine["chain"]=self.resolve_certificate_chain(active or {},machine["certificate_details"])
                    if machine["status"] in ("PASS","PASS WITH NOTE") and machine["chain"]["status"]=="FAIL": machine["status"]="FAIL"
                except Exception as error:
                    machine["error"]=f"{type(error).__name__}: {error}"
                output["machines"].append(machine)
            states=[m.get("status","UNKNOWN") for m in output["machines"]]
            output["status"]="FAIL" if "FAIL" in states else "UNKNOWN" if not states or "UNKNOWN" in states else "PASS WITH NOTE" if "PASS WITH NOTE" in states else "PASS"
        except Exception as error:
            output["error"]=f"{type(error).__name__}: {error}"
        return output

    def portal_connection_data(self, session):
        """Lightweight Portal authentication and Administrator API check."""
        token = self.ensure_portal_token(session)
        self.progress.emit("Portal: menguji authenticated Portal Self...")
        info = self.request_json(
            session, "GET", f"{self.portal_base_url}/sharing/rest/portals/self",
            params={"token": token, "f": "json", "_ts": str(int(time.time() * 1000))},
            headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
        )
        self.progress.emit("Portal: menguji Administrator API dan membaca version...")
        root = self.request_json(
            session, "GET", self.portal_admin_url,
            params={"token": token, "f": "json", "_ts": str(int(time.time() * 1000))},
            headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
        )
        version = root.get("version", root.get("currentVersion", "Tidak tersedia"))
        self.progress.emit("Portal: lightweight connection test selesai.")
        return {
            "success": True, "connection_only": True,
            "base_url": self.portal_base_url,
            "version": version, "version_available": version != "Tidak tersedia",
            "name": info.get("name", "Tidak tersedia"),
            "id": info.get("id", "Tidak tersedia"),
            "access": info.get("access", "Tidak diuji pada Test Connection"),
            "all_ssl": info.get("allSSL", "Tidak diuji pada Test Connection"),
        }

    def server_connection_data(self, session):
        """Lightweight Server authentication and Administrator API check."""
        token = self.ensure_server_token(session)
        self.progress.emit("Server: menguji Administrator API dan membaca site version...")
        info = self.request_json(
            session, "GET", self.server_admin_url,
            params={"token": token, "f": "json", "_ts": str(int(time.time() * 1000))},
            headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
        )
        self.progress.emit("Server: lightweight connection test selesai.")
        return {
            "success": True, "connection_only": True,
            "version": info.get("currentVersion", info.get("version", "Tidak tersedia")),
            "full_version": info.get("fullVersion", "Tidak tersedia"),
            "authentication_tier": "Verified by administrator token",
        }

    def portal_data(self, session):
        token = self.ensure_portal_token(session)
        self.progress.emit("Portal: membaca informasi organisasi live...")
        try:
            info = self.request_json(
                session, "GET", f"{self.portal_base_url}/sharing/rest/portals/self",
                params={
                    "token": token,
                    "f": "json",
                    "_ts": str(int(time.time() * 1000)),
                },
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
            self.progress.emit("Portal: membaca security configuration live...")
            security = self.request_json(
                session, "GET", f"{self.portal_admin_url}/security/config",
                params={
                    "token": token,
                    "f": "json",
                    "_ts": str(int(time.time() * 1000)),
                },
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
            self.progress.emit("Portal: membaca System Properties live...")
            portal_system_properties = self.request_json(
                session, "GET", f"{self.portal_admin_url}/system/properties",
                params={
                    "token": token,
                    "f": "json",
                    "_ts": str(int(time.time() * 1000)),
                },
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
        except ArcGISRequestError as error:
            # ArcGIS menggunakan 498/499 untuk token invalid/required.
            if "498" not in str(error) and "499" not in str(error):
                raise
            self.progress.emit("Portal: menerima error 498/499; membuat ulang referer-bound token...")
            self.portal_token = None
            self.portal_token_expires = 0
            token = self.generate_portal_token(session)
            info = self.request_json(
                session, "GET", f"{self.portal_base_url}/sharing/rest/portals/self",
                params={
                    "token": token,
                    "f": "json",
                    "_ts": str(int(time.time() * 1000)),
                },
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
            security = self.request_json(
                session, "GET", f"{self.portal_admin_url}/security/config",
                params={
                    "token": token,
                    "f": "json",
                    "_ts": str(int(time.time() * 1000)),
                },
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
            self.progress.emit("Portal: membaca System Properties live...")
            portal_system_properties = self.request_json(
                session, "GET", f"{self.portal_admin_url}/system/properties",
                params={
                    "token": token,
                    "f": "json",
                    "_ts": str(int(time.time() * 1000)),
                },
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
        self.progress.emit("Portal: membaca New Member Default Settings live...")
        try:
            user_default_settings = self.request_json(
                session, "GET", f"{self.portal_base_url}/sharing/rest/portals/self/userDefaultSettings",
                params={"token": token, "f": "json", "_ts": str(int(time.time() * 1000))},
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
            user_license_types = self.request_json(
                session, "GET", f"{self.portal_base_url}/sharing/rest/portals/self/userLicenseTypes",
                params={"token": token, "f": "json", "_ts": str(int(time.time() * 1000))},
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
            portal_roles = self.request_json(
                session, "GET", f"{self.portal_base_url}/sharing/rest/portals/self/roles",
                params={"token": token, "start": 1, "num": 100, "f": "json", "_ts": str(int(time.time() * 1000))},
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
            user_default_settings_error = None
        except ArcGISRequestError as error:
            user_default_settings = {}
            user_license_types = {}
            portal_roles = {}
            user_default_settings_error = str(error).splitlines()[0]
            self.progress.emit("Portal: New Member Default Settings tidak dapat dibaca; status UNKNOWN.")

        portal_version = "Tidak tersedia"
        portal_version_available = False
        try:
            self.progress.emit("Portal: membaca version information untuk connection summary...")
            portal_root = self.request_json(
                session, "GET", self.portal_admin_url,
                params={
                    "token": token,
                    "f": "json",
                    "_ts": str(int(time.time() * 1000)),
                },
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
            portal_version = portal_root.get(
                "version", portal_root.get("currentVersion", "Tidak tersedia")
            )
            portal_version_available = portal_version != "Tidak tersedia"
        except ArcGISRequestError:
            # Connection dan hardening endpoint utama tetap valid. Version bersifat
            # display-only sehingga kegagalannya menghasilkan status parsial, bukan gagal.
            self.progress.emit(
                "Portal: version information tidak tersedia; connection summary ditandai parsial."
            )
        try:
            portal_web_adaptors = self.discover_web_adaptors(
                session, self.portal_admin_url, token, "portal"
            )
            portal_web_adaptors_error = None
        except ArcGISRequestError as error:
            portal_web_adaptors = []
            portal_web_adaptors_error = str(error).splitlines()[0]
            self.progress.emit("Portal: Web Adaptor inventory tidak dapat dibaca; status HTTPS parsial.")
        public_certificates = self.collect_web_tier_certificates(portal_web_adaptors, "Portal") if self.run_network_probes else []
        portal_certificate = self.portal_certificate_inventory(session, token) if self.run_network_probes else {"status": "NOT_RUN"}
        return {
            "success": True,
            "base_url": self.portal_base_url,
            "version": portal_version,
            "version_available": portal_version_available,
            "name": info.get("name", "Tidak tersedia"),
            "id": info.get("id", "Tidak tersedia"),
            "access": info.get("access", "Tidak tersedia"),
            "can_share_bing_public_present": "canShareBingPublic" in info,
            "can_share_bing_public": info.get("canShareBingPublic"),
            "all_ssl": info.get("allSSL", "Tidak tersedia"),
            "web_adaptors": portal_web_adaptors,
            "web_adaptors_error": portal_web_adaptors_error,
            "public_certificates": public_certificates,
            "portal_certificate": portal_certificate,
            "portal_directory_disabled": find_value(security, "disableServicesDirectory"),
            "automatic_account_creation": find_value(security, "enableAutomaticAccountCreation"),
            "system_properties": portal_system_properties,
            "portal_servlet_properties": {
                name: {
                    "present": name in portal_system_properties,
                    "value": portal_system_properties.get(name),
                }
                for name in (
                    "disableLegendServlet",
                    "disablePrintServlet",
                    "disableWFSServlet",
                )
            },
            "builtin_self_creation_present": "disableSignup" in portal_system_properties,
            "builtin_self_creation_disabled": portal_system_properties.get("disableSignup"),
            "public_profile_setting_present": "updateUserProfileDisabled" in info,
            "public_profile_update_disabled": info.get("updateUserProfileDisabled"),
            "user_default_settings": user_default_settings,
            "user_default_settings_error": user_default_settings_error,
            "user_license_types": user_license_types,
            "portal_roles": portal_roles,
            "portal_properties": info.get("portalProperties", {}),
            "social_media_links_present": (
                isinstance(info.get("portalProperties"), dict)
                and "showSocialMediaLinks" in info.get("portalProperties", {})
            ),
            "social_media_links_raw": (
                info.get("portalProperties", {}).get("showSocialMediaLinks")
                if isinstance(info.get("portalProperties"), dict) else None
            ),
        }

    def server_data(self, session):
        token = self.ensure_server_token(session)
        try:
            self.progress.emit("Server: membaca site information live...")
            info = self.request_json(
                session, "GET", self.server_admin_url,
                params={
                    "token": token,
                    "f": "json",
                    "_ts": str(int(time.time() * 1000)),
                },
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
            self.progress.emit("Server: membaca security configuration live...")
            security = self.request_json(
                session, "GET", f"{self.server_admin_url}/security/config",
                params={
                    "token": token,
                    "f": "json",
                    "_ts": str(int(time.time() * 1000)),
                },
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
            self.progress.emit("Server: membaca Services Directory live...")
            services_directory = self.request_json(
                session, "GET",
                f"{self.server_admin_url}/system/handlers/rest/servicesdirectory",
                params={
                    "token": token,
                    "f": "json",
                    "_ts": str(int(time.time() * 1000)),
                },
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
            self.progress.emit("Server: membaca System Properties live...")
            system_properties = self.request_json(
                session, "GET", f"{self.server_admin_url}/system/properties",
                params={
                    "token": token,
                    "f": "json",
                    "_ts": str(int(time.time() * 1000)),
                },
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
            self.progress.emit("Server: membaca Token Manager Configuration live...")
            token_configuration = self.request_json(
                session, "GET", f"{self.server_admin_url}/security/tokens",
                params={
                    "token": token,
                    "f": "json",
                    "_ts": str(int(time.time() * 1000)),
                },
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
        except ArcGISRequestError as error:
            if "498" not in str(error) and "499" not in str(error):
                raise
            self.progress.emit("Server: token ditolak; memperbarui token lalu mencoba sekali lagi...")
            self.server_token = None
            self.server_token_expires = 0
            token = self.generate_server_token(session)
            info = self.request_json(
                session, "GET", self.server_admin_url,
                params={
                    "token": token,
                    "f": "json",
                    "_ts": str(int(time.time() * 1000)),
                },
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
            security = self.request_json(
                session, "GET", f"{self.server_admin_url}/security/config",
                params={
                    "token": token,
                    "f": "json",
                    "_ts": str(int(time.time() * 1000)),
                },
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
            services_directory = self.request_json(
                session, "GET",
                f"{self.server_admin_url}/system/handlers/rest/servicesdirectory",
                params={
                    "token": token,
                    "f": "json",
                    "_ts": str(int(time.time() * 1000)),
                },
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
            self.progress.emit("Server: membaca System Properties live...")
            system_properties = self.request_json(
                session, "GET", f"{self.server_admin_url}/system/properties",
                params={
                    "token": token,
                    "f": "json",
                    "_ts": str(int(time.time() * 1000)),
                },
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
            self.progress.emit("Server: membaca Token Manager Configuration live...")
            token_configuration = self.request_json(
                session, "GET", f"{self.server_admin_url}/security/tokens",
                params={
                    "token": token,
                    "f": "json",
                    "_ts": str(int(time.time() * 1000)),
                },
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
            )
        try:
            server_web_adaptors = self.discover_web_adaptors(
                session, self.server_admin_url, token, "server"
            )
            server_web_adaptors_error = None
        except ArcGISRequestError as error:
            server_web_adaptors = []
            server_web_adaptors_error = str(error).splitlines()[0]
            self.progress.emit("Server: Web Adaptor inventory tidak dapat dibaca; status HTTPS parsial.")
        public_certificates = self.collect_web_tier_certificates(server_web_adaptors, "Server") if self.run_network_probes else []
        server_machine_certificates = self.server_machine_certificates(session, token) if self.run_network_probes else {"status": "NOT_RUN", "machines": []}
        return {
            "success": True,
            "version": info.get("currentVersion", info.get("version", "Tidak tersedia")),
            "full_version": info.get("fullVersion", "Tidak tersedia"),
            "protocol": find_value(security, "protocol", default="Tidak tersedia"),
            "http_enabled": find_value(security, "httpEnabled", default="Tidak tersedia"),
            "ssl_enabled": find_value(security, "sslEnabled", "httpsEnabled", default="Tidak tersedia"),
            "hsts_enabled": find_value(security, "HSTSEnabled", "hstsEnabled", default="Tidak tersedia"),
            "web_adaptors": server_web_adaptors,
            "web_adaptors_error": server_web_adaptors_error,
            "public_certificates": public_certificates,
            "server_machine_certificates": server_machine_certificates,
            "authentication_tier": find_value(security, "authenticationTier"),
            "services_directory_enabled": find_value(
                services_directory, "servicesDirEnabled", "servicesDirectoryEnabled", "enabled"
            ),
            "callback_functions_enabled": find_value(
                services_directory, "callbackFunctionsEnabled"
            ),
            "allowed_origins": find_value(services_directory, "allowedOrigins"),
            "system_properties": system_properties,
            "standardized_queries_present": "standardizedQueries" in system_properties,
            "standardized_queries_raw": system_properties.get("standardizedQueries"),
            "feature_service_xss_present": "featureServiceXSSFilter" in system_properties,
            "feature_service_xss_raw": system_properties.get("featureServiceXSSFilter"),
            "token_configuration": token_configuration,
            "allow_http_get_present": "allowHttpGet" in token_configuration.get("properties", {}),
            "allow_http_get_raw": token_configuration.get("properties", {}).get("allowHttpGet"),
        }

    def run(self):
        result = {
            "purpose": self.purpose,
            "mode": self.mode,
            "portal": None,
            "server": None,
        }
        session = self.session()
        try:
            if self.mode in ("portal", "all"):
                try:
                    result["portal"] = (
                        self.portal_connection_data(session)
                        if self.purpose == "connection" else self.portal_data(session)
                    )
                except Exception as error:
                    result["portal"] = {"success": False, "error": str(error)}
            if self.mode in ("server", "all"):
                try:
                    result["server"] = (
                        self.server_connection_data(session)
                        if self.purpose == "connection" else self.server_data(session)
                    )
                except Exception as error:
                    result["server"] = {"success": False, "error": str(error)}
            result["token_cache"] = {
                "portal_token": self.portal_token,
                "portal_token_expires": self.portal_token_expires,
                "server_token": self.server_token,
                "server_token_expires": self.server_token_expires,
            }
            self.completed.emit(result)
        finally:
            session.close()
            self.finished.emit()


class ServerDirectoryHardeningWorker(QObject):
    progress = Signal(str)
    completed = Signal(dict)
    finished = Signal()

    def __init__(
        self, action, control, server_admin_url, server_username, server_password,
        verify_ssl, server_token=None, server_token_expires=0,
        backup_file=None,
    ):
        super().__init__()
        self.action = action
        self.control = control
        self.server_admin_url = normalize_url(server_admin_url)
        self.server_username = server_username.strip()
        self.server_password = server_password
        self.verify_ssl = verify_ssl
        self.server_token = server_token
        self.server_token_expires = float(server_token_expires or 0)
        self.backup_file = backup_file

    def request_json(self, session, method, url, **kwargs):
        try:
            response = session.request(method, url, timeout=60, **kwargs)
        except requests.exceptions.RequestException as error:
            raise ArcGISRequestError(f"Request gagal.\nEndpoint: {url}\nDetail: {error}") from error
        if response.status_code >= 400:
            raise ArcGISRequestError(
                f"Endpoint mengembalikan HTTP {response.status_code}.\nEndpoint: {url}"
            )
        try:
            result = response.json()
        except ValueError as error:
            raise ArcGISRequestError(
                f"Endpoint tidak mengembalikan JSON.\nEndpoint: {url}"
            ) from error
        if isinstance(result, dict) and result.get("error"):
            err = result["error"]
            raise ArcGISRequestError(
                f"ArcGIS API menolak request.\nCode: {err.get('code', '-')}\n"
                f"Message: {err.get('message', '-')}\nDetails: {err.get('details', '-')}"
            )
        return result

    def token_valid(self):
        return bool(self.server_token) and self.server_token_expires > time.time() + 120

    def ensure_token(self, session):
        if self.token_valid():
            self.progress.emit("Apply: menggunakan Server token yang masih valid dari memory...")
            return self.server_token
        self.progress.emit("Apply: membuat Server token baru...")
        result = self.request_json(
            session, "POST", f"{self.server_admin_url}/generateToken",
            data={
                "username": self.server_username,
                "password": self.server_password,
                "client": "requestip",
                "expiration": "30",
                "f": "json",
            },
        )
        token = result.get("token")
        if not token:
            raise ArcGISRequestError("ArcGIS Server tidak mengembalikan administrator token.")
        expires_ms = result.get("expires", 0)
        self.server_token = token
        self.server_token_expires = (
            float(expires_ms) / 1000 if expires_ms else time.time() + 25 * 60
        )
        return token

    def read_services_directory(self, session, token):
        return self.request_json(
            session,
            "GET",
            f"{self.server_admin_url}/system/handlers/rest/servicesdirectory",
            params={
                    "token": token,
                    "f": "json",
                    "_ts": str(int(time.time() * 1000)),
                },
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
        )

    def read_system_properties(self, session, token):
        return self.request_json(
            session,
            "GET",
            f"{self.server_admin_url}/system/properties",
            params={
                "token": token,
                "f": "json",
                "_ts": str(int(time.time() * 1000)),
            },
            headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
        )

    def update_system_properties(self, session, token, properties):
        # Endpoint ini memperbarui kumpulan Server System Properties. Selalu kirim
        # object hasil GET+merge agar property existing tidak hilang.
        return self.request_json(
            session,
            "POST",
            f"{self.server_admin_url}/system/properties/update",
            data={
                "properties": json.dumps(properties, ensure_ascii=False),
                "token": token,
                "f": "json",
            },
        )

    def standardized_queries_state(self, properties):
        # ArcGIS efektif mengaktifkan standardized queries saat property tidak ada,
        # tetapi baseline utility mensyaratkan bukti konfigurasi eksplisit.
        if "standardizedQueries" not in properties:
            return False, "Not configured (effective default enabled)"
        raw = properties.get("standardizedQueries")
        value = str(raw).strip().lower()
        if value == "true":
            return True, "Enabled (explicit)"
        if value == "false":
            return False, "Disabled (explicit)"
        return None, f"Unknown ({raw})"

    def create_system_properties_backup(self, current):
        execution_id = datetime.now().strftime("%Y%m%d-%H%M%S")
        base_dir = Path.cwd() / "backups" / execution_id
        base_dir.mkdir(parents=True, exist_ok=False)
        backup_file = base_dir / "server-system-properties-before.json"
        metadata_file = base_dir / "metadata.json"
        backup_file.write_text(
            json.dumps(current, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        metadata = {
            "execution_id": execution_id,
            "created_at": datetime.now().astimezone().isoformat(),
            "component": "ArcGIS Server",
            "control": "enable_standardized_queries",
            "server_admin_url": self.server_admin_url,
            "property_was_present": "standardizedQueries" in current,
            "baseline_requires_explicit_true": True,
            "before": {"standardizedQueries": current.get("standardizedQueries")},
            "target": {"standardizedQueries": "true"},
            "status": "backup_created",
        }
        metadata_file.write_text(
            json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        return backup_file, metadata_file, execution_id

    def apply_standardized_queries(self, session, token):
        self.progress.emit("Apply: membaca seluruh Server System Properties terbaru...")
        current = self.read_system_properties(session, token)
        compliant, current_text = self.standardized_queries_state(current)
        if compliant is True:
            return {
                "success": True,
                "already_compliant": True,
                "verified_value": current_text,
                "server_token": self.server_token,
                "server_token_expires": self.server_token_expires,
            }
        if compliant is None:
            raise ArcGISRequestError(
                "Nilai standardizedQueries tidak dikenali. Apply dihentikan agar tidak "
                f"mengubah konfigurasi ambigu: {current.get('standardizedQueries')}"
            )

        self.progress.emit("Apply: membuat backup lengkap Server System Properties...")
        backup_file, metadata_file, execution_id = (
            self.create_system_properties_backup(current)
        )
        target = dict(current)
        target["standardizedQueries"] = "true"

        self.progress.emit(
            "Apply: merge standardizedQueries=true tanpa menghapus property lain..."
        )
        apply_result = self.update_system_properties(session, token, target)
        self.progress.emit("Verify: membaca ulang Server System Properties live...")
        verified = self.read_system_properties(session, token)
        verified_ok, verified_text = self.standardized_queries_state(verified)
        self.update_metadata(
            metadata_file,
            apply_response=apply_result,
            verified_value=verified.get("standardizedQueries"),
            verified_state=verified_text,
            status="verified_success" if verified_ok is True else "verification_failed",
            completed_at=datetime.now().astimezone().isoformat(),
        )
        if verified_ok is not True:
            raise ArcGISRequestError(
                "Update mendapat respons, tetapi live verification Standardized Queries "
                f"gagal. Kondisi aktual: {verified_text}"
            )
        return {
            "success": True,
            "already_compliant": False,
            "execution_id": execution_id,
            "backup_file": str(backup_file),
            "metadata_file": str(metadata_file),
            "verified_value": verified_text,
            "server_token": self.server_token,
            "server_token_expires": self.server_token_expires,
        }

    def rollback_standardized_queries(self, session, token):
        if not self.backup_file:
            raise ArcGISRequestError("Backup Standardized Queries tidak tersedia.")
        backup_file = Path(self.backup_file)
        if not backup_file.exists():
            raise ArcGISRequestError(f"Backup file tidak ditemukan: {backup_file}")
        previous = json.loads(backup_file.read_text(encoding="utf-8"))
        if not isinstance(previous, dict):
            raise ArcGISRequestError("Backup Server System Properties bukan JSON object.")

        # Rollback exact backup: jika property dahulu tidak ada, object yang dikirim
        # juga tidak memiliki standardizedQueries. Tidak pernah memaksakan false.
        self.progress.emit(
            "Rollback: mengembalikan seluruh Server System Properties dari backup..."
        )
        rollback_result = self.update_system_properties(session, token, previous)
        self.progress.emit("Rollback: melakukan live verification exact backup...")
        verified = self.read_system_properties(session, token)
        verified_ok = verified == previous
        _, verified_text = self.standardized_queries_state(verified)
        if not verified_ok:
            raise ArcGISRequestError(
                "Rollback selesai, tetapi Server System Properties tidak identik dengan "
                "backup. Tidak ada perubahan lanjutan yang dilakukan."
            )
        metadata_file = backup_file.parent / "metadata.json"
        if metadata_file.exists():
            self.update_metadata(
                metadata_file,
                rollback_response=rollback_result,
                rollback_verified_state=verified_text,
                status="rolled_back",
                rolled_back_at=datetime.now().astimezone().isoformat(),
            )
        return {
            "success": True,
            "verified_value": verified_text,
            "backup_file": str(backup_file),
            "server_token": self.server_token,
            "server_token_expires": self.server_token_expires,
        }

    def feature_service_xss_state(self, properties):
        if "featureServiceXSSFilter" not in properties:
            return False, "Not configured (effective default input)"
        raw = properties.get("featureServiceXSSFilter")
        value = str(raw).strip().lower()
        if value == "input":
            return True, "Input scanning (Basic)"
        if value == "inputoutput":
            return True, "Input + output scanning"
        return None, f"Unknown ({raw})"

    def create_feature_service_xss_backup(self, current):
        # Backup hanya state milik kontrol ini. Property lain dibaca ulang saat
        # rollback agar perubahan kontrol lain tidak ikut dibatalkan.
        execution_id = datetime.now().strftime("%Y%m%d-%H%M%S")
        base_dir = Path.cwd() / "backups" / execution_id
        base_dir.mkdir(parents=True, exist_ok=False)
        backup_file = base_dir / "server-feature-service-xss-before.json"
        metadata_file = base_dir / "metadata.json"
        safe_backup = {
            "property_was_present": "featureServiceXSSFilter" in current,
            "featureServiceXSSFilter": current.get("featureServiceXSSFilter"),
        }
        backup_file.write_text(
            json.dumps(safe_backup, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        metadata = {
            "execution_id": execution_id,
            "created_at": datetime.now().astimezone().isoformat(),
            "component": "ArcGIS Server",
            "control": "verify_feature_service_xss_filter_default",
            "server_admin_url": self.server_admin_url,
            "before": safe_backup,
            "target": {"featureServiceXSSFilter": "input"},
            "scope": "default_for_new_feature_services",
            "other_system_properties_modified": False,
            "status": "backup_created",
        }
        metadata_file.write_text(
            json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        return backup_file, metadata_file, execution_id

    def apply_feature_service_xss(self, session, token, target_value="input"):
        target_value = str(target_value).strip()
        if target_value not in ("input", "inputOutput"):
            raise ArcGISRequestError(f"Target featureServiceXSSFilter tidak valid: {target_value}")

        self.progress.emit("Apply: membaca seluruh Server System Properties terbaru...")
        current = self.read_system_properties(session, token)
        compliant, current_text = self.feature_service_xss_state(current)
        current_value = str(current.get("featureServiceXSSFilter", "")).strip().lower()

        if compliant is None:
            raise ArcGISRequestError(
                "Nilai featureServiceXSSFilter tidak dikenali. Apply dihentikan agar tidak "
                f"mengubah konfigurasi ambigu: {current.get('featureServiceXSSFilter')}"
            )

        # Basic tidak pernah menurunkan inputOutput menjadi input.
        if target_value == "input" and current_value in ("input", "inputoutput"):
            return {
                "success": True,
                "already_compliant": True,
                "verified_value": current_text,
                "applied_profile": "Basic",
                "server_token": self.server_token,
                "server_token_expires": self.server_token_expires,
            }
        if target_value == "inputOutput" and current_value == "inputoutput":
            return {
                "success": True,
                "already_compliant": True,
                "verified_value": current_text,
                "applied_profile": "Advanced",
                "server_token": self.server_token,
                "server_token_expires": self.server_token_expires,
            }

        profile = "Advanced" if target_value == "inputOutput" else "Basic"
        self.progress.emit(f"Apply {profile}: membuat backup state Feature Service XSS Filter...")
        backup_file, metadata_file, execution_id = (
            self.create_feature_service_xss_backup(current)
        )
        self.update_metadata(
            metadata_file,
            selected_profile=profile,
            target={"featureServiceXSSFilter": target_value},
        )

        # Control isolation: salin konfigurasi live, ubah satu property saja.
        target = dict(current)
        target["featureServiceXSSFilter"] = target_value
        self.progress.emit(
            f"Apply {profile}: merge hanya featureServiceXSSFilter={target_value}; "
            "property lain dipertahankan..."
        )
        apply_result = self.update_system_properties(session, token, target)
        self.progress.emit("Verify: membaca ulang Server System Properties live...")
        verified = self.read_system_properties(session, token)
        verified_ok, verified_text = self.feature_service_xss_state(verified)
        exact_target = (
            str(verified.get("featureServiceXSSFilter", "")).strip().lower()
            == target_value.lower()
        )
        self.update_metadata(
            metadata_file,
            apply_response=apply_result,
            verified_value=verified.get("featureServiceXSSFilter"),
            verified_state=verified_text,
            status="verified_success" if verified_ok is True and exact_target else "verification_failed",
            completed_at=datetime.now().astimezone().isoformat(),
        )
        if verified_ok is not True or not exact_target:
            raise ArcGISRequestError(
                f"Update mendapat respons, tetapi live verification "
                f"featureServiceXSSFilter={target_value} gagal. Kondisi aktual: {verified_text}"
            )
        return {
            "success": True,
            "already_compliant": False,
            "execution_id": execution_id,
            "backup_file": str(backup_file),
            "metadata_file": str(metadata_file),
            "verified_value": verified_text,
            "applied_profile": profile,
            "server_token": self.server_token,
            "server_token_expires": self.server_token_expires,
        }

    def rollback_feature_service_xss(self, session, token):
        if not self.backup_file:
            raise ArcGISRequestError("Backup Feature Service XSS Filter tidak tersedia.")
        backup_file = Path(self.backup_file)
        if not backup_file.exists():
            raise ArcGISRequestError(f"Backup file tidak ditemukan: {backup_file}")
        previous = json.loads(backup_file.read_text(encoding="utf-8"))
        if not isinstance(previous, dict):
            raise ArcGISRequestError("Backup Feature Service XSS Filter bukan JSON object.")

        self.progress.emit(
            "Rollback: membaca System Properties live agar kontrol lain tetap dipertahankan..."
        )
        current = self.read_system_properties(session, token)
        target = dict(current)
        if previous.get("property_was_present"):
            target["featureServiceXSSFilter"] = previous.get("featureServiceXSSFilter")
        else:
            target.pop("featureServiceXSSFilter", None)
        self.progress.emit("Rollback: mengembalikan hanya state featureServiceXSSFilter...")
        rollback_result = self.update_system_properties(session, token, target)
        verified = self.read_system_properties(session, token)
        expected_present = bool(previous.get("property_was_present"))
        actual_present = "featureServiceXSSFilter" in verified
        if expected_present:
            verified_ok = (
                actual_present and
                str(verified.get("featureServiceXSSFilter")).strip().lower() ==
                str(previous.get("featureServiceXSSFilter")).strip().lower()
            )
        else:
            verified_ok = not actual_present
        _, verified_text = self.feature_service_xss_state(verified)
        if not verified_ok:
            raise ArcGISRequestError(
                "Rollback selesai, tetapi featureServiceXSSFilter tidak sesuai backup."
            )
        metadata_file = backup_file.parent / "metadata.json"
        if metadata_file.exists():
            self.update_metadata(
                metadata_file,
                rollback_response=rollback_result,
                rollback_verified_state=verified_text,
                status="rolled_back",
                rolled_back_at=datetime.now().astimezone().isoformat(),
            )
        return {
            "success": True,
            "verified_value": verified_text,
            "backup_file": str(backup_file),
            "server_token": self.server_token,
            "server_token_expires": self.server_token_expires,
        }

    def read_token_configuration(self, session, token):
        return self.request_json(
            session, "GET", f"{self.server_admin_url}/security/tokens",
            params={
                "token": token,
                "f": "json",
                "_ts": str(int(time.time() * 1000)),
            },
            headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
        )

    def update_token_configuration(self, session, token, configuration):
        return self.request_json(
            session, "POST", f"{self.server_admin_url}/security/tokens/update",
            data={
                "tokenManagerConfig": json.dumps(configuration, ensure_ascii=False),
                "token": token,
                "f": "json",
            },
        )

    def allow_http_get_state(self, configuration):
        properties = configuration.get("properties", {})
        if "allowHttpGet" not in properties:
            return False, "Not configured (default disabled)"
        raw = properties.get("allowHttpGet")
        value = str(raw).strip().lower()
        if value == "false":
            return True, "Disabled (explicit)"
        if value == "true":
            return False, "Enabled (explicit)"
        return None, f"Unknown ({raw})"

    def create_token_http_get_backup(self, configuration):
        # Shared key sengaja tidak disimpan ke disk. Rollback membaca konfigurasi
        # live terbaru lalu hanya merge nilai allowHttpGet sebelumnya.
        execution_id = datetime.now().strftime("%Y%m%d-%H%M%S")
        base_dir = Path.cwd() / "backups" / execution_id
        base_dir.mkdir(parents=True, exist_ok=False)
        backup_file = base_dir / "server-token-http-get-before.json"
        metadata_file = base_dir / "metadata.json"
        properties = configuration.get("properties", {})
        safe_backup = {
            "property_was_present": "allowHttpGet" in properties,
            "allowHttpGet": properties.get("allowHttpGet"),
        }
        backup_file.write_text(
            json.dumps(safe_backup, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        metadata = {
            "execution_id": execution_id,
            "created_at": datetime.now().astimezone().isoformat(),
            "component": "ArcGIS Server",
            "control": "disable_token_acquisition_via_http_get",
            "server_admin_url": self.server_admin_url,
            "before": safe_backup,
            "target": {"allowHttpGet": "false"},
            "sensitive_values_saved": False,
            "status": "backup_created",
        }
        metadata_file.write_text(
            json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        return backup_file, metadata_file, execution_id

    def test_token_post(self, session):
        result = self.request_json(
            session, "POST", f"{self.server_admin_url}/generateToken",
            data={
                "username": self.server_username,
                "password": self.server_password,
                "client": "requestip",
                "expiration": "5",
                "f": "json",
            },
        )
        return bool(result.get("token"))

    def test_token_get_rejected(self, session):
        # Request dikirim langsung agar respons penolakan dapat dievaluasi tanpa
        # menulis credential ke Execution Log ataupun exception detail.
        try:
            response = session.get(
                f"{self.server_admin_url}/generateToken",
                params={
                    "username": self.server_username,
                    "password": self.server_password,
                    "client": "requestip",
                    "expiration": "5",
                    "f": "json",
                },
                timeout=60,
            )
        except requests.exceptions.RequestException:
            return True
        try:
            result = response.json()
        except ValueError:
            return response.status_code >= 400
        # Jika GET menghasilkan token, hardening belum efektif.
        return not bool(result.get("token"))

    def apply_token_http_get(self, session, token):
        self.progress.emit("Apply: membaca seluruh Token Manager Configuration terbaru...")
        current = self.read_token_configuration(session, token)
        compliant, current_text = self.allow_http_get_state(current)
        if compliant is True:
            return {
                "success": True,
                "already_compliant": True,
                "verified_value": current_text,
                "server_token": self.server_token,
                "server_token_expires": self.server_token_expires,
            }
        if compliant is None:
            raise ArcGISRequestError(
                "Nilai allowHttpGet tidak dikenali. Apply dihentikan agar tidak mengubah "
                f"konfigurasi ambigu: {current.get('properties', {}).get('allowHttpGet')}"
            )
        self.progress.emit("Apply: membuat backup aman tanpa sharedKey...")
        backup_file, metadata_file, execution_id = self.create_token_http_get_backup(current)
        target = json.loads(json.dumps(current))
        target.setdefault("properties", {})["allowHttpGet"] = "false"
        self.progress.emit("Apply: merge allowHttpGet=false tanpa mengubah property token lain...")
        apply_result = self.update_token_configuration(session, token, target)
        self.progress.emit("Verify: membaca ulang Token Manager Configuration live...")
        verified = self.read_token_configuration(session, token)
        verified_ok, verified_text = self.allow_http_get_state(verified)
        # Jangan melakukan negative test GET dengan credential asli: tindakan itu
        # justru menaruh password pada URL/log, yaitu risiko yang sedang dicegah.
        self.progress.emit("Verify: positive test HTTP POST dengan credential di request body...")
        post_works = self.test_token_post(session)
        all_ok = verified_ok is True and post_works
        self.update_metadata(
            metadata_file,
            apply_response=apply_result,
            verified_state=verified_text,
            negative_get_test_skipped=True,
            negative_get_test_reason="Avoid placing credentials in URL/logs",
            post_request_succeeded=post_works,
            status="verified_success" if all_ok else "verification_failed",
            completed_at=datetime.now().astimezone().isoformat(),
        )
        if not all_ok:
            raise ArcGISRequestError(
                "Live verification Token HTTP GET gagal. "
                f"Config={verified_text}; POST works={post_works}"
            )
        return {
            "success": True,
            "already_compliant": False,
            "execution_id": execution_id,
            "backup_file": str(backup_file),
            "metadata_file": str(metadata_file),
            "verified_value": verified_text,
            "server_token": self.server_token,
            "server_token_expires": self.server_token_expires,
        }

    def rollback_token_http_get(self, session, token):
        if not self.backup_file:
            raise ArcGISRequestError("Backup Token HTTP GET tidak tersedia.")
        backup_file = Path(self.backup_file)
        if not backup_file.exists():
            raise ArcGISRequestError(f"Backup file tidak ditemukan: {backup_file}")
        previous = json.loads(backup_file.read_text(encoding="utf-8"))
        current = self.read_token_configuration(session, token)
        target = json.loads(json.dumps(current))
        properties = target.setdefault("properties", {})
        if previous.get("property_was_present"):
            properties["allowHttpGet"] = previous.get("allowHttpGet")
        else:
            properties.pop("allowHttpGet", None)
        self.progress.emit("Rollback: merge nilai allowHttpGet sebelumnya ke konfigurasi live...")
        rollback_result = self.update_token_configuration(session, token, target)
        verified = self.read_token_configuration(session, token)
        verified_properties = verified.get("properties", {})
        expected_present = bool(previous.get("property_was_present"))
        actual_present = "allowHttpGet" in verified_properties
        if expected_present:
            verified_ok = (
                actual_present
                and str(verified_properties.get("allowHttpGet")).lower()
                == str(previous.get("allowHttpGet")).lower()
            )
        else:
            verified_ok = not actual_present
        _, verified_text = self.allow_http_get_state(verified)
        if not verified_ok:
            raise ArcGISRequestError(
                "Rollback allowHttpGet selesai, tetapi live verification tidak sesuai backup."
            )
        metadata_file = backup_file.parent / "metadata.json"
        if metadata_file.exists():
            self.update_metadata(
                metadata_file,
                rollback_response=rollback_result,
                rollback_verified_state=verified_text,
                status="rolled_back",
                rolled_back_at=datetime.now().astimezone().isoformat(),
            )
        return {
            "success": True,
            "verified_value": verified_text,
            "backup_file": str(backup_file),
            "server_token": self.server_token,
            "server_token_expires": self.server_token_expires,
        }

    def build_edit_payload(
        self, current, token, callback_target=None, services_dir_target=None
    ):
        # Kirim ulang seluruh properti scalar agar setting lain tidak ter-reset.
        payload = {
            key: value
            for key, value in current.items()
            if isinstance(value, (str, int, float, bool))
            and key not in ("status", "success", "error")
        }
        if callback_target is not None:
            payload["callbackFunctionsEnabled"] = str(bool(callback_target)).lower()
        if services_dir_target is not None:
            payload["servicesDirEnabled"] = str(bool(services_dir_target)).lower()
        payload["token"] = token
        payload["f"] = "json"
        return payload

    def create_backup(self, current, control, property_name, target_value):
        execution_id = datetime.now().strftime("%Y%m%d-%H%M%S")
        base_dir = Path.cwd() / "backups" / execution_id
        base_dir.mkdir(parents=True, exist_ok=False)
        backup_file = base_dir / "server-services-directory-before.json"
        metadata_file = base_dir / "metadata.json"
        backup_file.write_text(
            json.dumps(current, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        metadata = {
            "execution_id": execution_id,
            "created_at": datetime.now().astimezone().isoformat(),
            "component": "ArcGIS Server",
            "control": control,
            "server_admin_url": self.server_admin_url,
            "before": {property_name: find_value(current, property_name)},
            "target": {property_name: target_value},
            "status": "backup_created",
        }
        metadata_file.write_text(
            json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        return backup_file, metadata_file, execution_id

    def update_metadata(self, metadata_file, **updates):
        metadata = json.loads(metadata_file.read_text(encoding="utf-8"))
        metadata.update(updates)
        metadata_file.write_text(
            json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8"
        )

    def apply_jsonp(self, session, token):
        self.progress.emit("Apply: membaca konfigurasi Services Directory terbaru...")
        current = self.read_services_directory(session, token)
        before = find_value(current, "callbackFunctionsEnabled")
        if str(before).lower() == "false":
            return {
                "success": True,
                "already_compliant": True,
                "verified_value": False,
                "server_token": self.server_token,
                "server_token_expires": self.server_token_expires,
            }

        self.progress.emit("Apply: membuat backup JSON...")
        backup_file, metadata_file, execution_id = self.create_backup(
            current, "disable_jsonp_callback", "callbackFunctionsEnabled", False
        )

        self.progress.emit("Apply: menonaktifkan JSONP callback functions...")
        payload = self.build_edit_payload(current, token, False)
        apply_result = self.request_json(
            session,
            "POST",
            f"{self.server_admin_url}/system/handlers/rest/servicesdirectory/edit",
            data=payload,
        )

        self.progress.emit("Verify: membaca ulang konfigurasi live...")
        verified = self.read_services_directory(session, token)
        verified_value = find_value(verified, "callbackFunctionsEnabled")
        verified_ok = str(verified_value).lower() == "false"
        self.update_metadata(
            metadata_file,
            apply_response=apply_result,
            verified_value=verified_value,
            status="verified_success" if verified_ok else "verification_failed",
            completed_at=datetime.now().astimezone().isoformat(),
        )
        if not verified_ok:
            raise ArcGISRequestError(
                "Apply mendapat respons, tetapi live verification gagal. "
                f"Nilai aktual callbackFunctionsEnabled: {verified_value}"
            )
        return {
            "success": True,
            "already_compliant": False,
            "execution_id": execution_id,
            "backup_file": str(backup_file),
            "metadata_file": str(metadata_file),
            "verified_value": verified_value,
            "server_token": self.server_token,
            "server_token_expires": self.server_token_expires,
        }

    def rollback_jsonp(self, session, token):
        if not self.backup_file:
            raise ArcGISRequestError("Backup file untuk rollback tidak tersedia.")
        backup_file = Path(self.backup_file)
        if not backup_file.exists():
            raise ArcGISRequestError(f"Backup file tidak ditemukan: {backup_file}")
        previous = json.loads(backup_file.read_text(encoding="utf-8"))
        target_value = find_value(previous, "callbackFunctionsEnabled")
        if str(target_value).lower() not in ("true", "false"):
            raise ArcGISRequestError(
                "Nilai callbackFunctionsEnabled pada backup tidak valid."
            )
        callback_target = str(target_value).lower() == "true"
        current = self.read_services_directory(session, token)
        payload = self.build_edit_payload(current, token, callback_target)
        self.progress.emit("Rollback: mengembalikan konfigurasi JSONP dari backup...")
        rollback_result = self.request_json(
            session,
            "POST",
            f"{self.server_admin_url}/system/handlers/rest/servicesdirectory/edit",
            data=payload,
        )
        self.progress.emit("Rollback: melakukan live verification...")
        verified = self.read_services_directory(session, token)
        verified_value = find_value(verified, "callbackFunctionsEnabled")
        verified_ok = str(verified_value).lower() == str(target_value).lower()
        if not verified_ok:
            raise ArcGISRequestError(
                "Rollback request selesai, tetapi live verification gagal. "
                f"Nilai aktual: {verified_value}; target backup: {target_value}"
            )
        metadata_file = backup_file.parent / "metadata.json"
        if metadata_file.exists():
            self.update_metadata(
                metadata_file,
                rollback_response=rollback_result,
                rollback_verified_value=verified_value,
                status="rolled_back",
                rolled_back_at=datetime.now().astimezone().isoformat(),
            )
        return {
            "success": True,
            "verified_value": verified_value,
            "backup_file": str(backup_file),
            "server_token": self.server_token,
            "server_token_expires": self.server_token_expires,
        }

    def apply_services_directory(self, session, token):
        self.progress.emit("Apply: membaca konfigurasi Services Directory terbaru...")
        current = self.read_services_directory(session, token)
        before = find_value(
            current, "servicesDirEnabled", "servicesDirectoryEnabled", "enabled"
        )
        if str(before).lower() == "false":
            return {
                "success": True,
                "already_compliant": True,
                "verified_value": False,
                "server_token": self.server_token,
                "server_token_expires": self.server_token_expires,
            }

        self.progress.emit("Apply: membuat backup JSON untuk Services Directory...")
        backup_file, metadata_file, execution_id = self.create_backup(
            current, "disable_services_directory", "servicesDirEnabled", False
        )
        self.progress.emit("Apply: menonaktifkan Services Directory...")
        payload = self.build_edit_payload(
            current, token, services_dir_target=False
        )
        apply_result = self.request_json(
            session,
            "POST",
            f"{self.server_admin_url}/system/handlers/rest/servicesdirectory/edit",
            data=payload,
        )
        self.progress.emit("Verify: membaca ulang Services Directory secara live...")
        verified = self.read_services_directory(session, token)
        verified_value = find_value(
            verified, "servicesDirEnabled", "servicesDirectoryEnabled", "enabled"
        )
        verified_ok = str(verified_value).lower() == "false"
        self.update_metadata(
            metadata_file,
            apply_response=apply_result,
            verified_value=verified_value,
            status="verified_success" if verified_ok else "verification_failed",
            completed_at=datetime.now().astimezone().isoformat(),
        )
        if not verified_ok:
            raise ArcGISRequestError(
                "Apply mendapat respons, tetapi live verification Services Directory gagal. "
                f"Nilai aktual: {verified_value}"
            )
        return {
            "success": True,
            "already_compliant": False,
            "execution_id": execution_id,
            "backup_file": str(backup_file),
            "metadata_file": str(metadata_file),
            "verified_value": verified_value,
            "server_token": self.server_token,
            "server_token_expires": self.server_token_expires,
        }

    def rollback_services_directory(self, session, token):
        if not self.backup_file:
            raise ArcGISRequestError("Backup file Services Directory tidak tersedia.")
        backup_file = Path(self.backup_file)
        if not backup_file.exists():
            raise ArcGISRequestError(f"Backup file tidak ditemukan: {backup_file}")
        previous = json.loads(backup_file.read_text(encoding="utf-8"))
        target_value = find_value(
            previous, "servicesDirEnabled", "servicesDirectoryEnabled", "enabled"
        )
        if str(target_value).lower() not in ("true", "false"):
            raise ArcGISRequestError("Nilai Services Directory pada backup tidak valid.")
        services_target = str(target_value).lower() == "true"
        current = self.read_services_directory(session, token)
        payload = self.build_edit_payload(
            current, token, services_dir_target=services_target
        )
        self.progress.emit("Rollback: mengembalikan Services Directory dari backup...")
        rollback_result = self.request_json(
            session,
            "POST",
            f"{self.server_admin_url}/system/handlers/rest/servicesdirectory/edit",
            data=payload,
        )
        self.progress.emit("Rollback: melakukan live verification Services Directory...")
        verified = self.read_services_directory(session, token)
        verified_value = find_value(
            verified, "servicesDirEnabled", "servicesDirectoryEnabled", "enabled"
        )
        verified_ok = str(verified_value).lower() == str(target_value).lower()
        if not verified_ok:
            raise ArcGISRequestError(
                "Rollback Services Directory selesai, tetapi live verification gagal. "
                f"Nilai aktual: {verified_value}; target backup: {target_value}"
            )
        metadata_file = backup_file.parent / "metadata.json"
        if metadata_file.exists():
            self.update_metadata(
                metadata_file,
                rollback_response=rollback_result,
                rollback_verified_value=verified_value,
                status="rolled_back",
                rolled_back_at=datetime.now().astimezone().isoformat(),
            )
        return {
            "success": True,
            "verified_value": verified_value,
            "backup_file": str(backup_file),
            "server_token": self.server_token,
            "server_token_expires": self.server_token_expires,
        }

    def run(self):
        session = requests.Session()
        session.verify = self.verify_ssl
        result = {"action": self.action, "success": False}
        try:
            token = self.ensure_token(session)
            if self.control == "jsonp" and self.action == "apply":
                result = self.apply_jsonp(session, token)
            elif self.control == "jsonp" and self.action == "rollback":
                result = self.rollback_jsonp(session, token)
            elif self.control == "services_directory" and self.action == "apply":
                result = self.apply_services_directory(session, token)
            elif self.control == "services_directory" and self.action == "rollback":
                result = self.rollback_services_directory(session, token)
            elif self.control == "standardized_queries" and self.action == "apply":
                result = self.apply_standardized_queries(session, token)
            elif self.control == "standardized_queries" and self.action == "rollback":
                result = self.rollback_standardized_queries(session, token)
            elif self.control == "feature_service_xss" and self.action == "apply":
                result = self.apply_feature_service_xss(session, token, "input")
            elif self.control == "feature_service_xss" and self.action == "apply_advanced":
                result = self.apply_feature_service_xss(session, token, "inputOutput")
            elif self.control == "feature_service_xss" and self.action == "rollback":
                result = self.rollback_feature_service_xss(session, token)
            elif self.control == "token_http_get" and self.action == "apply":
                result = self.apply_token_http_get(session, token)
            elif self.control == "token_http_get" and self.action == "rollback":
                result = self.rollback_token_http_get(session, token)
            else:
                raise ArcGISRequestError(
                    f"Action/control tidak dikenal: {self.action}/{self.control}"
                )
            result["action"] = self.action
            result["control"] = self.control
        except Exception as error:
            result = {
                "action": self.action,
                "control": self.control,
                "success": False,
                "error": str(error),
                "server_token": self.server_token,
                "server_token_expires": self.server_token_expires,
            }
        finally:
            session.close()
            self.completed.emit(result)
            self.finished.emit()


class PortalDirectoryHardeningWorker(QObject):
    progress = Signal(str)
    completed = Signal(dict)
    finished = Signal()

    def __init__(
        self, action, control, portal_admin_url, portal_username, portal_password,
        verify_ssl, portal_token=None, portal_token_expires=0,
        backup_file=None, selected_properties=None, portal_version=None,
    ):
        super().__init__()
        self.action = action
        self.control = control
        self.portal_admin_url = normalize_url(portal_admin_url)
        self.portal_base_url = derive_portal_base_url(self.portal_admin_url)
        self.portal_username = portal_username.strip()
        self.portal_password = portal_password
        self.verify_ssl = verify_ssl
        self.portal_token = portal_token
        self.portal_token_expires = float(portal_token_expires or 0)
        self.backup_file = backup_file
        self.selected_properties = tuple(selected_properties or ())
        self.portal_version = str(portal_version or "").strip()

    def request_json(self, session, method, url, **kwargs):
        try:
            response = session.request(method, url, timeout=60, **kwargs)
        except requests.exceptions.RequestException as error:
            raise ArcGISRequestError(
                f"Portal request gagal.\nEndpoint: {url}\nDetail: {error}"
            ) from error
        if response.status_code >= 400:
            message = "Portal endpoint menolak request."
            if response.status_code == 403:
                message = "Akses Portal Administrator API ditolak (HTTP 403)."
            elif response.status_code in (498, 499):
                message = f"Portal token ditolak (HTTP {response.status_code})."
            raise ArcGISRequestError(
                f"{message}\nHTTP {response.status_code}\nEndpoint: {url}"
            )
        try:
            result = response.json()
        except ValueError as error:
            raise ArcGISRequestError(
                f"Portal endpoint tidak mengembalikan JSON.\nEndpoint: {url}"
            ) from error
        if isinstance(result, dict) and result.get("error"):
            err = result["error"]
            code = err.get("code", "-")
            if code in (498, 499):
                summary = "Portal token ditolak oleh endpoint."
            elif code == 403:
                summary = "Akun/token tidak memiliki akses ke endpoint Portal Administrator API."
            else:
                summary = "Portal API menolak request."
            raise ArcGISRequestError(
                f"{summary}\nCode: {code}\n"
                f"Message: {err.get('message', '-')}\nDetails: {err.get('details', '-')}"
            )
        return result

    def token_valid(self):
        return bool(self.portal_token) and self.portal_token_expires > time.time() + 120

    def ensure_token(self, session):
        if self.token_valid():
            self.progress.emit("Portal Apply: menggunakan token yang masih valid dari memory...")
            return self.portal_token
        self.progress.emit("Portal Apply: membuat Portal token baru...")
        result = self.request_json(
            session, "POST", f"{self.portal_base_url}/sharing/rest/generateToken",
            data={
                "username": self.portal_username,
                "password": self.portal_password,
                "client": "referer",
                "referer": f"{self.portal_base_url}/home/",
                "expiration": "30",
                "f": "json",
            },
        )
        token = result.get("token")
        if not token:
            raise ArcGISRequestError("Portal tidak mengembalikan administrator token.")
        expires_ms = result.get("expires", 0)
        self.portal_token = token
        self.portal_token_expires = (
            float(expires_ms) / 1000 if expires_ms else time.time() + 25 * 60
        )
        return token

    def read_portal_system_properties(self, session, token):
        return self.request_json(
            session, "GET", f"{self.portal_admin_url}/system/properties",
            params={
                "token": token,
                "f": "json",
                "_ts": str(int(time.time() * 1000)),
            },
            headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
        )

    def update_portal_system_properties(self, session, token, properties):
        # HAR Portal 11.5 menunjukkan UI mengirim seluruh object properties.
        # Gunakan GET + merge + POST agar property lain tidak hilang.
        return self.request_json(
            session, "POST", f"{self.portal_admin_url}/system/properties/update",
            data={
                "properties": json.dumps(properties, ensure_ascii=False),
                "token": token,
                "f": "json",
            },
        )

    @staticmethod
    def portal_servlet_property_names():
        return (
            "disableLegendServlet",
            "disablePrintServlet",
            "disableWFSServlet",
        )
    def portal_servlet_compatibility(self):
        version = self.portal_version
        if version.startswith("11.5"):
            return "SUPPORTED"
        if version.startswith("12.1"):
            return "NOT_SUPPORTED"
        return "UNVALIDATED"
    def validate_portal_servlet_operation(self):
        compatibility = self.portal_servlet_compatibility()
        if compatibility != "SUPPORTED":
            raise ArcGISRequestError(
                "Portal Servlet Hardening diblokir secara fail-safe. "
                f"Portal version={self.portal_version or 'unknown'}; "
                f"compatibility={compatibility}. Control hanya divalidasi pada 11.5; "
                "ArcGIS Enterprise 12.1 menolak property servlet."
            )
        allowed = set(self.portal_servlet_property_names())
        selected = tuple(dict.fromkeys(self.selected_properties))
        if not selected or not set(selected).issubset(allowed):
            raise ArcGISRequestError("Pilihan Portal Servlet kosong atau tidak valid.")
        return selected
    @staticmethod
    def portal_servlet_value(properties, name):
        if name not in properties:
            return None, "Not configured"
        raw = properties.get(name)
        value = str(raw).strip().lower()
        if value == "true":
            return True, "Disabled (explicit)"
        if value == "false":
            return False, "Enabled (explicit)"
        return None, f"Unknown ({raw})"
    def create_portal_servlet_backup(self, current, selected):
        execution_id = datetime.now().strftime("%Y%m%d-%H%M%S")
        base_dir = Path.cwd() / "backups" / execution_id
        base_dir.mkdir(parents=True, exist_ok=False)
        backup_file = base_dir / "portal-servlets-before.json"
        metadata_file = base_dir / "metadata.json"
        safe_backup = {
            "portal_version": self.portal_version,
            "selected_properties": {
                name: {
                    "property_was_present": name in current,
                    "value": current.get(name),
                }
                for name in selected
            },
        }
        backup_file.write_text(
            json.dumps(safe_backup, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        metadata = {
            "execution_id": execution_id,
            "created_at": datetime.now().astimezone().isoformat(),
            "component": "Portal for ArcGIS",
            "control": "disable_portal_servlets",
            "portal_admin_url": self.portal_admin_url,
            "portal_version": self.portal_version,
            "selected_properties": list(selected),
            "before": safe_backup["selected_properties"],
            "target": {name: True for name in selected},
            "scope": "selected_portal_servlet_properties_only",
            "status": "backup_created",
        }
        metadata_file.write_text(
            json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        return backup_file, metadata_file, execution_id
    def portal_servlet_summary(self, properties, selected=None):
        names = selected or self.portal_servlet_property_names()
        labels = {
            "disableLegendServlet": "Legend",
            "disablePrintServlet": "Print",
            "disableWFSServlet": "WFS",
        }
        parts = []
        for name in names:
            _ok, text = self.portal_servlet_value(properties, name)
            parts.append(f"{labels[name]}: {text}")
        return "; ".join(parts)
    def apply_portal_servlets(self, session, token):
        selected = self.validate_portal_servlet_operation()
        self.progress.emit("Portal Servlet Apply: GET latest Portal System Properties...")
        current = self.read_portal_system_properties(session, token)
        for name in selected:
            state, _text = self.portal_servlet_value(current, name)
            if state is None and name in current:
                raise ArcGISRequestError(
                    f"Nilai {name} tidak dikenali. Apply dihentikan agar tidak menebak konfigurasi."
                )
        pending = tuple(
            name for name in selected
            if self.portal_servlet_value(current, name)[0] is not True
        )
        if not pending:
            return {
                "success": True,
                "already_compliant": True,
                "verified_value": self.portal_servlet_summary(current, selected),
                "verified_properties": current,
                "selected_properties": list(selected),
                "portal_token": self.portal_token,
                "portal_token_expires": self.portal_token_expires,
            }
        self.progress.emit("Portal Servlet Apply: membuat scoped backup pilihan operator...")
        backup_file, metadata_file, execution_id = self.create_portal_servlet_backup(
            current, pending
        )
        target = dict(current)
        for name in pending:
            target[name] = True
        self.progress.emit(
            "Portal Servlet Apply: merge hanya property terpilih; property lain dipertahankan..."
        )
        apply_result = self.update_portal_system_properties(session, token, target)
        token = self.wait_for_portal_and_reauthenticate(
            session, apply_result.get("recheckAfterSeconds", 20)
        )
        self.progress.emit("Portal Servlet Verify: membaca System Properties live...")
        verified = self.read_portal_system_properties(session, token)
        failed = [
            name for name in pending
            if self.portal_servlet_value(verified, name)[0] is not True
        ]
        self.update_metadata(
            metadata_file,
            apply_response=apply_result,
            verified={name: verified.get(name) for name in pending},
            status="verified_success" if not failed else "verification_failed",
            completed_at=datetime.now().astimezone().isoformat(),
        )
        if failed:
            raise ArcGISRequestError(
                "Live verification Portal Servlet gagal untuk: " + ", ".join(failed)
            )
        return {
            "success": True,
            "already_compliant": False,
            "execution_id": execution_id,
            "backup_file": str(backup_file),
            "metadata_file": str(metadata_file),
            "verified_value": self.portal_servlet_summary(verified, selected),
            "verified_properties": verified,
            "selected_properties": list(selected),
            "portal_token": self.portal_token,
            "portal_token_expires": self.portal_token_expires,
        }
    def rollback_portal_servlets(self, session, token):
        self.validate_portal_servlet_operation()
        if not self.backup_file:
            raise ArcGISRequestError("Backup Portal Servlet tidak tersedia.")
        backup_file = Path(self.backup_file)
        if not backup_file.exists():
            raise ArcGISRequestError(f"Backup file tidak ditemukan: {backup_file}")
        previous = json.loads(backup_file.read_text(encoding="utf-8"))
        entries = previous.get("selected_properties") if isinstance(previous, dict) else None
        if not isinstance(entries, dict) or not entries:
            raise ArcGISRequestError("Backup Portal Servlet tidak valid.")
        allowed = set(self.portal_servlet_property_names())
        if not set(entries).issubset(allowed):
            raise ArcGISRequestError("Backup memuat property Portal Servlet yang tidak dikenal.")
        self.progress.emit(
            "Portal Servlet Rollback: GET latest agar property lain tetap dipertahankan..."
        )
        current = self.read_portal_system_properties(session, token)
        target = dict(current)
        for name, entry in entries.items():
            if not isinstance(entry, dict) or "property_was_present" not in entry:
                raise ArcGISRequestError(f"Backup state {name} tidak valid.")
            if entry.get("property_was_present"):
                target[name] = entry.get("value")
            else:
                target.pop(name, None)
        rollback_result = self.update_portal_system_properties(session, token, target)
        token = self.wait_for_portal_and_reauthenticate(
            session, rollback_result.get("recheckAfterSeconds", 20)
        )
        verified = self.read_portal_system_properties(session, token)
        failed = []
        for name, entry in entries.items():
            expected_present = bool(entry.get("property_was_present"))
            if expected_present:
                ok = name in verified and str(verified.get(name)).strip().lower() == str(
                    entry.get("value")
                ).strip().lower()
            else:
                ok = name not in verified
            if not ok:
                failed.append(name)
        if failed:
            raise ArcGISRequestError(
                "Live verification rollback Portal Servlet gagal untuk: " + ", ".join(failed)
            )
        metadata_file = backup_file.parent / "metadata.json"
        if metadata_file.exists():
            self.update_metadata(
                metadata_file,
                rollback_response=rollback_result,
                rollback_verified={name: verified.get(name) for name in entries},
                status="rolled_back",
                rolled_back_at=datetime.now().astimezone().isoformat(),
            )
        return {
            "success": True,
            "verified_value": self.portal_servlet_summary(verified, tuple(entries)),
            "verified_properties": verified,
            "selected_properties": list(entries),
            "backup_file": str(backup_file),
            "portal_token": self.portal_token,
            "portal_token_expires": self.portal_token_expires,
        }
    def builtin_self_creation_state(self, properties):
        if "disableSignup" not in properties:
            return False, "Not configured (default disabled)"
        raw = properties.get("disableSignup")
        value = str(raw).strip().lower()
        if value == "true":
            return True, "Disabled (explicit)"
        if value == "false":
            return False, "Enabled (explicit)"
        return None, f"Unknown ({raw})"

    def create_portal_system_properties_backup(self, current):
        execution_id = datetime.now().strftime("%Y%m%d-%H%M%S")
        base_dir = Path.cwd() / "backups" / execution_id
        base_dir.mkdir(parents=True, exist_ok=False)
        backup_file = base_dir / "portal-system-properties-before.json"
        metadata_file = base_dir / "metadata.json"
        backup_file.write_text(
            json.dumps(current, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        metadata = {
            "execution_id": execution_id,
            "created_at": datetime.now().astimezone().isoformat(),
            "component": "Portal for ArcGIS",
            "control": "disable_builtin_account_self_creation",
            "portal_admin_url": self.portal_admin_url,
            "property_was_present": "disableSignup" in current,
            "before": {"disableSignup": current.get("disableSignup")},
            "target": {"disableSignup": True},
            "status": "backup_created",
        }
        metadata_file.write_text(
            json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        return backup_file, metadata_file, execution_id

    def wait_for_portal_and_reauthenticate(self, session, requested_wait):
        wait_seconds = max(1, min(int(requested_wait or 20), 120))
        self.progress.emit(
            f"Portal restart/propagation: menunggu {wait_seconds} detik sesuai respons API..."
        )
        time.sleep(wait_seconds)
        deadline = time.time() + 300
        last_error = None
        while time.time() < deadline:
            try:
                self.progress.emit("Portal health check: mencoba autentikasi ulang...")
                self.portal_token = None
                self.portal_token_expires = 0
                token = self.ensure_token(session)
                self.read_portal_system_properties(session, token)
                self.progress.emit("Portal health check: Portal tersedia kembali.")
                return token
            except Exception as error:
                last_error = error
                self.progress.emit("Portal belum siap; mencoba kembali dalam 10 detik...")
                time.sleep(10)
        raise ArcGISRequestError(
            "Portal belum kembali tersedia dalam 5 menit setelah update. "
            f"Error terakhir: {last_error}"
        )

    def apply_builtin_self_creation(self, session, token):
        self.progress.emit("Portal Apply: membaca seluruh Portal System Properties terbaru...")
        current = self.read_portal_system_properties(session, token)
        compliant, current_text = self.builtin_self_creation_state(current)
        if compliant is True:
            return {
                "success": True,
                "already_compliant": True,
                "verified_value": current_text,
                "portal_token": self.portal_token,
                "portal_token_expires": self.portal_token_expires,
            }
        if compliant is None:
            raise ArcGISRequestError(
                "Nilai disableSignup tidak dikenali. Apply dihentikan agar tidak mengubah "
                f"konfigurasi ambigu: {current.get('disableSignup')}"
            )
        self.progress.emit("Portal Apply: membuat backup lengkap Portal System Properties...")
        backup_file, metadata_file, execution_id = (
            self.create_portal_system_properties_backup(current)
        )
        target = dict(current)
        target["disableSignup"] = True
        self.progress.emit("Portal Apply: merge disableSignup=true tanpa menghapus property lain...")
        apply_result = self.update_portal_system_properties(session, token, target)
        token = self.wait_for_portal_and_reauthenticate(
            session, apply_result.get("recheckAfterSeconds", 20)
        )
        self.progress.emit("Portal Verify: membaca ulang System Properties live...")
        verified = self.read_portal_system_properties(session, token)
        verified_ok, verified_text = self.builtin_self_creation_state(verified)
        self.update_metadata(
            metadata_file,
            apply_response=apply_result,
            verified_value=verified.get("disableSignup"),
            verified_state=verified_text,
            status="verified_success" if verified_ok is True else "verification_failed",
            completed_at=datetime.now().astimezone().isoformat(),
        )
        if verified_ok is not True:
            raise ArcGISRequestError(
                "Apply mendapat respons, tetapi live verification disableSignup gagal. "
                f"Kondisi aktual: {verified_text}"
            )
        return {
            "success": True,
            "already_compliant": False,
            "execution_id": execution_id,
            "backup_file": str(backup_file),
            "metadata_file": str(metadata_file),
            "verified_value": verified_text,
            "portal_token": self.portal_token,
            "portal_token_expires": self.portal_token_expires,
        }

    def rollback_builtin_self_creation(self, session, token):
        if not self.backup_file:
            raise ArcGISRequestError("Backup Built-In Account Self-Creation tidak tersedia.")
        backup_file = Path(self.backup_file)
        if not backup_file.exists():
            raise ArcGISRequestError(f"Backup file tidak ditemukan: {backup_file}")
        previous = json.loads(backup_file.read_text(encoding="utf-8"))
        if not isinstance(previous, dict):
            raise ArcGISRequestError("Backup Portal System Properties bukan JSON object.")
        self.progress.emit("Portal Rollback: mengembalikan exact System Properties dari backup...")
        rollback_result = self.update_portal_system_properties(session, token, previous)
        token = self.wait_for_portal_and_reauthenticate(
            session, rollback_result.get("recheckAfterSeconds", 20)
        )
        self.progress.emit("Portal Rollback: melakukan live verification exact backup...")
        verified = self.read_portal_system_properties(session, token)
        verified_ok = verified == previous
        _, verified_text = self.builtin_self_creation_state(verified)
        if not verified_ok:
            raise ArcGISRequestError(
                "Rollback selesai, tetapi Portal System Properties tidak identik dengan backup."
            )
        metadata_file = backup_file.parent / "metadata.json"
        if metadata_file.exists():
            self.update_metadata(
                metadata_file,
                rollback_response=rollback_result,
                rollback_verified_state=verified_text,
                status="rolled_back",
                rolled_back_at=datetime.now().astimezone().isoformat(),
            )
        return {
            "success": True,
            "verified_value": verified_text,
            "backup_file": str(backup_file),
            "portal_token": self.portal_token,
            "portal_token_expires": self.portal_token_expires,
        }

    def read_portal_self(self, session, token):
        return self.request_json(
            session, "GET", f"{self.portal_base_url}/sharing/rest/portals/self",
            params={
                "token": token,
                "f": "json",
                "_ts": str(int(time.time() * 1000)),
            },
            headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
        )

    def update_portal_self_property(self, session, token, property_name, value):
        return self.request_json(
            session, "POST", f"{self.portal_base_url}/sharing/rest/portals/self/update",
            data={
                property_name: str(bool(value)).lower(),
                "token": token,
                "f": "json",
            },
        )

    def public_profile_state(self, portal_self):
        if "updateUserProfileDisabled" not in portal_self:
            return None, "Unknown (property unavailable)"
        raw = portal_self.get("updateUserProfileDisabled")
        value = str(raw).strip().lower()
        if value == "true":
            return True, "Disabled (explicit)"
        if value == "false":
            return False, "Enabled (explicit)"
        return None, f"Unknown ({raw})"

    def create_public_profile_backup(self, portal_self):
        execution_id = datetime.now().strftime("%Y%m%d-%H%M%S")
        base_dir = Path.cwd() / "backups" / execution_id
        base_dir.mkdir(parents=True, exist_ok=False)
        backup_file = base_dir / "portal-public-profile-before.json"
        metadata_file = base_dir / "metadata.json"
        safe_backup = {
            "property_was_present": "updateUserProfileDisabled" in portal_self,
            "updateUserProfileDisabled": portal_self.get("updateUserProfileDisabled"),
        }
        backup_file.write_text(
            json.dumps(safe_backup, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        metadata = {
            "execution_id": execution_id,
            "created_at": datetime.now().astimezone().isoformat(),
            "component": "Portal for ArcGIS",
            "control": "disable_public_user_profile_sharing",
            "portal_admin_url": self.portal_admin_url,
            "before": safe_backup,
            "target": {"updateUserProfileDisabled": True},
            "sensitive_values_saved": False,
            "status": "backup_created",
        }
        metadata_file.write_text(
            json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        return backup_file, metadata_file, execution_id

    def apply_public_profile_sharing(self, session, token):
        self.progress.emit("Portal Apply: membaca Portal Self terbaru...")
        current = self.read_portal_self(session, token)
        compliant, current_text = self.public_profile_state(current)
        if compliant is True:
            return {
                "success": True,
                "already_compliant": True,
                "verified_value": current_text,
                "portal_token": self.portal_token,
                "portal_token_expires": self.portal_token_expires,
            }
        if compliant is None:
            raise ArcGISRequestError(
                "Property updateUserProfileDisabled tidak tersedia atau nilainya tidak dikenali. "
                "Apply dihentikan agar tidak menebak konfigurasi Portal."
            )
        self.progress.emit("Portal Apply: membuat backup state profile sharing...")
        backup_file, metadata_file, execution_id = self.create_public_profile_backup(current)
        self.progress.emit("Portal Apply: mengirim partial update updateUserProfileDisabled=true...")
        apply_result = self.update_portal_self_property(
            session, token, "updateUserProfileDisabled", True
        )
        if not apply_result.get("success"):
            raise ArcGISRequestError("Portal tidak mengonfirmasi keberhasilan update profil.")
        self.progress.emit("Portal Verify: membaca ulang Portal Self live...")
        verified = self.read_portal_self(session, token)
        verified_ok, verified_text = self.public_profile_state(verified)
        self.update_metadata(
            metadata_file,
            apply_response=apply_result,
            verified_value=verified.get("updateUserProfileDisabled"),
            verified_state=verified_text,
            status="verified_success" if verified_ok is True else "verification_failed",
            completed_at=datetime.now().astimezone().isoformat(),
        )
        if verified_ok is not True:
            raise ArcGISRequestError(
                "Apply mendapat respons sukses, tetapi live verification "
                f"updateUserProfileDisabled gagal: {verified_text}"
            )
        return {
            "success": True,
            "already_compliant": False,
            "execution_id": execution_id,
            "backup_file": str(backup_file),
            "metadata_file": str(metadata_file),
            "verified_value": verified_text,
            "portal_token": self.portal_token,
            "portal_token_expires": self.portal_token_expires,
        }

    def rollback_public_profile_sharing(self, session, token):
        if not self.backup_file:
            raise ArcGISRequestError("Backup Public User Profile Sharing tidak tersedia.")
        backup_file = Path(self.backup_file)
        if not backup_file.exists():
            raise ArcGISRequestError(f"Backup file tidak ditemukan: {backup_file}")
        previous = json.loads(backup_file.read_text(encoding="utf-8"))
        if not previous.get("property_was_present"):
            raise ArcGISRequestError(
                "Backup tidak memiliki nilai updateUserProfileDisabled. Rollback dihentikan "
                "agar tidak menebak default Portal."
            )
        previous_raw = previous.get("updateUserProfileDisabled")
        if str(previous_raw).strip().lower() not in ("true", "false"):
            raise ArcGISRequestError("Nilai profile sharing pada backup tidak valid.")
        previous_value = str(previous_raw).strip().lower() == "true"
        self.progress.emit("Portal Rollback: mengembalikan state profile sharing dari backup...")
        rollback_result = self.update_portal_self_property(
            session, token, "updateUserProfileDisabled", previous_value
        )
        if not rollback_result.get("success"):
            raise ArcGISRequestError("Portal tidak mengonfirmasi keberhasilan rollback profil.")
        verified = self.read_portal_self(session, token)
        actual = verified.get("updateUserProfileDisabled")
        verified_ok = str(actual).strip().lower() == str(previous_value).lower()
        _, verified_text = self.public_profile_state(verified)
        if not verified_ok:
            raise ArcGISRequestError(
                "Rollback selesai, tetapi live verification profile sharing tidak sesuai backup."
            )
        metadata_file = backup_file.parent / "metadata.json"
        if metadata_file.exists():
            self.update_metadata(
                metadata_file,
                rollback_response=rollback_result,
                rollback_verified_state=verified_text,
                status="rolled_back",
                rolled_back_at=datetime.now().astimezone().isoformat(),
            )
        return {
            "success": True,
            "verified_value": verified_text,
            "backup_file": str(backup_file),
            "portal_token": self.portal_token,
            "portal_token_expires": self.portal_token_expires,
        }

    def social_media_links_state(self, portal_self):
        portal_properties = portal_self.get("portalProperties")
        if not isinstance(portal_properties, dict):
            return None, "Unknown (portalProperties unavailable)"
        if "showSocialMediaLinks" not in portal_properties:
            return False, "Not configured (default disabled)"
        raw = portal_properties.get("showSocialMediaLinks")
        value = str(raw).strip().lower()
        if value == "false":
            return True, "Disabled (explicit)"
        if value == "true":
            return False, "Enabled (explicit)"
        return None, f"Unknown ({raw})"

    def create_social_media_links_backup(self, portal_properties):
        execution_id = datetime.now().strftime("%Y%m%d-%H%M%S")
        base_dir = Path.cwd() / "backups" / execution_id
        base_dir.mkdir(parents=True, exist_ok=False)
        backup_file = base_dir / "portal-properties-before.json"
        metadata_file = base_dir / "metadata.json"
        backup_file.write_text(
            json.dumps(portal_properties, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        metadata = {
            "execution_id": execution_id,
            "created_at": datetime.now().astimezone().isoformat(),
            "component": "Portal for ArcGIS",
            "control": "disable_show_social_media_links",
            "portal_admin_url": self.portal_admin_url,
            "property_was_present": "showSocialMediaLinks" in portal_properties,
            "before": {
                "showSocialMediaLinks": portal_properties.get("showSocialMediaLinks")
            },
            "target": {"showSocialMediaLinks": False},
            "status": "backup_created",
        }
        metadata_file.write_text(
            json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        return backup_file, metadata_file, execution_id

    def update_portal_properties(self, session, token, portal_properties):
        # HAR Portal 11.5 menunjukkan UI mengirim seluruh portalProperties.
        return self.request_json(
            session, "POST", f"{self.portal_base_url}/sharing/rest/portals/self/update",
            data={
                "portalProperties": json.dumps(portal_properties, ensure_ascii=False),
                "token": token,
                "f": "json",
            },
        )

    def apply_social_media_links(self, session, token):
        self.progress.emit("Portal Apply: membaca seluruh portalProperties terbaru...")
        current_self = self.read_portal_self(session, token)
        compliant, current_text = self.social_media_links_state(current_self)
        if compliant is True:
            return {
                "success": True,
                "already_compliant": True,
                "verified_value": current_text,
                "portal_token": self.portal_token,
                "portal_token_expires": self.portal_token_expires,
            }
        if compliant is None:
            raise ArcGISRequestError(
                "portalProperties tidak tersedia atau showSocialMediaLinks tidak dikenali. "
                "Apply dihentikan agar tidak menebak konfigurasi Portal."
            )
        current_properties = current_self.get("portalProperties", {})
        self.progress.emit("Portal Apply: membuat backup lengkap portalProperties...")
        backup_file, metadata_file, execution_id = (
            self.create_social_media_links_backup(current_properties)
        )
        target = json.loads(json.dumps(current_properties))
        target["showSocialMediaLinks"] = False
        self.progress.emit(
            "Portal Apply: merge showSocialMediaLinks=false tanpa menghapus property lain..."
        )
        apply_result = self.update_portal_properties(session, token, target)
        if not apply_result.get("success"):
            raise ArcGISRequestError("Portal tidak mengonfirmasi update social media links.")
        self.progress.emit("Portal Verify: membaca ulang portalProperties live...")
        verified_self = self.read_portal_self(session, token)
        verified_ok, verified_text = self.social_media_links_state(verified_self)
        self.update_metadata(
            metadata_file,
            apply_response=apply_result,
            verified_value=(
                verified_self.get("portalProperties", {}).get("showSocialMediaLinks")
                if isinstance(verified_self.get("portalProperties"), dict) else None
            ),
            verified_state=verified_text,
            status="verified_success" if verified_ok is True else "verification_failed",
            completed_at=datetime.now().astimezone().isoformat(),
        )
        if verified_ok is not True:
            raise ArcGISRequestError(
                "Apply mendapat respons sukses, tetapi live verification "
                f"showSocialMediaLinks gagal: {verified_text}"
            )
        return {
            "success": True,
            "already_compliant": False,
            "execution_id": execution_id,
            "backup_file": str(backup_file),
            "metadata_file": str(metadata_file),
            "verified_value": verified_text,
            "portal_token": self.portal_token,
            "portal_token_expires": self.portal_token_expires,
        }

    def rollback_social_media_links(self, session, token):
        if not self.backup_file:
            raise ArcGISRequestError("Backup Social Media Links tidak tersedia.")
        backup_file = Path(self.backup_file)
        if not backup_file.exists():
            raise ArcGISRequestError(f"Backup file tidak ditemukan: {backup_file}")
        previous = json.loads(backup_file.read_text(encoding="utf-8"))
        if not isinstance(previous, dict):
            raise ArcGISRequestError("Backup portalProperties bukan JSON object.")
        self.progress.emit("Portal Rollback: mengembalikan exact portalProperties dari backup...")
        rollback_result = self.update_portal_properties(session, token, previous)
        if not rollback_result.get("success"):
            raise ArcGISRequestError("Portal tidak mengonfirmasi rollback social media links.")
        verified_self = self.read_portal_self(session, token)
        verified_properties = verified_self.get("portalProperties")
        verified_ok = isinstance(verified_properties, dict) and verified_properties == previous
        _, verified_text = self.social_media_links_state(verified_self)
        if not verified_ok:
            raise ArcGISRequestError(
                "Rollback selesai, tetapi portalProperties tidak identik dengan backup."
            )
        metadata_file = backup_file.parent / "metadata.json"
        if metadata_file.exists():
            self.update_metadata(
                metadata_file,
                rollback_response=rollback_result,
                rollback_verified_state=verified_text,
                status="rolled_back",
                rolled_back_at=datetime.now().astimezone().isoformat(),
            )
        return {
            "success": True,
            "verified_value": verified_text,
            "backup_file": str(backup_file),
            "portal_token": self.portal_token,
            "portal_token_expires": self.portal_token_expires,
        }

    def anonymous_access_state(self, portal_self):
        raw = portal_self.get("access")
        value = str(raw).strip().lower()
        if value == "private":
            return True, "Disabled (private)"
        if value == "public":
            return False, "Enabled (public)"
        return None, f"Unknown ({raw})"

    def create_anonymous_access_backup(self, portal_self):
        execution_id = datetime.now().strftime("%Y%m%d-%H%M%S")
        base_dir = Path.cwd() / "backups" / execution_id
        base_dir.mkdir(parents=True, exist_ok=False)
        backup_file = base_dir / "portal-anonymous-access-before.json"
        metadata_file = base_dir / "metadata.json"
        safe_backup = {
            "access": portal_self.get("access"),
            "canShareBingPublic_present": "canShareBingPublic" in portal_self,
            "canShareBingPublic": portal_self.get("canShareBingPublic"),
        }
        backup_file.write_text(
            json.dumps(safe_backup, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        metadata = {
            "execution_id": execution_id,
            "created_at": datetime.now().astimezone().isoformat(),
            "component": "Portal for ArcGIS",
            "control": "disable_anonymous_access",
            "portal_admin_url": self.portal_admin_url,
            "before": safe_backup,
            "target": {"access": "private", "canShareBingPublic": False},
            "sensitive_values_saved": False,
            "status": "backup_created",
        }
        metadata_file.write_text(
            json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        return backup_file, metadata_file, execution_id

    def update_anonymous_access(self, session, token, access_value, include_bing=False):
        payload = {
            "access": access_value,
            "token": token,
            "f": "json",
        }
        if include_bing:
            payload["canShareBingPublic"] = "false"
        return self.request_json(
            session, "POST", f"{self.portal_base_url}/sharing/rest/portals/self/update",
            data=payload,
        )

    def anonymous_access_probe(self, session):
        # Evidence tambahan tanpa token. Konfigurasi authenticated tetap menjadi
        # sumber keputusan compliance karena Portal Self anonymous dapat mengembalikan
        # respons terbatas yang berbeda antar deployment/proxy.
        try:
            response = session.get(
                f"{self.portal_base_url}/sharing/rest/portals/self",
                params={"f": "json", "_ts": str(int(time.time() * 1000))},
                headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
                timeout=60,
            )
            try:
                body = response.json()
            except ValueError:
                return {
                    "http_status": response.status_code,
                    "result": "NON_JSON_OR_LOGIN_RESPONSE",
                }
            if isinstance(body, dict) and body.get("error"):
                return {
                    "http_status": response.status_code,
                    "result": "RESTRICTED",
                    "error_code": body.get("error", {}).get("code"),
                }
            return {
                "http_status": response.status_code,
                "result": "LIMITED_RESPONSE_REVIEW",
                "anonymous_access_value": body.get("access") if isinstance(body, dict) else None,
            }
        except requests.exceptions.RequestException as error:
            return {"result": "REQUEST_FAILED", "detail": type(error).__name__}

    def apply_anonymous_access(self, session, token):
        self.progress.emit("Portal Apply: membaca Portal Self terbaru...")
        current = self.read_portal_self(session, token)
        compliant, current_text = self.anonymous_access_state(current)
        if compliant is True:
            return {
                "success": True,
                "already_compliant": True,
                "verified_value": current_text,
                "server_evidence": "Authenticated access already private",
                "portal_token": self.portal_token,
                "portal_token_expires": self.portal_token_expires,
            }
        if compliant is None:
            raise ArcGISRequestError(
                "Nilai Portal access tidak dikenali. Apply dihentikan agar tidak menebak konfigurasi."
            )
        self.progress.emit("Portal Apply: membuat backup state anonymous access...")
        backup_file, metadata_file, execution_id = self.create_anonymous_access_backup(current)
        self.progress.emit("Portal Apply: mengirim access=private dan canShareBingPublic=false...")
        apply_result = self.update_anonymous_access(
            session, token, "private", include_bing=True
        )
        if not apply_result.get("success"):
            raise ArcGISRequestError("Portal tidak mengonfirmasi update anonymous access.")
        self.progress.emit("Portal Verify: membaca ulang Portal Self dengan token...")
        verified = self.read_portal_self(session, token)
        verified_ok, verified_text = self.anonymous_access_state(verified)
        self.progress.emit("Portal Verify: menjalankan anonymous evidence probe tanpa token...")
        anonymous_probe = self.anonymous_access_probe(session)
        self.update_metadata(
            metadata_file,
            apply_response=apply_result,
            verified_value=verified.get("access"),
            verified_state=verified_text,
            anonymous_probe=anonymous_probe,
            status="verified_success" if verified_ok is True else "verification_failed",
            completed_at=datetime.now().astimezone().isoformat(),
        )
        if verified_ok is not True:
            raise ArcGISRequestError(
                "Apply mendapat respons sukses, tetapi live verification access=private gagal. "
                f"Kondisi aktual: {verified_text}"
            )
        return {
            "success": True,
            "already_compliant": False,
            "execution_id": execution_id,
            "backup_file": str(backup_file),
            "metadata_file": str(metadata_file),
            "verified_value": verified_text,
            "anonymous_probe": anonymous_probe,
            "portal_token": self.portal_token,
            "portal_token_expires": self.portal_token_expires,
        }

    def rollback_anonymous_access(self, session, token):
        if not self.backup_file:
            raise ArcGISRequestError("Backup Anonymous Access tidak tersedia.")
        backup_file = Path(self.backup_file)
        if not backup_file.exists():
            raise ArcGISRequestError(f"Backup file tidak ditemukan: {backup_file}")
        previous = json.loads(backup_file.read_text(encoding="utf-8"))
        access_value = str(previous.get("access", "")).strip().lower()
        if access_value not in ("public", "private"):
            raise ArcGISRequestError("Nilai access pada backup tidak valid.")
        # Mengikuti HAR Portal 11.5: target private mengirim canShareBingPublic=false;
        # target public hanya mengirim access=public.
        include_bing = access_value == "private"
        self.progress.emit(f"Portal Rollback: mengembalikan access={access_value}...")
        rollback_result = self.update_anonymous_access(
            session, token, access_value, include_bing=include_bing
        )
        if not rollback_result.get("success"):
            raise ArcGISRequestError("Portal tidak mengonfirmasi rollback anonymous access.")
        verified = self.read_portal_self(session, token)
        actual = str(verified.get("access", "")).strip().lower()
        verified_ok = actual == access_value
        _, verified_text = self.anonymous_access_state(verified)
        anonymous_probe = self.anonymous_access_probe(session)
        if not verified_ok:
            raise ArcGISRequestError(
                "Rollback selesai, tetapi live verification access tidak sesuai backup."
            )
        metadata_file = backup_file.parent / "metadata.json"
        if metadata_file.exists():
            self.update_metadata(
                metadata_file,
                rollback_response=rollback_result,
                rollback_verified_state=verified_text,
                rollback_anonymous_probe=anonymous_probe,
                status="rolled_back",
                rolled_back_at=datetime.now().astimezone().isoformat(),
            )
        return {
            "success": True,
            "verified_value": verified_text,
            "anonymous_probe": anonymous_probe,
            "backup_file": str(backup_file),
            "portal_token": self.portal_token,
            "portal_token_expires": self.portal_token_expires,
        }

    def read_user_default_settings(self, session, token):
        return self.request_json(
            session, "GET", f"{self.portal_base_url}/sharing/rest/portals/self/userDefaultSettings",
            params={"token": token, "f": "json", "_ts": str(int(time.time() * 1000))},
            headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
        )
    def read_user_license_types(self, session, token):
        return self.request_json(
            session, "GET", f"{self.portal_base_url}/sharing/rest/portals/self/userLicenseTypes",
            params={"token": token, "f": "json", "_ts": str(int(time.time() * 1000))},
            headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
        )
    def read_portal_roles(self, session, token):
        return self.request_json(
            session, "GET", f"{self.portal_base_url}/sharing/rest/portals/self/roles",
            params={"token": token, "start": 1, "num": 100, "f": "json", "_ts": str(int(time.time() * 1000))},
            headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
        )
    @staticmethod
    def collection_items(data, *keys):
        if isinstance(data, list):
            return data
        if isinstance(data, dict):
            for key in keys:
                value = data.get(key)
                if isinstance(value, list):
                    return value
        return []
    def resolve_viewer_defaults(self, license_types, roles):
        license_items = self.collection_items(
            license_types, "userLicenseTypes", "userLicenseType", "results", "items"
        )
        role_items = self.collection_items(roles, "roles", "results", "items")
        viewer_license = next((item for item in license_items if isinstance(item, dict) and (
            str(item.get("id", item.get("userLicenseTypeId", ""))).lower() == "viewerut"
            or str(item.get("name", item.get("title", ""))).strip().lower() == "viewer"
        )), None)
        viewer_role = next((item for item in role_items if isinstance(item, dict) and (
            str(item.get("name", "")).strip().lower() == "viewer"
            or str(item.get("description", "")).strip().lower() == "viewer"
        )), None)
        license_id = None if not viewer_license else str(
            viewer_license.get("id", viewer_license.get("userLicenseTypeId", ""))
        ).strip()
        role_id = None if not viewer_role else str(viewer_role.get("id", "")).strip()
        if not license_id or not role_id:
            raise ArcGISRequestError(
                "Viewer user type atau Viewer role tidak berhasil ditemukan pada katalog Portal. "
                "Apply dihentikan agar tidak menebak technical ID."
            )
        return license_id, role_id
    @staticmethod
    def member_defaults_state(settings, viewer_license_id=None, viewer_role_id=None):
        settings = settings if isinstance(settings, dict) else {}
        license_value = settings.get("userLicenseType")
        role_value = settings.get("role")
        if viewer_license_id and viewer_role_id:
            compliant = (
                str(license_value).strip().lower() == str(viewer_license_id).strip().lower()
                and str(role_value).strip() == str(viewer_role_id).strip()
            )
        else:
            compliant = False
        license_text = license_value or "Not set"
        role_text = role_value or "Not set"
        return compliant, f"User type: {license_text}; Role: {role_text}"
    def create_member_defaults_backup(self, current, viewer_license_id, viewer_role_id):
        execution_id = datetime.now().strftime("%Y%m%d-%H%M%S")
        base_dir = Path.cwd() / "backups" / execution_id
        base_dir.mkdir(parents=True, exist_ok=False)
        backup_file = base_dir / "portal-new-member-defaults-before.json"
        metadata_file = base_dir / "metadata.json"
        safe_backup = {
            "userLicenseType_was_present": "userLicenseType" in current,
            "userLicenseType": current.get("userLicenseType"),
            "role_was_present": "role" in current,
            "role": current.get("role"),
        }
        backup_file.write_text(json.dumps(safe_backup, indent=2), encoding="utf-8")
        metadata_file.write_text(json.dumps({
            "execution_id": execution_id,
            "created_at": datetime.now().astimezone().isoformat(),
            "component": "Portal for ArcGIS",
            "control": "configure_new_member_default_role_as_viewer",
            "before": safe_backup,
            "target": {"userLicenseType": viewer_license_id, "role": viewer_role_id},
            "scope": "new_members_only",
            "restart_expected": False,
            "status": "backup_created",
        }, indent=2), encoding="utf-8")
        return backup_file, metadata_file, execution_id
    def update_member_defaults(self, session, token, user_license_type, role):
        return self.request_json(
            session, "POST", f"{self.portal_base_url}/sharing/rest/portals/self/setUserDefaultSettings",
            data={"userLicenseType": user_license_type, "role": role, "token": token, "f": "json"},
        )
    def apply_member_defaults_viewer(self, session, token):
        self.progress.emit("New Member Defaults: GET latest settings dan resolve Viewer IDs...")
        current = self.read_user_default_settings(session, token)
        viewer_license_id, viewer_role_id = self.resolve_viewer_defaults(
            self.read_user_license_types(session, token), self.read_portal_roles(session, token)
        )
        compliant, current_text = self.member_defaults_state(
            current, viewer_license_id, viewer_role_id
        )
        if compliant:
            return {"success": True, "already_compliant": True, "verified_value": current_text,
                    "verified_settings": current, "portal_token": self.portal_token,
                    "portal_token_expires": self.portal_token_expires}
        backup_file, metadata_file, execution_id = self.create_member_defaults_backup(
            current, viewer_license_id, viewer_role_id
        )
        self.progress.emit("New Member Defaults: set Viewer user type dan Viewer role...")
        response = self.update_member_defaults(
            session, token, viewer_license_id, viewer_role_id
        )
        if not response.get("success"):
            raise ArcGISRequestError("Portal tidak mengonfirmasi update New Member Default Settings.")
        verified = self.read_user_default_settings(session, token)
        verified_ok, verified_text = self.member_defaults_state(
            verified, viewer_license_id, viewer_role_id
        )
        self.update_metadata(metadata_file, apply_response=response,
            verified=verified, status="verified_success" if verified_ok else "verification_failed",
            completed_at=datetime.now().astimezone().isoformat())
        if not verified_ok:
            raise ArcGISRequestError(
                "Update mendapat respons sukses, tetapi live verification Viewer defaults gagal."
            )
        return {"success": True, "already_compliant": False, "execution_id": execution_id,
                "backup_file": str(backup_file), "metadata_file": str(metadata_file),
                "verified_value": verified_text, "verified_settings": verified,
                "portal_token": self.portal_token, "portal_token_expires": self.portal_token_expires}
    def rollback_member_defaults_viewer(self, session, token):
        if not self.backup_file:
            raise ArcGISRequestError("Backup New Member Default Settings tidak tersedia.")
        backup_file = Path(self.backup_file)
        previous = json.loads(backup_file.read_text(encoding="utf-8"))
        if not previous.get("userLicenseType_was_present") or not previous.get("role_was_present"):
            raise ArcGISRequestError(
                "Rollback fail-safe: state sebelumnya adalah Not set. HAR belum memvalidasi payload "
                "untuk menghapus kedua default, sehingga utility tidak akan menebak operasi clear."
            )
        previous_license = previous.get("userLicenseType")
        previous_role = previous.get("role")
        response = self.update_member_defaults(
            session, token, previous_license, previous_role
        )
        if not response.get("success"):
            raise ArcGISRequestError("Portal tidak mengonfirmasi rollback New Member Default Settings.")
        verified = self.read_user_default_settings(session, token)
        ok = (str(verified.get("userLicenseType")) == str(previous_license)
              and str(verified.get("role")) == str(previous_role))
        if not ok:
            raise ArcGISRequestError("Rollback selesai, tetapi live verification tidak sesuai backup.")
        metadata_file = backup_file.parent / "metadata.json"
        if metadata_file.exists():
            self.update_metadata(metadata_file, rollback_response=response, rollback_verified=verified,
                status="rolled_back", rolled_back_at=datetime.now().astimezone().isoformat())
        return {"success": True, "verified_value": f"User type: {previous_license}; Role: {previous_role}",
                "verified_settings": verified, "backup_file": str(backup_file),
                "portal_token": self.portal_token, "portal_token_expires": self.portal_token_expires}

    def read_security_config(self, session, token):
        return self.request_json(
            session,
            "GET",
            f"{self.portal_admin_url}/security/config",
            params={
                "token": token,
                "f": "json",
                "_ts": str(int(time.time() * 1000)),
            },
            headers={"Cache-Control": "no-cache", "Pragma": "no-cache"},
        )

    def create_backup(self, current, control, property_name, target_value):
        execution_id = datetime.now().strftime("%Y%m%d-%H%M%S")
        base_dir = Path.cwd() / "backups" / execution_id
        base_dir.mkdir(parents=True, exist_ok=False)
        backup_file = base_dir / "portal-security-before.json"
        metadata_file = base_dir / "metadata.json"
        backup_file.write_text(
            json.dumps(current, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        metadata = {
            "execution_id": execution_id,
            "created_at": datetime.now().astimezone().isoformat(),
            "component": "Portal for ArcGIS",
            "control": control,
            "portal_admin_url": self.portal_admin_url,
            "before": {property_name: find_value(current, property_name)},
            "target": {property_name: target_value},
            "status": "backup_created",
        }
        metadata_file.write_text(
            json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8"
        )
        return backup_file, metadata_file, execution_id

    def update_metadata(self, metadata_file, **updates):
        metadata = json.loads(metadata_file.read_text(encoding="utf-8"))
        metadata.update(updates)
        metadata_file.write_text(
            json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8"
        )

    def update_security_config(self, session, token, config):
        # Portal API menerima konfigurasi keamanan lengkap sebagai securityConfig JSON.
        return self.request_json(
            session,
            "POST",
            f"{self.portal_admin_url}/security/config/update",
            data={
                "securityConfig": json.dumps(config, ensure_ascii=False),
                "token": token,
                "f": "json",
            },
        )

    def apply_portal_directory(self, session, token):
        self.progress.emit("Portal Apply: membaca security configuration terbaru...")
        current = self.read_security_config(session, token)
        before = find_value(current, "disableServicesDirectory")
        if str(before).lower() == "true":
            return {
                "success": True,
                "already_compliant": True,
                "verified_value": True,
                "portal_token": self.portal_token,
                "portal_token_expires": self.portal_token_expires,
            }

        self.progress.emit("Portal Apply: membuat backup JSON...")
        backup_file, metadata_file, execution_id = self.create_backup(
            current, "disable_portal_directory", "disableServicesDirectory", True
        )
        target_config = dict(current)
        target_config["disableServicesDirectory"] = True

        self.progress.emit("Portal Apply: menonaktifkan Portal Directory...")
        apply_result = self.update_security_config(session, token, target_config)
        self.progress.emit("Portal Verify: membaca ulang security configuration live...")
        verified = self.read_security_config(session, token)
        verified_value = find_value(verified, "disableServicesDirectory")
        verified_ok = str(verified_value).lower() == "true"
        self.update_metadata(
            metadata_file,
            apply_response=apply_result,
            verified_value=verified_value,
            status="verified_success" if verified_ok else "verification_failed",
            completed_at=datetime.now().astimezone().isoformat(),
        )
        if not verified_ok:
            raise ArcGISRequestError(
                "Portal Apply mendapat respons, tetapi live verification gagal. "
                f"Nilai aktual disableServicesDirectory: {verified_value}"
            )
        return {
            "success": True,
            "already_compliant": False,
            "execution_id": execution_id,
            "backup_file": str(backup_file),
            "metadata_file": str(metadata_file),
            "verified_value": verified_value,
            "portal_token": self.portal_token,
            "portal_token_expires": self.portal_token_expires,
        }

    def rollback_portal_directory(self, session, token):
        if not self.backup_file:
            raise ArcGISRequestError("Backup Portal Directory tidak tersedia.")
        backup_file = Path(self.backup_file)
        if not backup_file.exists():
            raise ArcGISRequestError(f"Backup file tidak ditemukan: {backup_file}")
        previous = json.loads(backup_file.read_text(encoding="utf-8"))
        target_value = find_value(previous, "disableServicesDirectory")
        if str(target_value).lower() not in ("true", "false"):
            raise ArcGISRequestError(
                "Nilai disableServicesDirectory pada backup tidak valid."
            )
        current = self.read_security_config(session, token)
        target_config = dict(current)
        target_config["disableServicesDirectory"] = (
            str(target_value).lower() == "true"
        )
        self.progress.emit("Portal Rollback: mengembalikan konfigurasi dari backup...")
        rollback_result = self.update_security_config(session, token, target_config)
        self.progress.emit("Portal Rollback: melakukan live verification...")
        verified = self.read_security_config(session, token)
        verified_value = find_value(verified, "disableServicesDirectory")
        verified_ok = str(verified_value).lower() == str(target_value).lower()
        if not verified_ok:
            raise ArcGISRequestError(
                "Portal rollback selesai, tetapi live verification gagal. "
                f"Nilai aktual: {verified_value}; target backup: {target_value}"
            )
        metadata_file = backup_file.parent / "metadata.json"
        if metadata_file.exists():
            self.update_metadata(
                metadata_file,
                rollback_response=rollback_result,
                rollback_verified_value=verified_value,
                status="rolled_back",
                rolled_back_at=datetime.now().astimezone().isoformat(),
            )
        return {
            "success": True,
            "verified_value": verified_value,
            "backup_file": str(backup_file),
            "portal_token": self.portal_token,
            "portal_token_expires": self.portal_token_expires,
        }

    def apply_automatic_account(self, session, token):
        self.progress.emit("Portal Apply: membaca security configuration terbaru...")
        current = self.read_security_config(session, token)
        before = find_value(current, "enableAutomaticAccountCreation")
        if str(before).lower() == "false":
            return {
                "success": True,
                "already_compliant": True,
                "verified_value": False,
                "portal_token": self.portal_token,
                "portal_token_expires": self.portal_token_expires,
            }

        self.progress.emit("Portal Apply: membuat backup JSON Automatic Enterprise Account Creation...")
        backup_file, metadata_file, execution_id = self.create_backup(
            current,
            "disable_automatic_account_creation",
            "enableAutomaticAccountCreation",
            False,
        )
        target_config = dict(current)
        target_config["enableAutomaticAccountCreation"] = False

        self.progress.emit("Portal Apply: menonaktifkan Automatic Enterprise Account Creation...")
        apply_result = self.update_security_config(session, token, target_config)
        self.progress.emit("Portal Verify: membaca ulang security configuration live...")
        verified = self.read_security_config(session, token)
        verified_value = find_value(verified, "enableAutomaticAccountCreation")
        verified_ok = str(verified_value).lower() == "false"
        self.update_metadata(
            metadata_file,
            apply_response=apply_result,
            verified_value=verified_value,
            status="verified_success" if verified_ok else "verification_failed",
            completed_at=datetime.now().astimezone().isoformat(),
        )
        if not verified_ok:
            raise ArcGISRequestError(
                "Portal Apply mendapat respons, tetapi live verification Automatic Account "
                f"Creation gagal. Nilai aktual: {verified_value}"
            )
        return {
            "success": True,
            "already_compliant": False,
            "execution_id": execution_id,
            "backup_file": str(backup_file),
            "metadata_file": str(metadata_file),
            "verified_value": verified_value,
            "portal_token": self.portal_token,
            "portal_token_expires": self.portal_token_expires,
        }

    def rollback_automatic_account(self, session, token):
        if not self.backup_file:
            raise ArcGISRequestError("Backup Automatic Enterprise Account Creation tidak tersedia.")
        backup_file = Path(self.backup_file)
        if not backup_file.exists():
            raise ArcGISRequestError(f"Backup file tidak ditemukan: {backup_file}")
        previous = json.loads(backup_file.read_text(encoding="utf-8"))
        target_value = find_value(previous, "enableAutomaticAccountCreation")
        if str(target_value).lower() not in ("true", "false"):
            raise ArcGISRequestError(
                "Nilai enableAutomaticAccountCreation pada backup tidak valid."
            )
        current = self.read_security_config(session, token)
        target_config = dict(current)
        target_config["enableAutomaticAccountCreation"] = (
            str(target_value).lower() == "true"
        )
        self.progress.emit(
            "Portal Rollback: mengembalikan Automatic Enterprise Account Creation dari backup..."
        )
        rollback_result = self.update_security_config(session, token, target_config)
        self.progress.emit("Portal Rollback: melakukan live verification...")
        verified = self.read_security_config(session, token)
        verified_value = find_value(verified, "enableAutomaticAccountCreation")
        verified_ok = str(verified_value).lower() == str(target_value).lower()
        if not verified_ok:
            raise ArcGISRequestError(
                "Rollback Automatic Enterprise Account Creation selesai, tetapi live verification gagal. "
                f"Nilai aktual: {verified_value}; target backup: {target_value}"
            )
        metadata_file = backup_file.parent / "metadata.json"
        if metadata_file.exists():
            self.update_metadata(
                metadata_file,
                rollback_response=rollback_result,
                rollback_verified_value=verified_value,
                status="rolled_back",
                rolled_back_at=datetime.now().astimezone().isoformat(),
            )
        return {
            "success": True,
            "verified_value": verified_value,
            "backup_file": str(backup_file),
            "portal_token": self.portal_token,
            "portal_token_expires": self.portal_token_expires,
        }

    def run(self):
        session = requests.Session()
        session.verify = self.verify_ssl
        session.headers.update({
            "User-Agent": "ArcGIS-Hardening-Utility/0.6",
            "Referer": f"{self.portal_base_url}/home/",
        })
        result = {
            "action": self.action,
            "control": self.control,
            "success": False,
        }
        try:
            token = self.ensure_token(session)
            if self.control == "portal_directory" and self.action == "apply":
                result = self.apply_portal_directory(session, token)
            elif self.control == "portal_directory" and self.action == "rollback":
                result = self.rollback_portal_directory(session, token)
            elif self.control == "automatic_account" and self.action == "apply":
                result = self.apply_automatic_account(session, token)
            elif self.control == "automatic_account" and self.action == "rollback":
                result = self.rollback_automatic_account(session, token)
            elif self.control == "builtin_self_creation" and self.action == "apply":
                result = self.apply_builtin_self_creation(session, token)
            elif self.control == "builtin_self_creation" and self.action == "rollback":
                result = self.rollback_builtin_self_creation(session, token)
            elif self.control == "public_profile_sharing" and self.action == "apply":
                result = self.apply_public_profile_sharing(session, token)
            elif self.control == "public_profile_sharing" and self.action == "rollback":
                result = self.rollback_public_profile_sharing(session, token)
            elif self.control == "social_media_links" and self.action == "apply":
                result = self.apply_social_media_links(session, token)
            elif self.control == "social_media_links" and self.action == "rollback":
                result = self.rollback_social_media_links(session, token)
            elif self.control == "anonymous_access" and self.action == "apply":
                result = self.apply_anonymous_access(session, token)
            elif self.control == "anonymous_access" and self.action == "rollback":
                result = self.rollback_anonymous_access(session, token)
            elif self.control == "portal_servlets" and self.action == "apply":
                result = self.apply_portal_servlets(session, token)
            elif self.control == "portal_servlets" and self.action == "rollback":
                result = self.rollback_portal_servlets(session, token)
            elif self.control == "member_defaults_viewer" and self.action == "apply":
                result = self.apply_member_defaults_viewer(session, token)
            elif self.control == "member_defaults_viewer" and self.action == "rollback":
                result = self.rollback_member_defaults_viewer(session, token)
            else:
                raise ArcGISRequestError(
                    f"Action/control Portal tidak dikenal: {self.action}/{self.control}"
                )
            result["action"] = self.action
            result["control"] = self.control
        except Exception as error:
            result = {
                "action": self.action,
                "control": self.control,
                "success": False,
                "error": str(error),
                "portal_token": self.portal_token,
                "portal_token_expires": self.portal_token_expires,
            }
        finally:
            session.close()
            self.completed.emit(result)
            self.finished.emit()


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.worker = None
        self.thread = None
        self.portal_result = None
        self.server_result = None
        self.last_portal_directory_backup_file = None
        self.last_automatic_account_backup_file = None
        self.last_builtin_self_creation_backup_file = None
        self.last_public_profile_sharing_backup_file = None
        self.last_social_media_links_backup_file = None
        self.last_anonymous_access_backup_file = None
        self.last_portal_servlets_backup_file = None
        self.last_member_defaults_viewer_backup_file = None
        self.portal_servlet_selection = ()
        self.last_jsonp_backup_file = None
        self.last_services_directory_backup_file = None
        self.last_standardized_queries_backup_file = None
        self.last_feature_service_xss_backup_file = None
        self.last_feature_service_xss_backup_file = None
        self.last_token_http_get_backup_file = None
        self.portal_token = None
        self.portal_token_expires = 0
        self.server_token = None
        self.server_token_expires = 0
        self.connected = False
        self.last_assessment_at = None
        self.trusted_origins = []
        self.control_rows = {}
        self.action_buttons = {}
        self.portal_directory_row = None
        self.automatic_account_row = None
        self.builtin_self_creation_row = None
        self.public_profile_sharing_row = None
        self.social_media_links_row = None
        self.anonymous_access_row = None
        self.portal_servlets_row = None
        self.member_defaults_viewer_row = None
        self.https_enforcement_row = None
        self.jsonp_row = None
        self.services_directory_row = None
        self.standardized_queries_row = None
        self.feature_service_xss_row = None
        self.feature_service_xss_row = None
        self.token_http_get_row = None
        self.last_portal_directory_backup_file = None
        self.last_automatic_account_backup_file = None
        self.last_builtin_self_creation_backup_file = None
        self.last_public_profile_sharing_backup_file = None
        self.last_social_media_links_backup_file = None
        self.last_anonymous_access_backup_file = None
        self.last_jsonp_backup_file = None
        self.last_services_directory_backup_file = None
        self.last_standardized_queries_backup_file = None
        self.last_token_http_get_backup_file = None
        self.hardening_worker = None
        self.hardening_thread = None
        self.portal_token = None
        self.portal_token_expires = 0
        self.server_token = None
        self.server_token_expires = 0
        self.setWindowTitle("ArcGIS Enterprise Hardening Utility")
        self.resize(1500, 900)
        self.setMinimumSize(1000, 700)
        self.settings = QSettings("itspamud", "ArcGIS Enterprise Hardening Utility")
        self.theme_preference = self.settings.value("appearance/theme", "System", type=str)
        if self.theme_preference not in ("System", "Dark", "Light"):
            self.theme_preference = "System"
        self.current_effective_theme = "Dark"
        self.build_ui()
        self.apply_theme(self.theme_preference, save=False)

    def build_ui(self):
        root = QWidget()
        self.setCentralWidget(root)
        layout = QVBoxLayout(root)
        layout.setContentsMargins(20, 18, 20, 10)
        header_layout = QHBoxLayout()
        header_layout.setContentsMargins(0, 0, 0, 0)
        header_layout.setSpacing(16)
        header_text_layout = QVBoxLayout()
        header_text_layout.setContentsMargins(0, 0, 0, 0)
        header_text_layout.setSpacing(4)
        title = QLabel("ArcGIS Enterprise Hardening Utility")
        title.setObjectName("titleLabel")
        subtitle = QLabel("Connection dan live hardening analysis untuk Portal dan ArcGIS Server.")
        subtitle.setObjectName("subtitleLabel")
        header_text_layout.addWidget(title)
        header_text_layout.addWidget(subtitle)
        header_layout.addLayout(header_text_layout, 1)

        theme_layout = QHBoxLayout()
        theme_layout.setContentsMargins(0, 2, 0, 0)
        theme_layout.setSpacing(7)
        theme_label = QLabel("Theme:")
        theme_label.setObjectName("themeLabel")
        self.theme_combo = QComboBox()
        self.theme_combo.setObjectName("themeCombo")
        self.theme_combo.addItems(["System", "Dark", "Light"])
        self.theme_combo.setFixedWidth(132)
        self.theme_combo.setToolTip(
            "Choose application appearance. System follows the Windows color scheme."
        )
        self.theme_combo.setCurrentText(self.theme_preference)
        self.theme_combo.currentTextChanged.connect(self.on_theme_changed)
        theme_layout.addWidget(theme_label)
        theme_layout.addWidget(self.theme_combo)
        header_layout.addLayout(theme_layout)
        layout.addLayout(header_layout)

        self.tabs = QTabWidget()
        self.connection_tab = QWidget()
        self.hardening_tab = QWidget()
        self.log_tab = QWidget()
        self.tabs.addTab(self.connection_tab, "Connection")
        self.tabs.addTab(self.hardening_tab, "Hardening Controls")
        self.tabs.addTab(self.log_tab, "Execution Log")
        self.tabs.setTabEnabled(1, False)
        layout.addWidget(self.tabs)
        self.build_connection_tab()
        self.build_hardening_tab()
        self.build_log_tab()
        self.status = QStatusBar()
        self.status.showMessage("Ready")
        self.setStatusBar(self.status)

    def card(self, title_text, description):
        frame = QFrame()
        frame.setObjectName("card")
        layout = QVBoxLayout(frame)
        title = QLabel(title_text)
        title.setObjectName("sectionTitle")
        desc = QLabel(description)
        desc.setObjectName("sectionDescription")
        desc.setWordWrap(True)
        layout.addWidget(title)
        layout.addWidget(desc)
        return frame, layout

    def build_connection_tab(self):
        outer = QVBoxLayout(self.connection_tab)
        outer.setContentsMargins(12, 12, 12, 12)
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setFrameShape(QFrame.Shape.NoFrame)
        page = QWidget()
        scroll.setWidget(page)
        forms = QVBoxLayout(page)
        forms.setSpacing(16)
        outer.addWidget(scroll, 1)

        portal_card, portal_layout = self.card(
            "Portal for ArcGIS", "Web Adaptor atau direct port 7443."
        )
        portal_form = QFormLayout()
        self.portal_admin_url = QLineEdit()
        self.portal_admin_url.setPlaceholderText(
            "https://domain/arcgis/portaladmin atau https://portal:7443/arcgis/portaladmin"
        )
        self.portal_username = QLineEdit()
        self.portal_password = QLineEdit()
        self.portal_password.setEchoMode(QLineEdit.EchoMode.Password)
        self.show_portal_password = QCheckBox("Tampilkan Portal password")
        self.show_portal_password.toggled.connect(
            lambda checked: self.portal_password.setEchoMode(
                QLineEdit.EchoMode.Normal if checked else QLineEdit.EchoMode.Password
            )
        )
        portal_form.addRow("Portal Admin URL:", self.portal_admin_url)
        portal_form.addRow("Portal Username:", self.portal_username)
        portal_form.addRow("Portal Password:", self.portal_password)
        portal_form.addRow("", self.show_portal_password)
        portal_layout.addLayout(portal_form)
        self.portal_status = QLabel("Portal status: Belum diuji.")
        self.portal_status.setObjectName("resultLabel")
        self.portal_status.setWordWrap(True)
        self.portal_status.setTextFormat(Qt.TextFormat.RichText)
        portal_layout.addWidget(self.portal_status)
        forms.addWidget(portal_card)

        server_card, server_layout = self.card(
            "ArcGIS Server", "Web Adaptor atau direct port 6443. Gunakan Primary Site Administrator."
        )
        server_form = QFormLayout()
        self.server_admin_url = QLineEdit()
        self.server_admin_url.setPlaceholderText(
            "https://domain/server/admin atau https://server:6443/arcgis/admin"
        )
        self.server_username = QLineEdit()
        self.server_password = QLineEdit()
        self.server_password.setEchoMode(QLineEdit.EchoMode.Password)
        self.show_server_password = QCheckBox("Tampilkan Server password")
        self.show_server_password.toggled.connect(
            lambda checked: self.server_password.setEchoMode(
                QLineEdit.EchoMode.Normal if checked else QLineEdit.EchoMode.Password
            )
        )
        server_form.addRow("Server Admin URL:", self.server_admin_url)
        server_form.addRow("Server Username:", self.server_username)
        server_form.addRow("Server Password:", self.server_password)
        server_form.addRow("", self.show_server_password)
        server_layout.addLayout(server_form)
        self.server_status = QLabel("Server status: Belum diuji.")
        self.server_status.setObjectName("resultLabel")
        self.server_status.setWordWrap(True)
        self.server_status.setTextFormat(Qt.TextFormat.RichText)
        server_layout.addWidget(self.server_status)
        forms.addWidget(server_card)
        forms.addStretch()

        footer = QFrame()
        footer.setObjectName("fixedFooter")
        footer_layout = QVBoxLayout(footer)
        self.verify_ssl = QCheckBox("Verifikasi sertifikat TLS/SSL")
        footer_layout.addWidget(self.verify_ssl)
        buttons = QHBoxLayout()
        self.clear_button = QPushButton("Clear")
        self.test_portal_button = QPushButton("Test Portal")
        self.test_server_button = QPushButton("Test Server")
        self.test_all_button = QPushButton("Test All")
        self.test_all_button.setObjectName("primaryButton")
        self.clear_button.clicked.connect(self.clear_fields)
        self.test_portal_button.clicked.connect(lambda: self.start_request("portal", "connection"))
        self.test_server_button.clicked.connect(lambda: self.start_request("server", "connection"))
        self.test_all_button.clicked.connect(lambda: self.start_request("all", "connection"))
        buttons.addWidget(self.clear_button)
        buttons.addStretch()
        buttons.addWidget(self.test_portal_button)
        buttons.addWidget(self.test_server_button)
        buttons.addWidget(self.test_all_button)
        footer_layout.addLayout(buttons)
        self.activity = QLabel("Activity: Ready")
        self.activity.setObjectName("activityLabel")
        footer_layout.addWidget(self.activity)
        outer.addWidget(footer, 0)

        for field in (
            self.portal_admin_url, self.portal_username,
            self.server_admin_url, self.server_username,
        ):
            field.textChanged.connect(self.invalidate_connection)

    def build_hardening_tab(self):
        layout = QVBoxLayout(self.hardening_tab)
        header = QLabel("ArcGIS Enterprise Hardening Controls")
        header.setObjectName("sectionTitle")
        note = QLabel(
            "Setiap kontrol memiliki action sendiri: Preview, Apply, dan Rollback. "
            "Apply aktif hanya untuk kontrol NON-COMPLIANT yang sudah didukung; "
            "Rollback aktif jika backup kontrol tersedia."
        )
        note.setObjectName("sectionDescription")
        note.setWordWrap(True)
        layout.addWidget(header)
        layout.addWidget(note)

        self.last_refreshed = QLabel("Data source: belum dianalisis")
        self.last_refreshed.setObjectName("activityLabel")
        layout.addWidget(self.last_refreshed)

        self.controls_table = QTableWidget(0, 7)
        self.controls_table.setHorizontalHeaderLabels(
            ["Component", "Control", "Current", "Target", "Status", "Risk", "Actions"]
        )
        header_view = self.controls_table.horizontalHeader()
        header_view.setStretchLastSection(False)
        for column in range(7):
            header_view.setSectionResizeMode(column, QHeaderView.ResizeMode.Fixed)
        self.controls_table.verticalHeader().setVisible(False)
        # Control, Current, dan Target dipertahankan satu baris. Saat ruang sempit,
        # Qt memakai ellipsis; tooltip selalu menyimpan nilai lengkap.
        self.controls_table.setWordWrap(False)
        self.controls_table.setTextElideMode(Qt.TextElideMode.ElideRight)
        self.controls_table.setHorizontalScrollBarPolicy(
            Qt.ScrollBarPolicy.ScrollBarAsNeeded
        )
        self.table_resize_timer = QTimer(self)
        self.table_resize_timer.setSingleShot(True)
        self.table_resize_timer.setInterval(120)
        self.table_resize_timer.timeout.connect(
            self.update_responsive_table_columns
        )
        self.controls_table.setEditTriggers(QTableWidget.EditTrigger.NoEditTriggers)
        self.controls_table.setSelectionMode(QTableWidget.SelectionMode.NoSelection)
        layout.addWidget(self.controls_table)

        footer = QHBoxLayout()
        self.analyze_button = QPushButton("Analyze Hardening")
        self.analyze_button.setObjectName("primaryButton")
        self.analyze_button.clicked.connect(lambda: self.start_request("all", "analyze"))
        footer.addWidget(self.analyze_button)
        footer.addStretch()
        self.export_assessment_button = QPushButton("Export Assessment")
        self.export_assessment_button.setToolTip(
            "Export hasil Analyze terakhir ke Word, Excel, dan JSON evidence package."
        )
        self.export_assessment_button.setEnabled(False)
        self.export_assessment_button.clicked.connect(self.export_current_assessment)
        footer.addWidget(self.export_assessment_button)
        layout.addLayout(footer)

    def build_log_tab(self):
        layout = QVBoxLayout(self.log_tab)
        layout.setContentsMargins(12, 12, 12, 12)
        self.log_label = QPlainTextEdit()
        self.log_label.setObjectName("executionLog")
        self.log_label.setReadOnly(True)
        self.log_label.setLineWrapMode(QPlainTextEdit.LineWrapMode.NoWrap)
        self.log_label.setPlaceholderText("Execution log akan muncul di sini.")
        self.log_label.setVerticalScrollBarPolicy(
            Qt.ScrollBarPolicy.ScrollBarAlwaysOn
        )
        self.log_label.setHorizontalScrollBarPolicy(
            Qt.ScrollBarPolicy.ScrollBarAsNeeded
        )
        layout.addWidget(self.log_label, 1)

    def append_log(self, message):
        stamp = datetime.now().strftime("%H:%M:%S")
        scrollbar = self.log_label.verticalScrollBar()
        # Auto-follow hanya jika pengguna sebelumnya berada dekat bagian paling bawah.
        follow_tail = scrollbar.value() >= max(0, scrollbar.maximum() - 4)
        self.log_label.appendPlainText(f"[{stamp}] {message}")
        if follow_tail:
            scrollbar.setValue(scrollbar.maximum())

    def validate_url(self, value, ending, name):
        parsed = urlparse(value)
        if parsed.scheme.lower() != "https" or not parsed.hostname:
            return f"{name} harus berupa URL HTTPS yang valid."
        if not parsed.path.rstrip("/").lower().endswith(ending):
            return f"{name} harus berakhir dengan {ending}."
        return None

    def validate_mode(self, mode):
        if mode in ("portal", "all"):
            if not all([
                self.portal_admin_url.text().strip(),
                self.portal_username.text().strip(),
                self.portal_password.text(),
            ]):
                return "Lengkapi seluruh data Portal."
            error = self.validate_url(
                self.portal_admin_url.text(), "/portaladmin", "Portal Admin URL"
            )
            if error:
                return error
        if mode in ("server", "all"):
            if not all([
                self.server_admin_url.text().strip(),
                self.server_username.text().strip(),
                self.server_password.text(),
            ]):
                return "Lengkapi seluruh data ArcGIS Server."
            error = self.validate_url(
                self.server_admin_url.text(), "/admin", "Server Admin URL"
            )
            if error:
                return error
        return None

    def set_busy(self, busy):
        for widget in (
            self.portal_admin_url, self.portal_username, self.portal_password,
            self.server_admin_url, self.server_username, self.server_password,
            self.show_portal_password, self.show_server_password, self.verify_ssl,
            self.clear_button, self.test_portal_button, self.test_server_button,
            self.test_all_button, self.analyze_button, self.export_assessment_button,
        ):
            widget.setEnabled(not busy)

        if busy:
            for buttons in self.action_buttons.values():
                for button in buttons.values():
                    button.setEnabled(False)
        else:
            self.refresh_action_buttons()

    def start_request(self, mode, purpose):
        if self.thread is not None:
            return
        error = self.validate_mode(mode)
        if error:
            QMessageBox.warning(self, "Data belum valid", error)
            return
        self.set_busy(True)
        if purpose == "analyze":
            self.last_refreshed.setText("Data source: refreshing live configuration...")
            self.append_log("Live Analyze dimulai.")
        self.thread = QThread()
        self.worker = ConnectionWorker(
            mode, purpose,
            self.portal_admin_url.text(), self.portal_username.text(), self.portal_password.text(),
            self.server_admin_url.text(), self.server_username.text(), self.server_password.text(),
            self.verify_ssl.isChecked(),
            portal_token=self.portal_token,
            portal_token_expires=self.portal_token_expires,
            server_token=self.server_token,
            server_token_expires=self.server_token_expires,
            run_network_probes=(purpose == "analyze"),
        )
        self.worker.moveToThread(self.thread)
        self.thread.started.connect(self.worker.run)
        self.worker.progress.connect(self.on_progress)
        self.worker.completed.connect(self.on_completed)
        self.worker.finished.connect(self.thread.quit)
        self.worker.finished.connect(self.worker.deleteLater)
        self.thread.finished.connect(self.thread.deleteLater)
        self.thread.finished.connect(self.on_worker_finished)
        self.thread.start()

    def on_progress(self, message):
        self.activity.setText(f"Activity: {message}")
        self.status.showMessage(message)
        self.append_log(message)

    def component_summary(self, name, result):
        if result is None:
            return f"{name}: TIDAK DIUJI"
        if result.get("success"):
            return f"{name}: BERHASIL"
        first_line = result.get("error", "Penyebab tidak tersedia.").splitlines()[0]
        return f"{name}: GAGAL\nPenyebab: {first_line}"

    def on_completed(self, result):
        portal = result.get("portal")
        server = result.get("server")
        purpose = result.get("purpose")
        token_cache = result.get("token_cache", {})
        self.portal_token = token_cache.get("portal_token")
        self.portal_token_expires = token_cache.get("portal_token_expires", 0)
        self.server_token = token_cache.get("server_token")
        self.server_token_expires = token_cache.get("server_token_expires", 0)
        portal_ok = bool(portal and portal.get("success"))
        server_ok = bool(server and server.get("success"))

        if purpose == "analyze":
            if portal_ok and server_ok:
                self.portal_result = portal
                self.server_result = server
                self.update_connection_results()
                self.render_controls()
                self.last_assessment_at = datetime.now().astimezone()
                self.export_assessment_button.setEnabled(True)
                stamp = self.last_assessment_at.strftime("%d %b %Y %H:%M:%S")
                self.last_refreshed.setText(
                    f"Data source: LIVE | Last analyzed: {stamp} | Portal: OK | Server: OK"
                )
                self.status.showMessage("Live Analyze selesai", 10000)
                self.activity.setText("Activity: Live analysis selesai. Data terbaru ditampilkan.")
                self.append_log("Live Analyze selesai; tabel diperbarui dari data terbaru tanpa cache.")
            else:
                summary = "\n\n".join([
                    self.component_summary("Portal for ArcGIS", portal),
                    self.component_summary("ArcGIS Server", server),
                ])
                self.last_refreshed.setText("Data source: refresh gagal; tabel lama tidak diubah")
                QMessageBox.critical(
                    self, "Analyze Hardening gagal",
                    "Konfigurasi live tidak berhasil dibaca dari seluruh komponen.\n\n"
                    f"{summary}\n\nTabel lama tidak diperbarui."
                )
                self.append_log("Live Analyze gagal; tabel lama dipertahankan.")
            return

        if portal_ok:
            self.portal_result = portal
        if server_ok:
            self.server_result = server
        self.update_connection_results()

        tested = [item for item in (portal, server) if item is not None]
        all_success = tested and all(item.get("success") for item in tested)
        if len(tested) == 2 and all_success:
            self.connected = True
            self.tabs.setTabEnabled(1, True)
            QMessageBox.information(
                self, "Test All berhasil",
                "Portal for ArcGIS: BERHASIL\n\nArcGIS Server: BERHASIL\n\n"
                "Hardening Controls sudah dibuka."
            )
        elif all_success:
            name = "Portal for ArcGIS" if portal is not None else "ArcGIS Server"
            QMessageBox.information(
                self, f"Test {name} berhasil", f"Komponen yang diuji: {name}\nStatus: BERHASIL"
            )
        else:
            summary = "\n\n".join(
                self.component_summary(name, item)
                for name, item in (("Portal for ArcGIS", portal), ("ArcGIS Server", server))
                if item is not None
            )
            QMessageBox.warning(self, "Pengujian koneksi tidak berhasil", summary)

    def connection_badge_html(self, state):
        dark = self.current_effective_theme == "Dark"
        palettes = {
            "success": (
                "#9be7bf" if dark else "#d9f2e3",
                "#124f2c" if dark else "#215c35",
                "BERHASIL",
            ),
            "partial": (
                "#f4d58d" if dark else "#f7e8bd",
                "#634600" if dark else "#6a4b00",
                "BERHASIL SEBAGIAN",
            ),
            "failed": (
                "#f3a6a6" if dark else "#f6d4d4",
                "#6f1d1d" if dark else "#8b2424",
                "GAGAL",
            ),
        }
        _background, foreground, label = palettes.get(state, palettes["partial"])
        return (
            f'<span style="color:{foreground}; font-weight:700;">'
            f'&#9679;&nbsp;{label}</span>'
        )

    @staticmethod
    def connection_value(value):
        if value is True:
            return "Enabled"
        if value is False:
            return "Disabled"
        text = str(value)
        return text if text and text != "None" else "Tidak tersedia"

    def update_connection_results(self):
        checked_at = datetime.now().strftime("%d %b %Y %H:%M:%S")

        if self.portal_result and self.portal_result.get("success"):
            p = self.portal_result
            portal_partial = not p.get("version_available", False)
            portal_state = "partial" if portal_partial else "success"
            token_valid = bool(self.portal_token) and self.portal_token_expires > time.time()
            access = str(p.get("access", "Tidak tersedia")).capitalize()
            https_only = self.connection_value(p.get("all_ssl", "Tidak tersedia"))
            self.portal_status.setText(
                '<p style="margin:0 0 12px 0;">'
                '<b>PORTAL CONNECTION</b>&nbsp;&nbsp;&nbsp;'
                f'{self.connection_badge_html(portal_state)}</p>'
                '<table cellspacing="0" cellpadding="2">'
                f'<tr><td><b>Version</b></td><td>&nbsp;:&nbsp; {p.get("version", "Tidak tersedia")}</td></tr>'
                f'<tr><td><b>Portal access</b></td><td>&nbsp;:&nbsp; {access}</td></tr>'
                f'<tr><td><b>HTTPS only</b></td><td>&nbsp;:&nbsp; {https_only}</td></tr>'
                '<tr><td><b>Administrator API</b></td><td>&nbsp;:&nbsp; Accessible</td></tr>'
                f'<tr><td><b>Token status</b></td><td>&nbsp;:&nbsp; {"Valid" if token_valid else "Expired / unavailable"}</td></tr>'
                f'<tr><td><b>Last checked</b></td><td>&nbsp;:&nbsp; {checked_at}</td></tr>'
                '</table>'
            )

        if self.server_result and self.server_result.get("success"):
            s = self.server_result
            required = (s.get("version"), s.get("full_version"), s.get("authentication_tier"))
            server_partial = any(
                value in (None, "", "Tidak tersedia") for value in required
            )
            server_state = "partial" if server_partial else "success"
            token_valid = bool(self.server_token) and self.server_token_expires > time.time()
            version_text = (
                f'{s.get("version", "Tidak tersedia")} / '
                f'{s.get("full_version", "Tidak tersedia")}'
            )
            self.server_status.setText(
                '<p style="margin:0 0 12px 0;">'
                '<b>SERVER CONNECTION</b>&nbsp;&nbsp;&nbsp;'
                f'{self.connection_badge_html(server_state)}</p>'
                '<table cellspacing="0" cellpadding="2">'
                f'<tr><td><b>Version / Full version</b></td><td>&nbsp;:&nbsp; {version_text}</td></tr>'
                f'<tr><td><b>Authentication tier</b></td><td>&nbsp;:&nbsp; {s.get("authentication_tier", "Tidak tersedia")}</td></tr>'
                '<tr><td><b>Administrator API</b></td><td>&nbsp;:&nbsp; Accessible</td></tr>'
                f'<tr><td><b>Token status</b></td><td>&nbsp;:&nbsp; {"Valid" if token_valid else "Expired / unavailable"}</td></tr>'
                f'<tr><td><b>Last checked</b></td><td>&nbsp;:&nbsp; {checked_at}</td></tr>'
                '</table>'
            )

    def on_worker_finished(self):
        # Bersihkan referensi thread terlebih dahulu. refresh_action_buttons()
        # menganggap aplikasi masih busy selama self.thread belum None.
        self.worker = None
        self.thread = None
        self.set_busy(False)

    def compliance(self, current, target):
        if isinstance(current, bool):
            return "COMPLIANT" if current == target else "NON-COMPLIANT"
        value = str(current).lower()
        if value in ("true", "false"):
            return "COMPLIANT" if (value == "true") == target else "NON-COMPLIANT"
        return "UNKNOWN"

    def backup_for_control(self, control_id):
        mapping = {
            "portal_directory": self.last_portal_directory_backup_file,
            "automatic_account": self.last_automatic_account_backup_file,
            "builtin_self_creation": self.last_builtin_self_creation_backup_file,
            "public_profile_sharing": self.last_public_profile_sharing_backup_file,
            "social_media_links": self.last_social_media_links_backup_file,
            "anonymous_access": self.last_anonymous_access_backup_file,
            "portal_servlets": self.last_portal_servlets_backup_file,
            "member_defaults_viewer": self.last_member_defaults_viewer_backup_file,
            "jsonp": self.last_jsonp_backup_file,
            "services_directory": self.last_services_directory_backup_file,
            "standardized_queries": self.last_standardized_queries_backup_file,
            "feature_service_xss": self.last_feature_service_xss_backup_file,
            "token_http_get": self.last_token_http_get_backup_file,
        }
        return mapping.get(control_id)

    def control_supported(self, control_id):
        return control_id in (
            "portal_directory", "automatic_account", "builtin_self_creation",
            "public_profile_sharing", "social_media_links", "anonymous_access", "portal_servlets",
            "member_defaults_viewer", "jsonp",
            "services_directory", "standardized_queries", "feature_service_xss", "token_http_get"
        )

    @staticmethod
    def wrap_table_words(value, words_per_line):
        """Wrap plain-text table content after a fixed number of whitespace words."""
        raw = str(value)
        # Preserve intentional newlines, but wrap each segment independently.
        wrapped_segments = []
        for segment in raw.splitlines() or [raw]:
            words = segment.split()
            if not words:
                wrapped_segments.append("")
                continue
            lines = [
                " ".join(words[index:index + words_per_line])
                for index in range(0, len(words), words_per_line)
            ]
            wrapped_segments.append("\n".join(lines))
        return "\n".join(wrapped_segments)

    def add_control(self, control_id, component, control, current, target, status, risk):
        row = self.controls_table.rowCount()
        self.controls_table.insertRow(row)
        self.control_rows[control_id] = row
        if control_id == "portal_directory":
            self.portal_directory_row = row
        elif control_id == "automatic_account":
            self.automatic_account_row = row
        elif control_id == "builtin_self_creation":
            self.builtin_self_creation_row = row
        elif control_id == "public_profile_sharing":
            self.public_profile_sharing_row = row
        elif control_id == "social_media_links":
            self.social_media_links_row = row
        elif control_id == "anonymous_access":
            self.anonymous_access_row = row
        elif control_id == "portal_servlets":
            self.portal_servlets_row = row
        elif control_id == "member_defaults_viewer":
            self.member_defaults_viewer_row = row
        elif control_id == "https_enforcement":
            self.https_enforcement_row = row
        elif control_id == "jsonp":
            self.jsonp_row = row
        elif control_id == "services_directory":
            self.services_directory_row = row
        elif control_id == "standardized_queries":
            self.standardized_queries_row = row
        elif control_id == "feature_service_xss":
            self.feature_service_xss_row = row
        elif control_id == "token_http_get":
            self.token_http_get_row = row

        # Responsive columns need plain single-line values. Collapse intentional
        # newlines/spaces for display, while preserving the original value in tooltip.
        display_control = " ".join(str(control).split())
        display_current = " ".join(bool_text(current).split())
        display_target = " ".join(bool_text(target).split())
        display_status = (
            "COMPLIANT\nWITH NOTE"
            if status == "COMPLIANT WITH NOTE"
            else self.wrap_table_words(status, 1)
        )
        display_values = [
            str(component), display_control, display_current,
            display_target, display_status, str(risk),
        ]
        tooltip_values = [
            str(component), str(control), bool_text(current),
            bool_text(target), str(status), str(risk),
        ]
        for col, value in enumerate(display_values):
            item = QTableWidgetItem(str(value))
            item.setToolTip(tooltip_values[col])
            item.setTextAlignment(
                Qt.AlignmentFlag.AlignLeft | Qt.AlignmentFlag.AlignVCenter
            )
            if col == 4:
                # Simpan status asli untuk logic tombol; teks tampil boleh mengandung newline.
                item.setData(Qt.ItemDataRole.UserRole, status)
                item.setForeground(self.semantic_status_color(status))
            self.controls_table.setItem(row, col, item)

        action_widget = QWidget()
        action_layout = QHBoxLayout(action_widget)
        action_layout.setContentsMargins(2, 2, 2, 2)
        action_layout.setSpacing(3)

        preview_button = QPushButton("Preview")
        preview_button.setObjectName("rowPreviewButton")
        preview_button.setToolTip("Lihat current value, target, dampak, dan safety workflow.")
        preview_button.clicked.connect(
            lambda _checked=False, cid=control_id: self.preview_control(cid)
        )

        apply_button = QPushButton("Apply")
        apply_button.setObjectName("rowApplyButton")
        apply_button.clicked.connect(
            lambda _checked=False, cid=control_id: self.apply_control(cid)
        )

        rollback_button = QPushButton("Rollback")
        rollback_button.setObjectName("rowRollbackButton")
        rollback_button.clicked.connect(
            lambda _checked=False, cid=control_id: self.rollback_control(cid)
        )

        action_layout.addWidget(preview_button)
        action_layout.addWidget(apply_button)
        action_layout.addWidget(rollback_button)
        self.controls_table.setCellWidget(row, 6, action_widget)
        status_lines = display_status.count("\n") + 1
        self.controls_table.setRowHeight(row, max(44, 20 * status_lines + 10))

        self.action_buttons[control_id] = {
            "preview": preview_button,
            "apply": apply_button,
            "rollback": rollback_button,
        }

    @staticmethod
    def normalize_arcgis_https_protocol(value, http_enabled=None, ssl_enabled=None):
        text = str(value).strip().lower().replace("_", " ").replace("-", " ")
        if text in ("https", "https only"):
            return "PASS", "HTTPS Only", "protocol"
        if "http" in text and "https" in text:
            return "FAIL", "HTTP and HTTPS", "protocol"
        if text in ("http", "http only"):
            return "CRITICAL", "HTTP Only", "protocol"
        http_text = str(http_enabled).strip().lower()
        ssl_text = str(ssl_enabled).strip().lower()
        if http_text in ("true", "false") and ssl_text in ("true", "false"):
            if http_text == "false" and ssl_text == "true":
                return "PASS", "HTTPS Only", "httpEnabled/sslEnabled"
            if http_text == "true" and ssl_text == "true":
                return "FAIL", "HTTP and HTTPS", "httpEnabled/sslEnabled"
            if http_text == "true" and ssl_text == "false":
                return "CRITICAL", "HTTP Only", "httpEnabled/sslEnabled"
        return "UNKNOWN", str(value), "unavailable"
    @staticmethod
    def https_probe_status(probe):
        http_result = str((probe or {}).get("http_result", "UNKNOWN"))
        https_result = str((probe or {}).get("https_result", "UNKNOWN"))
        if http_result == "REDIRECTS_TO_NON_HTTPS":
            return "CRITICAL"
        if http_result == "PLAINTEXT_CONTENT_AVAILABLE":
            return "FAIL"
        if https_result != "REACHABLE":
            return "UNKNOWN"
        if http_result == "REDIRECTS_TO_HTTPS":
            return "PASS"
        if http_result in ("CONNECT_TIMEOUT", "NOT_EXPOSED_FROM_TEST_LOCATION"):
            return "PASS WITH SCOPE NOTE"
        return "UNKNOWN"
    def https_enforcement_evidence(self):
        portal = self.portal_result or {}
        server = self.server_result or {}
        portal_ssl = str(portal.get("all_ssl", "")).strip().lower()
        portal_policy = "PASS" if portal_ssl == "true" else "FAIL" if portal_ssl == "false" else "UNKNOWN"
        server_policy, server_protocol_text, protocol_source = self.normalize_arcgis_https_protocol(
            server.get("protocol", "Tidak tersedia"),
            server.get("http_enabled"), server.get("ssl_enabled"),
        )
        records = []
        for component, result in (("Portal", portal), ("Server", server)):
            for index, adaptor in enumerate(result.get("web_adaptors", []), start=1):
                entry = dict(adaptor)
                entry["component"] = component
                entry["number"] = index
                entry["status"] = self.https_probe_status(entry.get("probe", {}))
                records.append(entry)
        checks = [portal_policy, server_policy] + [item["status"] for item in records]
        inventory_errors = bool(portal.get("web_adaptors_error") or server.get("web_adaptors_error"))
        if "CRITICAL" in checks:
            overall = "CRITICAL"
        elif "FAIL" in checks:
            overall = "NON-COMPLIANT"
        elif not records or inventory_errors or "UNKNOWN" in checks:
            overall = "PARTIALLY VERIFIED"
        elif "PASS WITH SCOPE NOTE" in checks:
            overall = "COMPLIANT WITH NOTE"
        else:
            overall = "COMPLIANT"
        return {
            "portal_policy": portal_policy,
            "portal_all_ssl": portal.get("all_ssl", "Tidak tersedia"),
            "server_policy": server_policy,
            "server_protocol": server_protocol_text,
            "server_protocol_source": protocol_source,
            "server_http_enabled": server.get("http_enabled", "Tidak tersedia"),
            "server_ssl_enabled": server.get("ssl_enabled", "Tidak tersedia"),
            "server_hsts_enabled": server.get("hsts_enabled", "Tidak tersedia"),
            "records": records,
            "portal_inventory_error": portal.get("web_adaptors_error"),
            "server_inventory_error": server.get("web_adaptors_error"),
            "overall": overall,
        }
    @staticmethod
    def https_result_explanation(result, component="endpoint"):
        explanations = {
            "REDIRECTS_TO_HTTPS": (
                "Request HTTP dialihkan ke HTTPS. Perilaku ini memenuhi baseline selama "
                "endpoint HTTPS dapat dijangkau dan tidak terjadi downgrade kembali ke HTTP."
            ),
            "REDIRECTS_TO_NON_HTTPS": (
                "Request HTTP memang dialihkan, tetapi tujuan redirect masih menggunakan HTTP. "
                "Redirect hanya mengubah path/format URL dan belum memindahkan komunikasi ke HTTPS."
            ),
            "PLAINTEXT_CONTENT_AVAILABLE": (
                "Konten masih dapat dilayani melalui HTTP tanpa dialihkan ke HTTPS. "
                "Komunikasi dapat berlangsung tanpa enkripsi."
            ),
            "NOT_EXPOSED_FROM_TEST_LOCATION": (
                "Endpoint HTTP tidak dapat dijangkau dari komputer operator, sedangkan hasil ini "
                "tidak membuktikan akses dari VLAN, proxy, load balancer, atau lokasi jaringan lain."
            ),
            "CONNECT_TIMEOUT": (
                "Endpoint HTTP tidak dapat dijangkau dari lokasi pengujian dalam 4 detik. "
                "Jika HTTPS tersedia, kondisi ini konsisten dengan baseline 443-only, dengan scope note."
            ),
            "READ_TIMEOUT": (
                "Koneksi terbentuk, tetapi respons tidak selesai dalam 8 detik. Periksa Web Adaptor, "
                "web server, proxy, load balancer, dan performa endpoint."
            ),
            "TLS_VERIFICATION_FAILED": (
                "Endpoint HTTPS tersedia, tetapi certificate tidak dapat diverifikasi dari komputer operator. "
                "Periksa trust chain, hostname, masa berlaku, dan CA certificate."
            ),
            "REACHABLE": (
                "Endpoint HTTPS memberikan respons. HTTP 401 atau 403 tetap berarti HTTPS tersedia, "
                "tetapi akses anonymous ditolak."
            ),
            "NOT_RUN": "Network probe belum dijalankan. Jalankan Analyze Hardening untuk memperoleh evidence.",
            "UNKNOWN": "Evidence belum cukup untuk menentukan perilaku HTTPS secara pasti.",
        }
        if str(result).startswith("HTTP_"):
            code = str(result).split("_", 1)[1]
            if code == "500":
                return (
                    "Endpoint HTTP mengembalikan HTTP 500. Ini merupakan error web-tier/application, "
                    "bukan enforcement HTTPS yang eksplisit berupa redirect atau penolakan koneksi."
                )
            return f"Endpoint HTTP mengembalikan status {code}; hasil perlu ditinjau lebih lanjut."
        return explanations.get(str(result), f"Hasil teknis {result} belum memiliki interpretasi khusus.")
    @staticmethod
    def https_status_summary(status):
        summaries = {
            "PASS": "Mandatory check memenuhi baseline.",
            "PASS WITH SCOPE NOTE": (
                "HTTPS tersedia dan HTTP tidak dapat dijangkau dari lokasi pengujian. "
                "Hasil memenuhi baseline 443-only dengan catatan cakupan jaringan."
            ),
            "FAIL": "Mandatory check tidak memenuhi baseline dan memerlukan perbaikan.",
            "CRITICAL": "Ditemukan perilaku HTTP yang berisiko tinggi dan perlu segera diperbaiki.",
            "UNKNOWN": "Evidence belum cukup untuk memberikan keputusan PASS atau FAIL.",
            "COMPLIANT": "Seluruh mandatory check yang ditemukan memenuhi baseline.",
            "COMPLIANT WITH NOTE": (
                "Mandatory check memenuhi baseline. HTTP tidak terjangkau dari komputer operator, "
                "tetapi exposure dari jalur jaringan lain belum diverifikasi langsung."
            ),
            "NON-COMPLIANT": "Sedikitnya satu mandatory check tidak memenuhi baseline.",
            "PARTIALLY VERIFIED": "Sebagian evidence belum tersedia atau belum dapat dipastikan.",
        }
        return summaries.get(str(status), "Status memerlukan peninjauan.")
    @staticmethod
    def report_escape(value):
        return html.escape(str(value if value is not None else "-"))
    def report_status_html(self, status):
        colors = {
            "PASS": "#19a463", "COMPLIANT": "#19a463",
            "PASS WITH SCOPE NOTE": "#19a463", "PASS WITH NOTE": "#19a463", "COMPLIANT WITH NOTE": "#19a463",
            "FAIL": "#d64040", "NON-COMPLIANT": "#d64040",
            "CRITICAL": "#c62828", "UNKNOWN": "#b7791f",
            "PARTIALLY VERIFIED": "#b7791f",
        }
        color = colors.get(str(status), "#667085")
        return f'<strong style="color:{color};">{self.report_escape(status)}</strong>'
    def report_kv_table(self, rows):
        # QTextBrowser mengikuti HTML subset Qt. Width attribute pada setiap cell
        # lebih konsisten daripada CSS colgroup/table-layout untuk tabel terpisah.
        label_width = 230
        colon_width = 18
        body = []
        for label, value in rows:
            body.append(
                '<tr>'
                f'<td width="{label_width}" style="padding:3px 10px 3px 0;vertical-align:top;">'
                f'<strong>{self.report_escape(label)}</strong></td>'
                f'<td width="{colon_width}" align="center" '
                'style="padding:3px 0;vertical-align:top;">:</td>'
                f'<td style="padding:3px 0 3px 10px;vertical-align:top;">{value}</td>'
                '</tr>'
            )
        return (
            '<table width="100%" cellspacing="0" cellpadding="0" border="0">'
            + ''.join(body)
            + '</table>'
        )

    @staticmethod
    def report_heading(text, level=2):
        sizes = {1: "22px", 2: "17px", 3: "14px"}
        margin = "0 0 12px 0" if level == 1 else "22px 0 9px 0"
        return (
            f'<div style="font-size:{sizes.get(level, "14px")};font-weight:700;'
            f'margin:{margin};">{html.escape(str(text))}</div>'
        )
    def show_readable_report(self, title, content, width=850, height=650, rich=False):
        dialog = QDialog(self)
        dialog.setWindowTitle(title)
        dialog.resize(width, height)
        layout = QVBoxLayout(dialog)
        if rich:
            report = QTextBrowser()
            report.setOpenExternalLinks(False)
            report.setHtml(content)
        else:
            report = QPlainTextEdit()
            report.setLineWrapMode(QPlainTextEdit.LineWrapMode.WidgetWidth)
            report.setPlainText("\n\n".join(content))
        report.setReadOnly(True)
        layout.addWidget(report, 1)
        buttons = QDialogButtonBox(QDialogButtonBox.StandardButton.Close)
        buttons.rejected.connect(dialog.reject)
        buttons.clicked.connect(dialog.accept)
        layout.addWidget(buttons)
        dialog.exec()
    def https_result_detail(self, result, endpoint_status=None):
        result = str(result)
        endpoint_status = str(endpoint_status or "UNKNOWN")
        details = {
            "REDIRECTS_TO_HTTPS": (
                "Request HTTP dialihkan ke HTTPS.",
                "Redirect target menggunakan secure scheme dan endpoint HTTPS dapat diuji.",
                "Evidence hanya mewakili jalur dari komputer operator.",
                "Tidak ada mandatory remediation."
            ),
            "REDIRECTS_TO_NON_HTTPS": (
                "Request HTTP dialihkan ke URL yang masih menggunakan HTTP.",
                "Redirect hanya mengubah path atau format URL dan belum menaikkan koneksi ke HTTPS.",
                "Request awal tetap menggunakan transport tanpa enkripsi.",
                "Ubah redirect target ke https:// atau nonaktifkan exposure HTTP."
            ),
            "PLAINTEXT_CONTENT_AVAILABLE": (
                "Konten masih tersedia langsung melalui HTTP.",
                "Endpoint mengembalikan respons sukses tanpa redirect ke HTTPS.",
                "Request dan response dapat ditransmisikan tanpa enkripsi.",
                "Terapkan redirect ke HTTPS atau nonaktifkan listener HTTP."
            ),
            "NOT_EXPOSED_FROM_TEST_LOCATION": (
                "Endpoint HTTP tidak dapat dijangkau dari komputer operator.",
                "Koneksi HTTP ditolak atau tidak tersedia dari jalur pengujian ini.",
                "Belum membuktikan kondisi dari VLAN, proxy, load balancer, atau lokasi lain.",
                "Tidak ada mandatory remediation jika HTTPS tersedia; validasi jalur jaringan lain bila relevan."
            ),
            "CONNECT_TIMEOUT": (
                "Endpoint HTTP tidak dapat dijangkau dari lokasi pengujian dalam batas waktu koneksi.",
                "Jika HTTPS berstatus REACHABLE, kondisi ini konsisten dengan baseline 443-only; "
                "port 80 dapat tidak memiliki listener atau difilter sebelum request mencapai Web Adaptor.",
                "Tools tidak membaca rule firewall, listener Tomcat/IIS, atau load balancer secara langsung "
                "dan belum membuktikan exposure TCP/80 dari network segment lain.",
                "Tidak ada mandatory remediation dari probe ini. Konfirmasi rule TCP/80 dan listener "
                "jika full infrastructure evidence diperlukan."
            ),
            "READ_TIMEOUT": (
                "Koneksi terbentuk tetapi respons tidak selesai dalam 8 detik.",
                "Web tier menerima koneksi namun tidak menyelesaikan response tepat waktu.",
                "Bisa dipengaruhi performa, proxy, load balancer, atau kondisi sementara.",
                "Periksa log dan performa endpoint lalu jalankan Analyze kembali."
            ),
            "TLS_VERIFICATION_FAILED": (
                "Certificate HTTPS tidak dapat diverifikasi oleh komputer operator.",
                "Trust chain, hostname, masa berlaku, atau CA certificate tidak memenuhi validasi.",
                "Client lain dapat menolak koneksi atau menampilkan certificate warning.",
                "Periksa hostname, expiration, trusted CA, dan kelengkapan certificate chain."
            ),
            "REACHABLE": (
                "Endpoint HTTPS memberikan respons.",
                "Koneksi HTTPS berhasil; HTTP 401 atau 403 tetap berarti endpoint tersedia namun anonymous access ditolak.",
                "Reachability tidak menilai kualitas certificate bila TLS verification dinonaktifkan.",
                "Tidak ada tindakan untuk reachability; gunakan certificate control untuk penilaian trust."
            ),
            "NOT_RUN": (
                "Network probe belum dijalankan.",
                "Test Connection hanya membaca konfigurasi dan metadata.",
                "Belum ada evidence perilaku HTTP/HTTPS.",
                "Jalankan Analyze Hardening."
            ),
            "UNKNOWN": (
                "Evidence belum cukup untuk memberikan keputusan.",
                "Respons tidak sesuai pola yang dapat dinilai secara pasti.",
                "Kondisi aman atau tidak aman belum dapat dikonfirmasi.",
                "Tinjau evidence teknis dan jalankan Analyze kembali setelah perbaikan."
            ),
        }
        if result.startswith("HTTP_"):
            code = result.split("_", 1)[1]
            if code == "500":
                return (
                    "Endpoint HTTP mengembalikan internal server error.",
                    "HTTP 500 bukan enforcement HTTPS yang eksplisit berupa redirect atau penolakan koneksi.",
                    "Perilaku dapat berubah dan penyebabnya mungkin berada pada Web Adaptor, web server, atau proxy.",
                    "Periksa connector, context, virtual host, proxy, dan log; arahkan HTTP ke HTTPS atau nonaktifkan HTTP."
                )
            return (
                f"Endpoint HTTP mengembalikan status {code}.",
                "Respons belum cocok dengan pola redirect HTTPS atau connection refusal.",
                "Evidence belum cukup untuk memastikan enforcement.",
                "Tinjau listener dan routing web tier."
            )
        return details.get(result, details["UNKNOWN"])
    def https_detail_html(self, result, endpoint_status=None):
        meaning, reason, risk, action = self.https_result_detail(result, endpoint_status)
        return self.report_kv_table([
            ("Arti hasil", self.report_escape(meaning)),
            ("Alasan penilaian", self.report_escape(reason)),
            ("Risiko / batas evidence", self.report_escape(risk)),
            ("Tindakan relevan", self.report_escape(action)),
        ])
    def signed_ca_evidence(self):
        portal = self.portal_result or {}; server = self.server_result or {}
        raw = list(portal.get("public_certificates", [])) + list(server.get("public_certificates", []))
        listeners = {}
        for item in raw:
            key = tuple(item.get("listener_key") or (str(item.get("hostname","")).lower(), item.get("port",443)))
            entry = listeners.setdefault(key, dict(item, components=[], registered_urls=[]))
            if item.get("component") not in entry["components"]: entry["components"].append(item.get("component"))
            if item.get("registered_url") not in entry["registered_urls"]: entry["registered_urls"].append(item.get("registered_url"))
        web_tier = list(listeners.values())
        portal_cert = portal.get("portal_certificate", {"status":"UNKNOWN"})
        server_layer = server.get("server_machine_certificates", {"status":"UNKNOWN","machines":[]})
        states = [x.get("status","UNKNOWN") for x in web_tier] + [portal_cert.get("status","UNKNOWN"), server_layer.get("status","UNKNOWN")]
        overall = "NON-COMPLIANT" if "FAIL" in states else "PARTIALLY VERIFIED" if not web_tier or "UNKNOWN" in states or "NOT_RUN" in states else "COMPLIANT WITH NOTE" if "PASS WITH NOTE" in states else "COMPLIANT"
        return {"overall":overall,"web_tier":web_tier,"portal":portal_cert,"server":server_layer}
    def certificate_html(self, cert):
        cert = cert or {}
        return self.report_kv_table([
            ("Alias", self.report_escape(cert.get("alias","-"))), ("Entry type", self.report_escape(cert.get("entry_type","-"))),
            ("Subject", self.report_escape(cert.get("subject","-"))), ("Issuer", self.report_escape(cert.get("issuer","-"))),
            ("SAN", self.report_escape(", ".join(cert.get("sans",[])) or "-")), ("Valid from", self.report_escape(cert.get("valid_from","-"))),
            ("Valid until", self.report_escape(cert.get("valid_until","-"))), ("Days remaining", self.report_escape(cert.get("days_remaining","-"))),
            ("Key algorithm / size", self.report_escape(f"{cert.get('key_algorithm','-')} / {cert.get('key_size','-')}")),
            ("Signature algorithm", self.report_escape(cert.get("signature_algorithm","-"))), ("Key usage", self.report_escape(", ".join(map(str,cert.get("key_usage",[]))) or "-")),
            ("SHA-256 fingerprint", self.report_escape(cert.get("sha256","-"))),
        ])
    def preview_signed_ca_certificates(self):
        ev=self.signed_ca_evidence(); web=ev["web_tier"]; portal=ev["portal"]; machines=ev["server"].get("machines",[])
        web_ok=sum(x.get("status") in ("PASS","PASS WITH NOTE") for x in web); portal_ok=int(portal.get("status") in ("PASS","PASS WITH NOTE")); machine_ok=sum(x.get("status") in ("PASS","PASS WITH NOTE") for x in machines)
        parts=[self.report_heading("IMPLEMENT SIGNED CA CERTIFICATES",1),'<p><strong>Mode:</strong> Analyze, Preview, dan Guidance saja. Tools tidak mengubah certificate atau active alias.</p>',self.report_heading("RINGKASAN HASIL",2),self.report_kv_table([
            ("Overall status",self.report_status_html(ev["overall"])),("Web-tier listeners",f"{web_ok}/{len(web)} meet baseline"),("Portal",f"{portal_ok}/1 meet baseline"),("ArcGIS Server machines",f"{machine_ok}/{len(machines)} meet baseline"),("Tools restart expected","No"),("Manual remediation","Service restart or propagation may be required")])]
        for i,x in enumerate(web,1):
            parts += [self.report_heading(f"WEB-TIER TLS LISTENER {i}     {x.get('status')}",2),self.report_kv_table([("Hostname / port",self.report_escape(f"{x.get('hostname')}:{x.get('port')}")),("Used by",self.report_escape(", ".join(filter(None,x.get("components",[]))))),("Registered URLs",self.report_escape("; ".join(filter(None,x.get("registered_urls",[]))))),("Subject",self.report_escape(x.get("subject_cn","-"))),("Issuer",self.report_escape(x.get("issuer_cn","-"))),("SAN",self.report_escape(", ".join(x.get("sans",[])) or "-")),("Valid until",self.report_escape(x.get("not_after","-"))),("Days remaining",self.report_escape(x.get("days_remaining","-"))),("Trusted chain from operator",self.report_escape(x.get("trusted","-"))),("SHA-256",self.report_escape(x.get("sha256","-"))),("Error",self.report_escape(x.get("error","-")))])]
        parts += [self.report_heading(f"PORTAL ACTIVE CERTIFICATE     {portal.get('status','UNKNOWN')}",2),self.report_kv_table([("Active alias",self.report_escape(portal.get("active_alias","-"))),("Inventory aliases",self.report_escape(", ".join(portal.get("aliases",[])) or "-")),("SSL protocols",self.report_escape(portal.get("protocols","-"))),("HSTS",self.report_escape(portal.get("hsts","-"))),("Reasons",self.report_escape("; ".join(portal.get("reasons",[])) or portal.get("error","-")))]),self.certificate_html(portal.get("active_certificate"))]
        for i, machine in enumerate(portal.get("machines", []), 1):
            parts += [self.report_heading(f"PORTAL MACHINE {i} SAN VERIFICATION     {machine.get('san_status','UNKNOWN')}",3), self.report_kv_table([
                ("Registered machine name", self.report_escape(machine.get("machine_name","-"))),
                ("Native Admin URL", self.report_escape(machine.get("admin_url","-"))),
                ("Platform", self.report_escape(machine.get("platform","-"))),
                ("SAN comparison target", self.report_escape(machine.get("san_target","Unavailable"))),
                ("SAN target source", self.report_escape(machine.get("san_target_source","Unavailable"))),
                ("Target type", self.report_escape(machine.get("san_target_kind","UNAVAILABLE"))),
                ("Certificate SAN", self.report_escape(", ".join((portal.get("active_certificate") or {}).get("sans", [])) or "-")),
                ("SAN verification", self.report_status_html(machine.get("san_status","UNKNOWN"))),
                ("Reasons", self.report_escape("; ".join(machine.get("reasons", [])) or machine.get("error","-"))),
            ])]
        chain=portal.get("chain",{}); parts += [self.report_heading(f"PORTAL CERTIFICATE CHAIN     {chain.get('status','UNKNOWN')}",3),self.report_kv_table([("Intermediate alias",self.report_escape((chain.get("intermediate") or {}).get("alias","Not resolved"))),("Root alias",self.report_escape((chain.get("root") or {}).get("alias","Not resolved")))])]
        for i,m in enumerate(machines,1):
            parts += [self.report_heading(f"ARCGIS SERVER MACHINE {i}     {m.get('status','UNKNOWN')}",2),self.report_kv_table([("Machine",self.report_escape(m.get("machine_name"))),("Admin URL",self.report_escape(m.get("admin_url","-"))),("Native SSL enabled",self.report_escape(m.get("ssl_enabled","-"))),("Active alias",self.report_escape(m.get("active_alias","-"))),("Inventory aliases",self.report_escape(", ".join(m.get("aliases",[])) or "-")),("SAN comparison target",self.report_escape(m.get("san_target","Unavailable"))),("SAN target source",self.report_escape(m.get("san_target_source","Unavailable"))),("Target type",self.report_escape(m.get("san_target_kind","UNAVAILABLE"))),("SAN verification",self.report_status_html(m.get("san_status","UNKNOWN"))),("Reasons",self.report_escape("; ".join(m.get("reasons",[])) or m.get("error","-")))]),self.certificate_html(m.get("active_certificate"))]
            c=m.get("chain",{}); parts += [self.report_heading(f"MACHINE {i} CERTIFICATE CHAIN     {c.get('status','UNKNOWN')}",3),self.report_kv_table([("Intermediate alias",self.report_escape((c.get("intermediate") or {}).get("alias","Not resolved"))),("Root alias",self.report_escape((c.get("root") or {}).get("alias","Not resolved")))])]
        parts += [self.report_heading("APA YANG DIPERIKSA SAAT ANALYZE?",2),'<ol><li>Registered Web Adaptor Portal dan Server, lalu deduplikasi berdasarkan hostname dan effective TLS port.</li><li>Certificate yang disajikan setiap unique web-tier listener.</li><li>Portal Web Server SSL Certificate active alias dan detail certificate melalui Portal Administrator API.</li><li>Seluruh ArcGIS Server machine, native SSL setting, active alias, inventory, dan detail certificate per machine.</li><li>PrivateKeyEntry, issuer/subject, SAN hostname match, validity, expiry 60 hari, RSA key size, signature algorithm, key usage, dan SHA-256 fingerprint.</li><li>Chain metadata melalui hubungan leaf issuer = intermediate subject dan intermediate issuer = root subject. Root self-signed tidak dianggap kegagalan leaf.</li></ol>',self.report_heading("BATAS PEMERIKSAAN",2),'<ul><li>Vanity URL, WAF, load balancer VIP, reverse proxy alias, atau hostname yang tidak registered tidak ditemukan otomatis.</li><li>Web-tier trust mengikuti trust store dan jalur jaringan komputer operator.</li><li>Hubungan chain Portal/Server dinilai dari metadata API, bukan verifikasi signature kriptografis penuh.</li><li>Keberadaan certificate bawaan selfsignedcertificate tidak menjadi finding jika bukan active alias.</li><li>Tools tidak membaca private key, password PFX/P12, permission filesystem, atau proses renewal.</li><li>Tools menggunakan hostname FQDN pada native Admin URL sebagai primary SAN target dan registered machine FQDN sebagai fallback.</li><li>Jika hanya short hostname tersedia, SAN verification menjadi UNKNOWN. Tools tidak menebak DNS suffix atau FQDN yang mungkin berasal dari DNS internal, DNS alias, hosts file, atau nama machine saat instalasi.</li><li>Tools tidak membaca DNS configuration, hosts file, Windows domain membership, atau OS hostname configuration.</li><li>Kegagalan certificate inventory menghasilkan UNKNOWN dan tidak menggagalkan control lain.</li></ul>']
        self.show_readable_report("Preview Signed CA Certificates",''.join(parts),width=980,height=780,rich=True)
    def show_signed_ca_guidance(self):
        ev=self.signed_ca_evidence(); parts=[self.report_heading("SIGNED CA CERTIFICATE GUIDANCE",1)]; findings=[]
        for x in ev["web_tier"]:
            if x.get("status") not in ("PASS",): findings.append(("Web tier",f"{x.get('hostname')}:{x.get('port')}",x.get("status"),x.get("error") or "Certificate does not meet baseline","Perbaiki certificate binding dan complete chain pada IIS, Tomcat, proxy, WAF, atau load balancer; lalu Analyze ulang."))
        p=ev["portal"]
        if p.get("status") not in ("PASS",): findings.append(("Portal",p.get("active_alias","active alias unavailable"),p.get("status"),"; ".join(p.get("reasons",[])) or p.get("error") or "Portal certificate unverified","Import CA-signed leaf beserta chain dan assign sebagai Web Server SSL Certificate pada Portal. Pastikan SAN mencakup hostname atau FQDN native yang telah disepakati dan digunakan pada native Admin URL. Nama tersebut dapat berasal dari hostname machine, DNS internal, DNS alias, hosts file, atau nama machine yang diregistrasikan saat instalasi; lalu Analyze ulang."))
        for m in ev["server"].get("machines",[]):
            if m.get("status") not in ("PASS",): findings.append(("ArcGIS Server",m.get("machine_name"),m.get("status"),"; ".join(m.get("reasons",[])) or m.get("error") or "Machine certificate unverified","Import CA-signed leaf beserta intermediate/root pada machine ini dan assign active alias. Pastikan SAN mencakup hostname atau FQDN native yang telah disepakati dan digunakan pada native Admin URL. Nama tersebut dapat berasal dari hostname machine, DNS internal, DNS alias, hosts file, atau nama machine yang diregistrasikan saat instalasi; lalu Analyze ulang."))
        if not findings: parts.append('<p>Tidak ada mandatory remediation. Seluruh web-tier listener, Portal, dan ArcGIS Server machine memenuhi baseline.</p>')
        for component,target,status,problem,action in findings:
            parts += [self.report_heading(f"{component.upper()}     {status}",2),self.report_kv_table([("Target",self.report_escape(target)),("Masalah",self.report_escape(problem)),("Mengapa penting","Setiap TLS layer harus menggunakan active CA-signed certificate yang valid dan sesuai hostname."),("Tindakan",self.report_escape(action)),("Restart","Manual remediation dapat memerlukan restart atau propagation service."),("Verify","Jalankan Analyze Hardening kembali dan bandingkan active alias serta SHA-256 fingerprint.")])]
        self.show_readable_report("Signed CA Certificate Guidance",''.join(parts),width=980,height=740,rich=True)

    def render_controls(self):
        self.controls_table.setRowCount(0)
        self.last_assessment_at = None
        self.export_assessment_button.setEnabled(False)
        self.control_rows = {}
        self.action_buttons = {}
        self.portal_directory_row = None
        self.automatic_account_row = None
        self.builtin_self_creation_row = None
        self.public_profile_sharing_row = None
        self.social_media_links_row = None
        self.anonymous_access_row = None
        self.portal_servlets_row = None
        self.member_defaults_viewer_row = None
        self.https_enforcement_row = None
        self.jsonp_row = None
        self.services_directory_row = None
        self.standardized_queries_row = None
        self.token_http_get_row = None
        p = self.portal_result
        s = self.server_result

        https_evidence = self.https_enforcement_evidence()
        https_records = https_evidence["records"]
        passed_records = sum(
            1 for item in https_records
            if item["status"] in ("PASS", "PASS WITH SCOPE NOTE")
        )
        https_current = (
            f"Portal {https_evidence['portal_policy']}; "
            f"Server {https_evidence['server_policy']}; "
            f"Web Adaptors {passed_records}/{len(https_records)} meet baseline"
        )
        self.add_control(
            "https_enforcement", "Enterprise", "Verify HTTPS Enforcement",
            https_current, "HTTPS only on all registered entry points",
            https_evidence["overall"], "High"
        )
        cert_evidence = self.signed_ca_evidence()
        web_records = cert_evidence["web_tier"]
        portal_cert = cert_evidence["portal"]
        server_machines = cert_evidence["server"].get("machines", [])
        web_passed = sum(x.get("status") in ("PASS", "PASS WITH NOTE") for x in web_records)
        portal_passed = int(portal_cert.get("status") in ("PASS", "PASS WITH NOTE"))
        server_passed = sum(x.get("status") in ("PASS", "PASS WITH NOTE") for x in server_machines)
        self.add_control(
            "signed_ca_certificates", "Enterprise", "Implement Signed CA Certificates",
            f"Web tier {web_passed}/{len(web_records)}; Portal {portal_passed}/1; Server machines {server_passed}/{len(server_machines)}",
            "CA-signed on web tier, Portal, and every Server machine", cert_evidence["overall"], "High"
        )
        self.add_control(
            "portal_directory", "Portal", "Disable Portal Directory",
            p["portal_directory_disabled"], True,
            self.compliance(p["portal_directory_disabled"], True), "Low"
        )
        self.add_control(
            "automatic_account", "Portal", "Disable Automatic Enterprise Account Creation",
            p["automatic_account_creation"], False,
            self.compliance(p["automatic_account_creation"], False), "Medium"
        )
        defaults = p.get("user_default_settings", {})
        defaults_error = p.get("user_default_settings_error")
        license_items = PortalDirectoryHardeningWorker.collection_items(
            p.get("user_license_types", {}), "userLicenseTypes", "userLicenseType", "results", "items"
        )
        role_items = PortalDirectoryHardeningWorker.collection_items(
            p.get("portal_roles", {}), "roles", "results", "items"
        )
        viewer_license = next((item for item in license_items if isinstance(item, dict) and (
            str(item.get("id", item.get("userLicenseTypeId", ""))).lower() == "viewerut"
            or str(item.get("name", item.get("title", ""))).strip().lower() == "viewer"
        )), None)
        viewer_role = next((item for item in role_items if isinstance(item, dict) and
                            str(item.get("name", "")).strip().lower() == "viewer"), None)
        viewer_license_id = None if not viewer_license else str(
            viewer_license.get("id", viewer_license.get("userLicenseTypeId", ""))
        )
        viewer_role_id = None if not viewer_role else str(viewer_role.get("id", ""))
        current_license = defaults.get("userLicenseType") or "Not set"
        current_role_id = defaults.get("role") or "Not set"
        current_role_name = next((str(item.get("name")) for item in role_items
                                  if isinstance(item, dict) and str(item.get("id")) == str(current_role_id)), None)
        defaults_current = f"User type: {current_license}; Role: {current_role_name or current_role_id}"
        if defaults_error or not viewer_license_id or not viewer_role_id:
            defaults_status = "UNKNOWN"
        elif (str(current_license).lower() == str(viewer_license_id).lower()
              and str(current_role_id) == str(viewer_role_id)):
            defaults_status = "COMPLIANT"
            defaults_current = "User type: Viewer; Role: Viewer"
        else:
            defaults_status = "NON-COMPLIANT"
        self.add_control(
            "member_defaults_viewer", "Portal", "Configure New Member Default Role as Viewer",
            defaults_current, "User type Viewer; Role Viewer", defaults_status, "Low"
        )
        builtin_present = p.get("builtin_self_creation_present", False)
        builtin_raw = p.get("builtin_self_creation_disabled")
        if not builtin_present:
            builtin_current = "Not configured (default disabled)"
            builtin_status = "NON-COMPLIANT"
        elif str(builtin_raw).strip().lower() == "true":
            builtin_current = "Disabled (explicit)"
            builtin_status = "COMPLIANT"
        elif str(builtin_raw).strip().lower() == "false":
            builtin_current = "Enabled (explicit)"
            builtin_status = "NON-COMPLIANT"
        else:
            builtin_current = f"Unknown ({builtin_raw})"
            builtin_status = "UNKNOWN"
        self.add_control(
            "builtin_self_creation", "Portal", "Disable Built-In Account Self-Creation (Portal Restart Expected)",
            builtin_current, "Disabled", builtin_status, "High"
        )
        profile_present = p.get("public_profile_setting_present", False)
        profile_raw = p.get("public_profile_update_disabled")
        if not profile_present:
            profile_current = "Unknown (property unavailable)"
            profile_status = "UNKNOWN"
        elif str(profile_raw).strip().lower() == "true":
            profile_current = "Disabled (explicit)"
            profile_status = "COMPLIANT"
        elif str(profile_raw).strip().lower() == "false":
            profile_current = "Enabled (explicit)"
            profile_status = "NON-COMPLIANT"
        else:
            profile_current = f"Unknown ({profile_raw})"
            profile_status = "UNKNOWN"
        self.add_control(
            "public_profile_sharing", "Portal", "Disable Public User Profile Sharing",
            profile_current, "Disabled", profile_status, "Medium"
        )
        social_present = p.get("social_media_links_present", False)
        social_raw = p.get("social_media_links_raw")
        if not isinstance(p.get("portal_properties"), dict):
            social_current = "Unknown (portalProperties unavailable)"
            social_status = "UNKNOWN"
        elif not social_present:
            social_current = "Not configured (default disabled)"
            social_status = "NON-COMPLIANT"
        elif str(social_raw).strip().lower() == "false":
            social_current = "Disabled (explicit)"
            social_status = "COMPLIANT"
        elif str(social_raw).strip().lower() == "true":
            social_current = "Enabled (explicit)"
            social_status = "NON-COMPLIANT"
        else:
            social_current = f"Unknown ({social_raw})"
            social_status = "UNKNOWN"
        self.add_control(
            "social_media_links", "Portal", "Disable Show Social Media Links",
            social_current, "Disabled", social_status, "Low"
        )
        access_raw = str(p.get("access", "")).strip().lower()
        if access_raw == "private":
            anonymous_current = "Disabled (private)\nBlocked (toggle OFF)"
            anonymous_status = "COMPLIANT"
        elif access_raw == "public":
            anonymous_current = "Enabled (public)\nAllowed (toggle ON)"
            anonymous_status = "NON-COMPLIANT"
        else:
            anonymous_current = f"Unknown ({p.get('access')})"
            anonymous_status = "UNKNOWN"
        self.add_control(
            "anonymous_access", "Portal", "Disable Anonymous Access",
            anonymous_current, "Private", anonymous_status, "High"
        )
        portal_version = str(p.get("version", "")).strip()
        servlet_states = p.get("portal_servlet_properties", {})
        servlet_labels = {
            "disableLegendServlet": "Legend",
            "disablePrintServlet": "Print",
            "disableWFSServlet": "WFS",
        }
        if portal_version.startswith("11.5"):
            compatibility = "SUPPORTED"
            servlet_parts = []
            servlet_all_disabled = True
            servlet_unknown = False
            for name, label in servlet_labels.items():
                entry = servlet_states.get(name, {})
                present = bool(entry.get("present"))
                raw = entry.get("value")
                value = str(raw).strip().lower() if present else ""
                if value == "true":
                    text = "Disabled"
                elif value == "false":
                    text = "Enabled"
                    servlet_all_disabled = False
                elif not present:
                    text = "Not configured"
                    servlet_all_disabled = False
                else:
                    text = f"Unknown ({raw})"
                    servlet_all_disabled = False
                    servlet_unknown = True
                servlet_parts.append(f"{label}: {text}")
            servlet_current = "; ".join(servlet_parts)
            servlet_status = (
                "UNKNOWN" if servlet_unknown
                else "COMPLIANT" if servlet_all_disabled
                else "NON-COMPLIANT"
            )
        elif portal_version.startswith("12.1"):
            compatibility = "NOT_SUPPORTED"
            servlet_current = "Unavailable in ArcGIS Enterprise 12.1"
            servlet_status = "NOT SUPPORTED"
        else:
            compatibility = "UNVALIDATED"
            servlet_current = f"Version {portal_version or 'unknown'} not validated"
            servlet_status = "UNVALIDATED"
        p["portal_servlet_compatibility"] = compatibility
        self.add_control(
            "portal_servlets", "Portal", "Disable Portal Servlets (Selective) (Portal Restart Expected)",
            servlet_current, "Operator selection", servlet_status, "Medium"
        )
        self.add_control(
            "services_directory", "Server", "Disable Services Directory",
            s["services_directory_enabled"], False,
            self.compliance(s["services_directory_enabled"], False), "Low"
        )
        self.add_control(
            "jsonp", "Server", "Disable JSONP Callback Functions",
            s["callback_functions_enabled"], False,
            self.compliance(s["callback_functions_enabled"], False), "Medium"
        )
        token_get_present = s.get("allow_http_get_present", False)
        token_get_raw = s.get("allow_http_get_raw")
        if not token_get_present:
            token_get_current = "Not configured (default disabled)"
            token_get_status = "NON-COMPLIANT"
        elif str(token_get_raw).strip().lower() == "false":
            token_get_current = "Disabled (explicit)"
            token_get_status = "COMPLIANT"
        elif str(token_get_raw).strip().lower() == "true":
            token_get_current = "Enabled (explicit)"
            token_get_status = "NON-COMPLIANT"
        else:
            token_get_current = f"Unknown ({token_get_raw})"
            token_get_status = "UNKNOWN"
        self.add_control(
            "token_http_get", "Server", "Disable Token Acquisition via HTTP GET",
            token_get_current, "Disabled", token_get_status, "High"
        )
        standardized_present = s.get("standardized_queries_present", False)
        standardized_raw = s.get("standardized_queries_raw")
        if not standardized_present:
            standardized_current = "Not configured\n(default enabled)"
            standardized_status = "NON-COMPLIANT"
        elif str(standardized_raw).strip().lower() == "true":
            standardized_current = "Enabled (explicit)"
            standardized_status = "COMPLIANT"
        elif str(standardized_raw).strip().lower() == "false":
            standardized_current = "Disabled (explicit)"
            standardized_status = "NON-COMPLIANT"
        else:
            standardized_current = f"Unknown ({standardized_raw})"
            standardized_status = "UNKNOWN"
        self.add_control(
            "standardized_queries", "Server", "Enable Standardized Queries",
            standardized_current, "Enabled", standardized_status, "High"
        )
        xss_present = s.get("feature_service_xss_present", False)
        xss_raw = s.get("feature_service_xss_raw")
        if not xss_present:
            xss_current = "Not configured (default input)"
            xss_status = "NON-COMPLIANT"
        elif str(xss_raw).strip().lower() == "input":
            xss_current = "Input scanning (Basic)"
            xss_status = "COMPLIANT"
        elif str(xss_raw).strip().lower() == "inputoutput":
            xss_current = "Input + output scanning"
            xss_status = "COMPLIANT"
        else:
            xss_current = f"Unknown ({xss_raw})"
            xss_status = "UNKNOWN"
        self.add_control(
            "feature_service_xss", "Server", "Verify Feature Service XSS Filter Default",
            xss_current, "Basic / Advanced", xss_status, "High"
        )

        origins = s["allowed_origins"]
        origin_status = "NEEDS CONFIGURATION" if str(origins).strip() == "*" else "REVIEW"
        self.add_control(
            "allowed_origins", "Server", "Restrict Allowed Origins",
            origins, "Not configured", origin_status, "Medium"
        )
        self.refresh_action_buttons()
        self.refresh_semantic_colors()
        QTimer.singleShot(0, self.update_responsive_table_columns)

    def update_responsive_table_columns(self):
        """Allocate table width responsively without sacrificing critical columns."""
        if not hasattr(self, "controls_table") or self.controls_table is None:
            return

        # Stable columns: protect labels, status, risk, and three action buttons.
        fixed_widths = {
            0: 96,   # Component
            4: 154,  # Status
            5: 80,   # Risk
            6: 268,  # Actions
        }
        for column, width in fixed_widths.items():
            self.controls_table.setColumnWidth(column, width)

        viewport_width = self.controls_table.viewport().width()
        fixed_total = sum(fixed_widths.values())
        # Include grid lines and a small safety margin for the vertical scrollbar.
        flexible_available = max(0, viewport_width - fixed_total - 16)

        minimums = {1: 250, 2: 185, 3: 125}
        minimum_total = sum(minimums.values())
        if flexible_available <= minimum_total:
            # Keep readable minima; horizontal scrollbar becomes the fallback.
            for column, width in minimums.items():
                self.controls_table.setColumnWidth(column, width)
            return

        # Share remaining space: Control 55%, Current 28%, Target 17%.
        control_width = max(minimums[1], int(flexible_available * 0.55))
        current_width = max(minimums[2], int(flexible_available * 0.28))
        target_width = max(
            minimums[3], flexible_available - control_width - current_width
        )
        self.controls_table.setColumnWidth(1, control_width)
        self.controls_table.setColumnWidth(2, current_width)
        self.controls_table.setColumnWidth(3, target_width)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        if hasattr(self, "table_resize_timer"):
            self.table_resize_timer.start()

    def refresh_action_buttons(self):
        busy = self.hardening_thread is not None or self.thread is not None
        for control_id, buttons in self.action_buttons.items():
            row = self.control_rows.get(control_id)
            status = ""
            if row is not None and self.controls_table.item(row, 4):
                status_item = self.controls_table.item(row, 4)
                status = status_item.data(Qt.ItemDataRole.UserRole) or (
                    status_item.text().replace("\n", " ")
                )

            preview_enabled = not busy
            apply_enabled = (
                not busy
                and self.control_supported(control_id)
                and status == "NON-COMPLIANT"
            )
            if control_id == "feature_service_xss":
                xss_present = self.server_result.get("feature_service_xss_present", False)
                xss_raw = str(
                    self.server_result.get("feature_service_xss_raw", "")
                ).strip().lower()
                # Missing: Basic/Advanced available. Basic: upgrade to Advanced available.
                # Advanced or unknown: no further Apply action.
                apply_enabled = (
                    not busy
                    and (not xss_present or xss_raw == "input")
                )
            if control_id == "portal_servlets":
                compatibility = (self.portal_result or {}).get(
                    "portal_servlet_compatibility", "UNVALIDATED"
                )
                apply_enabled = (
                    not busy
                    and compatibility == "SUPPORTED"
                    and status in ("NON-COMPLIANT", "COMPLIANT")
                )
            rollback_enabled = (
                not busy and bool(self.backup_for_control(control_id))
            )
            if control_id == "portal_servlets":
                rollback_enabled = rollback_enabled and (
                    (self.portal_result or {}).get("portal_servlet_compatibility") == "SUPPORTED"
                )

            if control_id in ("https_enforcement", "signed_ca_certificates"):
                apply_enabled = not busy
                rollback_enabled = False
                buttons["apply"].setText("Guidance")
                buttons["apply"].setToolTip(
                    "Tampilkan guidance untuk web tier, Portal, dan setiap ArcGIS Server machine."
                    if control_id == "signed_ca_certificates"
                    else "Tampilkan manual remediation guidance sesuai IIS/Tomcat."
                )
                buttons["rollback"].setToolTip("Verify-only control tidak memiliki rollback otomatis.")
            buttons["preview"].setEnabled(preview_enabled)
            buttons["apply"].setEnabled(apply_enabled)
            buttons["rollback"].setEnabled(rollback_enabled)

            if status == "COMPLIANT":
                if control_id == "feature_service_xss":
                    buttons["apply"].setToolTip(
                        "Pilih profil Basic atau upgrade ke Advanced (inputOutput)."
                    )
                else:
                    buttons["apply"].setToolTip("Kontrol sudah compliant.")
            elif not self.control_supported(control_id):
                buttons["apply"].setToolTip(
                    "Apply untuk kontrol ini belum tersedia atau memerlukan konfigurasi."
                )
            elif status != "NON-COMPLIANT":
                buttons["apply"].setToolTip(
                    f"Apply tidak tersedia saat status {status or 'belum dianalisis'}."
                )
            else:
                if control_id == "feature_service_xss":
                    buttons["apply"].setToolTip("Pilih profil XSS Basic atau Advanced.")
                else:
                    buttons["apply"].setToolTip("Terapkan hardening untuk kontrol ini.")

            if rollback_enabled:
                buttons["rollback"].setToolTip(
                    f"Rollback menggunakan backup: {self.backup_for_control(control_id)}"
                )
            else:
                buttons["rollback"].setToolTip(
                    "Belum ada backup yang tersedia untuk kontrol ini."
                )

    def preview_control(self, control):
        if not self.portal_result or not self.server_result:
            QMessageBox.warning(
                self, "Belum dianalisis", "Jalankan Analyze Hardening terlebih dahulu."
            )
            return

        if control == "signed_ca_certificates":
            self.preview_signed_ca_certificates()
            self.append_log("Preview Implement Signed CA Certificates ditampilkan.")
            return
        if control == "https_enforcement":
            evidence = self.https_enforcement_evidence()
            records = evidence["records"]
            passed = sum(
                1 for item in records
                if item.get("status") in ("PASS", "PASS WITH SCOPE NOTE")
            )
            parts = [self.report_heading("VERIFY HTTPS ENFORCEMENT", 1)]
            parts.append('<p><strong>Mode:</strong> Verify-only. Tools tidak mengubah konfigurasi apa pun.</p>')
            parts.append(self.report_heading("RINGKASAN HASIL", 2))
            parts.append(self.report_kv_table([
                ("Overall status", self.report_status_html(evidence["overall"])),
                ("Arti status", self.report_escape(self.https_status_summary(evidence["overall"]))),
                ("Portal policy", f'{self.report_status_html(evidence["portal_policy"])} &nbsp; (allSSL={self.report_escape(evidence["portal_all_ssl"])})'),
                ("Server policy", f'{self.report_status_html(evidence["server_policy"])} &nbsp; ({self.report_escape(evidence["server_protocol"])})'),
                ("Server evidence", self.report_escape(evidence["server_protocol_source"])),
                ("Registered endpoints", f'{passed} dari {len(records)} memenuhi baseline'),
                ("Server HSTS setting", f'{self.report_escape(evidence["server_hsts_enabled"])} &nbsp; (optional strengthening)'),
            ]))
            if evidence.get("portal_inventory_error"):
                parts.append(self.report_heading("PORTAL WEB ADAPTOR INVENTORY — UNKNOWN", 2))
                parts.append(f'<p>{self.report_escape(evidence["portal_inventory_error"])}</p>')
            if evidence.get("server_inventory_error"):
                parts.append(self.report_heading("SERVER WEB ADAPTOR INVENTORY — UNKNOWN", 2))
                parts.append(f'<p>{self.report_escape(evidence["server_inventory_error"])}</p>')
            if not records:
                parts.append(self.report_heading("REGISTERED WEB ADAPTORS — UNKNOWN", 2))
                parts.append('<p>Tidak ada registered Web Adaptor yang berhasil ditemukan.</p>')
            for item in records:
                probe = item.get("probe", {})
                http_result = str(probe.get("http_result", "UNKNOWN"))
                https_result = str(probe.get("https_result", "UNKNOWN"))
                parts.append(self.report_heading(
                    f"{item['component'].upper()} WEB ADAPTOR {item['number']} — {item.get('status')}", 2
                ))
                parts.append(self.report_kv_table([
                    ("Arti status", self.report_escape(self.https_status_summary(item.get("status")))),
                    ("URL terdaftar", self.report_escape(item.get("url") or "Tidak tersedia")),
                    ("Platform", self.report_escape(item.get("web_server"))),
                    ("Operating system", self.report_escape(item.get("os"))),
                    ("HTTP / HTTPS port", f'{self.report_escape(item.get("http_port"))} / {self.report_escape(item.get("https_port"))}'),
                ]))
                parts.append(self.report_heading("HTTP TEST", 3))
                parts.append(self.report_kv_table([
                    ("Technical result", f'<code>{self.report_escape(http_result)}</code>'),
                    ("HTTP status", self.report_escape(probe.get("http_status") if probe.get("http_status") is not None else "-")),
                    ("Redirect target", self.report_escape(probe.get("http_location") or "-")),
                    ("Durasi", f'{float(probe.get("http_elapsed", 0)):.2f}s'),
                ]))
                parts.append(self.https_detail_html(http_result, item.get("status")))
                parts.append(self.report_heading("HTTPS TEST", 3))
                parts.append(self.report_kv_table([
                    ("Technical result", f'<code>{self.report_escape(https_result)}</code>'),
                    ("HTTPS status", self.report_escape(probe.get("https_status") if probe.get("https_status") is not None else "-")),
                    ("Durasi", f'{float(probe.get("https_elapsed", 0)):.2f}s'),
                    ("HSTS observed", "Yes" if probe.get("hsts") else "No"),
                ]))
                parts.append(self.https_detail_html(https_result, item.get("status")))
            parts.append(self.report_heading("APA YANG DIPERIKSA SAAT ANALYZE?", 2))
            checks = [
                ("Kebijakan HTTPS Portal", "Memeriksa nilai <code>allSSL</code> pada konfigurasi organisasi Portal.", "Target: <code>allSSL = true</code>."),
                ("Kebijakan HTTPS ArcGIS Server", "Memeriksa <code>protocol</code> atau kombinasi <code>httpEnabled</code> dan <code>sslEnabled</code>.", "Target: HTTP dinonaktifkan dan HTTPS diaktifkan."),
                ("Web Adaptor yang terdaftar", "Menginventarisasi seluruh Web Adaptor Portal dan ArcGIS Server.", "Target: seluruh registered entry point berhasil ditemukan."),
                ("Platform Web Adaptor", "Membaca IIS atau Tomcat, operating system, version, registered URL, serta port HTTP dan HTTPS.", "Target: metadata cukup untuk menentukan cakupan pengujian."),
                ("Perilaku endpoint HTTP", "Menguji apakah URL HTTP dialihkan ke HTTPS, masih menyajikan konten, dialihkan ke HTTP lain, tidak terjangkau, timeout, atau menghasilkan error.", "Target: redirect ke HTTPS atau HTTP tidak diekspos."),
                ("Ketersediaan endpoint HTTPS", "Memastikan endpoint HTTPS memberikan respons. HTTP 401 atau 403 tetap menunjukkan HTTPS tersedia, tetapi anonymous access ditolak.", "Target: endpoint HTTPS berstatus REACHABLE."),
                ("Header HSTS", "Memeriksa setting <code>HSTSEnabled</code> dan header <code>Strict-Transport-Security</code> pada response HTTPS.", "Target: optional strengthening setelah dependency review."),
                ("Penilaian gabungan", "Menggabungkan seluruh mandatory evidence menjadi overall status.", "Output: COMPLIANT, COMPLIANT WITH NOTE, NON-COMPLIANT, CRITICAL, atau PARTIALLY VERIFIED."),
            ]
            parts.append('<ol style="margin-top:4px;padding-left:24px;">' + ''.join(
                f'<li style="margin:0 0 12px 0;"><strong>{title}</strong><br>{desc}<br><span style="color:#667085;">{target}</span></li>'
                for title, desc, target in checks
            ) + '</ol>')
            parts.append(self.report_heading("BATAS PEMERIKSAAN", 2))
            limits = [
                "Pengujian jaringan hanya mewakili jalur dari komputer operator.",
                "Tidak ada username, password, token, atau login cookie yang dikirim dalam probe HTTP/HTTPS.",
                "Tools tidak membaca atau mengubah server.xml, web.xml, IIS, proxy, load balancer, WAF, firewall, atau certificate store secara langsung.",
                "Hasil hanya mencakup Web Adaptor yang terdaftar pada ArcGIS Enterprise.",
                "Public alias, reverse proxy, atau entry point yang tidak terdaftar perlu diperiksa terpisah.",
                "Guidance menggunakan evidence Analyze terakhir dan tidak menjalankan network probe ulang.",
            ]
            parts.append('<ul style="padding-left:22px;">' + ''.join(
                f'<li style="margin-bottom:7px;">{self.report_escape(value)}</li>' for value in limits
            ) + '</ul>')
            self.show_readable_report(
                "Preview HTTPS Enforcement", ''.join(parts), width=960, height=760, rich=True
            )
            self.append_log("Preview Verify HTTPS Enforcement ditampilkan.")
            return
        if control == "portal_servlets":
            version = str(self.portal_result.get("version", "Tidak tersedia"))
            compatibility = self.portal_result.get("portal_servlet_compatibility", "UNVALIDATED")
            states = self.portal_result.get("portal_servlet_properties", {})
            labels = {
                "disableLegendServlet": "Legend servlet",
                "disablePrintServlet": "Print servlet",
                "disableWFSServlet": "WFS servlet",
            }
            state_lines = []
            for name, label in labels.items():
                entry = states.get(name, {})
                if not entry.get("present"):
                    text = "Not configured"
                elif str(entry.get("value")).strip().lower() == "true":
                    text = "Disabled (explicit true)"
                elif str(entry.get("value")).strip().lower() == "false":
                    text = "Enabled (explicit false)"
                else:
                    text = f"Unknown ({entry.get('value')})"
                state_lines.append(f"- {label}: {text}")
            support_note = (
                "Apply tersedia dan selective pada Portal 11.5."
                if compatibility == "SUPPORTED"
                else "Apply diblokir. Portal 12.1 menolak property servlet."
                if compatibility == "NOT_SUPPORTED"
                else "Apply diblokir sampai versi ini divalidasi."
            )
            QMessageBox.information(
                self, "Preview Portal Servlet Hardening",
                "Component: Portal for ArcGIS\n"
                f"Portal version: {version}\nCompatibility: {compatibility}\n\n"
                "Current state:\n" + "\n".join(state_lines) + "\n\n"
                "Selectable properties:\n"
                "- disableLegendServlet\n- disablePrintServlet\n- disableWFSServlet\n\n"
                f"{support_note}\n\n"
                "Safety workflow:\n"
                "1. Operator memilih servlet secara individual.\n"
                "2. Apply melakukan GET latest System Properties.\n"
                "3. Backup hanya state property yang perlu diubah.\n"
                "4. Merge pilihan operator tanpa mengubah property lain.\n"
                "5. Tunggu restart/propagation, autentikasi ulang, dan live verify.\n"
                "6. Rollback merge previous state ke konfigurasi live terbaru.\n\n"
                "Belum ada konfigurasi yang diubah pada tahap preview."
            )
            self.append_log("Preview Portal Servlet Hardening ditampilkan.")
            return
        if control == "builtin_self_creation":
            present = self.portal_result.get("builtin_self_creation_present", False)
            raw = self.portal_result.get("builtin_self_creation_disabled")
            if not present:
                current = "Not configured (effective default is disabled)"
            elif str(raw).strip().lower() == "true":
                current = "Disabled (explicit true)"
            elif str(raw).strip().lower() == "false":
                current = "Enabled (explicit false)"
            else:
                current = f"Unknown ({raw})"
            QMessageBox.information(
                self, "Preview Built-In Account Self-Creation",
                "Component: Portal for ArcGIS\n"
                "Control: Disable Built-In Account Self-Creation\n"
                "Technical property: disableSignup\n\n"
                f"Current: {current}\nTarget : disableSignup = true\n\n"
                "Account type:\nPortal built-in identity store.\n\n"
                "How it works:\nJika aktif, halaman Sign In menyediakan Create an account. "
                "Pengguna dapat membuat username dan password yang disimpan serta dikelola Portal.\n\n"
                "Impact setelah hardening:\n"
                "- Create an account tidak tersedia.\n"
                "- Administrator tetap dapat membuat atau mengundang member.\n"
                "- Existing built-in accounts tidak dihapus atau dinonaktifkan.\n"
                "- Enterprise login dan Identity Provider tidak diubah.\n"
                "- Portal dapat restart/propagasi; utility mengikuti recheckAfterSeconds.\n\n"
                "Perbedaan dengan Automatic Enterprise Account Creation:\n"
                "- Enterprise: identitas sudah ada di AD/LDAP/SAML/OIDC; Portal hanya mendaftarkannya.\n"
                "- Built-in: pengguna membuat identitas baru langsung di Portal.\n\n"
                "Safety workflow:\n"
                "1. GET dan backup seluruh Portal System Properties.\n"
                "2. Merge hanya disableSignup=true.\n"
                "3. POST seluruh object agar property lain tetap ada.\n"
                "4. Tunggu Portal, autentikasi ulang, dan live verify.\n"
                "5. Rollback mengembalikan exact backup.\n\n"
                "Belum ada konfigurasi yang diubah pada tahap preview."
            )
            self.append_log("Preview Built-In Account Self-Creation ditampilkan.")
            return
        if control == "anonymous_access":
            access_raw = str(self.portal_result.get("access", "")).strip().lower()
            if access_raw == "private":
                current = "Disabled (private)"
            elif access_raw == "public":
                current = "Enabled (public)"
            else:
                current = f"Unknown ({self.portal_result.get('access')})"
            QMessageBox.information(
                self, "Preview Disable Anonymous Access",
                "Component: Portal for ArcGIS\n"
                "Control: Disable Anonymous Access\n"
                "Technical property: access\n\n"
                f"Current: {current}\nTarget : access = private\n\n"
                "Potential impact:\n"
                "- Pengguna tanpa login dapat kehilangan akses ke website Portal.\n"
                "- Public web maps, dashboards, dan web applications yang mengandalkan "
                "anonymous access dapat terdampak.\n"
                "- Sharing level item dan permission underlying service tidak otomatis diubah.\n"
                "- Web-tier authentication dapat memerlukan konfigurasi IIS/web server terpisah.\n\n"
                "Safety workflow:\n"
                "1. Backup access dan state canShareBingPublic.\n"
                "2. Partial update access=private dan canShareBingPublic=false.\n"
                "3. Authenticated live verification access=private.\n"
                "4. Anonymous evidence probe tanpa token.\n"
                "5. Rollback mengembalikan access sebelumnya sesuai payload Portal 11.5.\n\n"
                "Belum ada konfigurasi yang diubah pada tahap preview."
            )
            self.append_log("Preview Disable Anonymous Access ditampilkan.")
            return
        if control == "social_media_links":
            portal_properties = self.portal_result.get("portal_properties")
            present = self.portal_result.get("social_media_links_present", False)
            raw = self.portal_result.get("social_media_links_raw")
            if not isinstance(portal_properties, dict):
                current = "Unknown (portalProperties unavailable)"
            elif not present:
                current = "Not configured (effective default is disabled)"
            elif str(raw).strip().lower() == "false":
                current = "Disabled (explicit false)"
            elif str(raw).strip().lower() == "true":
                current = "Enabled (explicit true)"
            else:
                current = f"Unknown ({raw})"
            QMessageBox.information(
                self, "Preview Show Social Media Links",
                "Component: Portal for ArcGIS\n"
                "Control: Disable Show Social Media Links\n"
                "Technical property: portalProperties.showSocialMediaLinks\n\n"
                f"Current: {current}\nTarget : showSocialMediaLinks = false\n\n"
                "Effect:\n"
                "Link social media tidak ditampilkan pada halaman item dan group.\n\n"
                "Not affected:\n"
                "- Item dan group tidak dihapus.\n"
                "- Sharing level public/private tidak berubah.\n"
                "- Direct URL tetap berfungsi.\n"
                "- Login, role, dan privilege tidak berubah.\n"
                "- Experience Builder Share widget atau aplikasi custom dapat memiliki "
                "konfigurasi sharing sendiri.\n\n"
                "Safety workflow:\n"
                "1. GET seluruh portalProperties terbaru.\n"
                "2. Backup object portalProperties lengkap.\n"
                "3. Merge hanya showSocialMediaLinks=false.\n"
                "4. POST seluruh portalProperties agar setting lain tetap ada.\n"
                "5. Live verification tanpa cache.\n"
                "6. Rollback mengembalikan exact portalProperties backup.\n\n"
                "Belum ada konfigurasi yang diubah pada tahap preview."
            )
            self.append_log("Preview Show Social Media Links ditampilkan.")
            return
        if control == "public_profile_sharing":
            present = self.portal_result.get("public_profile_setting_present", False)
            raw = self.portal_result.get("public_profile_update_disabled")
            if not present:
                current = "Unknown (property unavailable)"
            elif str(raw).strip().lower() == "true":
                current = "Disabled (explicit true)"
            elif str(raw).strip().lower() == "false":
                current = "Enabled (explicit false)"
            else:
                current = f"Unknown ({raw})"
            QMessageBox.information(
                self, "Preview Public User Profile Sharing",
                "Component: Portal for ArcGIS\n"
                "Control: Disable Public User Profile Sharing\n"
                "Technical property: updateUserProfileDisabled\n\n"
                f"Current: {current}\nTarget : updateUserProfileDisabled = true\n\n"
                "Effect:\n"
                "Member tidak dapat mengubah informasi biografi dan menentukan visibility "
                "profil secara mandiri.\n\n"
                "Not affected:\n"
                "- Existing user tidak dihapus atau dinonaktifkan.\n"
                "- Login, role, privilege, dan Identity Provider tidak berubah.\n"
                "- Item, group, dan content sharing tidak otomatis berubah.\n"
                "- Data profil yang sudah ada tidak otomatis dihapus.\n\n"
                "Security benefit:\n"
                "Mengurangi paparan informasi personal, reconnaissance, dan social engineering.\n\n"
                "Safety workflow:\n"
                "1. Baca Portal Self terbaru dan backup previous state.\n"
                "2. Partial update updateUserProfileDisabled=true.\n"
                "3. Live verification melalui Portal Self tanpa cache.\n"
                "4. Rollback mengirim previous state dari backup.\n\n"
                "Belum ada konfigurasi yang diubah pada tahap preview."
            )
            self.append_log("Preview Public User Profile Sharing ditampilkan.")
            return
        if control == "member_defaults_viewer":
            defaults = self.portal_result.get("user_default_settings", {})
            license_value = defaults.get("userLicenseType") or "Not set"
            role_value = defaults.get("role") or "Not set"
            QMessageBox.information(
                self, "Preview New Member Default Role as Viewer",
                "Component: Portal for ArcGIS\n"
                "Control: Configure New Member Default Role as Viewer\n"
                "Restart expected: No\n"
                "Scope: New members only\n\n"
                f"Current user type: {license_value}\nCurrent role ID: {role_value}\n\n"
                "Target user type: Viewer (viewerUT)\nTarget role: Viewer (resolved live from Portal roles)\n\n"
                "Security benefit:\nNew members receive least-privilege Viewer defaults.\n\n"
                "Not affected:\n- Existing members, roles, licenses, groups, and content are not changed.\n"
                "- Administrators can still assign a different role when required.\n\n"
                "Safety workflow:\n1. GET latest userDefaultSettings.\n"
                "2. Resolve Viewer IDs from live userLicenseTypes and roles catalogs.\n"
                "3. Create scoped backup for userLicenseType and role.\n"
                "4. POST viewerUT and the resolved Viewer role ID.\n"
                "5. GET live settings and verify exact target.\n"
                "6. Rollback restores a previously configured pair; clearing Not set remains fail-safe blocked.\n\n"
                "Belum ada konfigurasi yang diubah pada tahap preview."
            )
            self.append_log("Preview New Member Default Role as Viewer ditampilkan.")
            return
        if control == "automatic_account":
            current = self.portal_result.get("automatic_account_creation")
            title = "Preview Automatic Enterprise Account Creation Hardening"
            control_name = "Disable Automatic Enterprise Account Creation"
            property_name = "enableAutomaticAccountCreation"
            target_text = "false"
            component_name = "Portal for ArcGIS"
            backup_scope = "Portal security configuration"
            impact = (
                "Berlaku untuk organization-specific identity dari AD, LDAP, SAML, OIDC, "
                "portal-tier, atau web-tier authentication. Pengguna sudah memiliki identitas "
                "di Identity Provider; Portal hanya mendaftarkan identitas tersebut sebagai "
                "member saat akses pertama. Setelah dinonaktifkan, administrator harus "
                "menambahkan member terlebih dahulu. Existing members tidak dihapus, integrasi "
                "Identity Provider tidak dimatikan, dan tidak ada password built-in baru yang dibuat. "
                "Ini berbeda dari Built-In Account Self-Creation, yaitu pengguna membuat username "
                "dan password langsung di Portal melalui Create an account."
            )
        elif control == "portal_directory":
            current = self.portal_result.get("portal_directory_disabled")
            title = "Preview Portal Directory Hardening"
            control_name = "Disable Portal Directory"
            property_name = "disableServicesDirectory"
            target_text = "true"
            component_name = "Portal for ArcGIS"
            backup_scope = "Portal security configuration"
            impact = (
                "Browsable HTML Portal Directory dinonaktifkan. JSON/REST access tetap "
                "tersedia dan permission item tetap berlaku."
            )
        elif control == "jsonp":
            current = self.server_result.get("callback_functions_enabled")
            title = "Preview JSONP Hardening"
            control_name = "Disable JSONP Callback Functions"
            property_name = "callbackFunctionsEnabled"
            target_text = "false"
            component_name = "ArcGIS Server"
            backup_scope = "Services Directory configuration"
            impact = "Legacy application yang masih menggunakan JSONP callback dapat berhenti bekerja."
        elif control == "services_directory":
            current = self.server_result.get("services_directory_enabled")
            title = "Preview Services Directory Hardening"
            control_name = "Disable ArcGIS Server Services Directory"
            property_name = "servicesDirEnabled"
            target_text = "false"
            component_name = "ArcGIS Server"
            backup_scope = "Services Directory configuration"
            impact = (
                "HTML Services Directory tidak dapat dibrowse, tetapi ArcGIS REST API dan "
                "services tetap tersedia sesuai permission yang berlaku."
            )
        elif control == "token_http_get":
            present = self.server_result.get("allow_http_get_present", False)
            raw = self.server_result.get("allow_http_get_raw")
            if not present:
                current = "Not configured (effective default is disabled)"
            elif str(raw).strip().lower() == "false":
                current = "Disabled (explicit false)"
            elif str(raw).strip().lower() == "true":
                current = "Enabled (explicit true)"
            else:
                current = f"Unknown ({raw})"
            QMessageBox.information(
                self, "Preview Token Acquisition via HTTP GET",
                "Component: ArcGIS Server\n"
                "Control: Disable Token Acquisition via HTTP GET\n"
                "Technical property: properties.allowHttpGet\n\n"
                f"Current: {current}\nTarget : allowHttpGet = false\n\n"
                "Security risk:\n"
                "HTTP GET menempatkan username dan password pada URL. URL dapat tersimpan "
                "pada browser history, reverse proxy, network device, dan access log.\n\n"
                "Impact setelah hardening:\n"
                "- Token tidak dapat diperoleh melalui HTTP GET.\n"
                "- Token melalui HTTP POST dengan credential di request body tetap berfungsi.\n"
                "- Aplikasi legacy yang masih memakai GET harus diperbarui.\n\n"
                "Safety workflow:\n"
                "1. GET seluruh Token Manager Configuration.\n"
                "2. Backup hanya state allowHttpGet; sharedKey tidak ditulis ke disk/log.\n"
                "3. Merge allowHttpGet=false ke konfigurasi live lengkap.\n"
                "4. Live verify konfigurasi.\n"
                "5. Positive test POST memastikan token service tetap berfungsi.\n"
                "6. Negative GET dengan credential asli sengaja tidak dikirim agar password tidak masuk URL/log.\n"
                "7. Rollback membaca konfigurasi live lalu merge nilai sebelumnya.\n\n"
                "Belum ada konfigurasi yang diubah pada tahap preview."
            )
            self.append_log("Preview Token Acquisition via HTTP GET ditampilkan.")
            return
        elif control == "feature_service_xss":
            present = self.server_result.get("feature_service_xss_present", False)
            raw = self.server_result.get("feature_service_xss_raw")
            if not present:
                current = "Not configured (effective default input)"
            elif str(raw).strip().lower() == "input":
                current = "Input scanning (Basic)"
            elif str(raw).strip().lower() == "inputoutput":
                current = "Input + output scanning (accepted stronger value)"
            else:
                current = f"Unknown ({raw})"
            QMessageBox.information(
                self, "Preview Feature Service XSS Filter Default",
                "Component: ArcGIS Server\n"
                "Control: Verify Feature Service XSS Filter Default\n"
                "Technical property: featureServiceXSSFilter\n\n"
                f"Current: {current}\nTarget : featureServiceXSSFilter = input\n\n"
                "Profiles:\n"
                "- Klik Apply lalu pilih Basic untuk featureServiceXSSFilter=input.\n"
                "- Klik Apply lalu pilih Advanced untuk featureServiceXSSFilter=inputOutput.\n"
                "- Keduanya dinilai COMPLIANT; Basic tidak pernah menurunkan inputOutput.\n\n"
                "Scope:\n"
                "Property ini menetapkan default untuk Feature Service baru. Existing Feature "
                "Services tidak otomatis berubah dan perlu diaudit terpisah melalui "
                "XSSPreventionEnabled.\n\n"
                "Potential impact:\n"
                "Input yang mengandung script atau markup mencurigakan dapat ditolak. Lakukan "
                "regression test pada editing, atribut HTML, popup, Field Maps, Survey123, "
                "Experience Builder, dan aplikasi custom.\n\n"
                "Control isolation:\n"
                "Apply dan Rollback hanya mengubah featureServiceXSSFilter. standardizedQueries "
                "dan System Properties lain dibaca ulang dan dipertahankan.\n\n"
                "Safety workflow:\n"
                "1. GET System Properties terbaru.\n"
                "2. Backup hanya previous state featureServiceXSSFilter.\n"
                "3. Merge featureServiceXSSFilter=input tanpa mengubah property lain.\n"
                "4. Live verification.\n"
                "5. Rollback merge previous state ke konfigurasi live terbaru.\n\n"
                "Belum ada konfigurasi yang diubah pada tahap preview."
            )
            self.append_log("Preview Feature Service XSS Filter Default ditampilkan.")
            return
        elif control == "standardized_queries":
            present = self.server_result.get("standardized_queries_present", False)
            raw = self.server_result.get("standardized_queries_raw")
            if not present:
                current = "Not configured (ArcGIS effective default is enabled)"
            elif str(raw).strip().lower() == "true":
                current = "Enabled (explicit true)"
            elif str(raw).strip().lower() == "false":
                current = "Disabled (explicit false)"
            else:
                current = f"Unknown ({raw})"
            title = "Preview Standardized Queries Hardening"
            control_name = "Enable Standardized Queries"
            property_name = "standardizedQueries"
            target_text = '"true"'
            component_name = "ArcGIS Server site"
            backup_scope = "seluruh Server System Properties"
            impact = (
                "Berlaku untuk seluruh ArcGIS Server site. Walaupun ArcGIS memperlakukan "
                "property yang kosong sebagai enabled secara default, baseline hardening ini "
                "mensyaratkan standardizedQueries=\"true\" tertulis eksplisit agar mudah "
                "diaudit dan tidak ambigu. Query dengan fungsi atau sintaks khusus DBMS, "
                "subquery tertentu, join antar-workspace, atau sumber OLE DB dapat ditolak. "
                "Update dilakukan dengan GET + merge + POST agar property lain tidak terhapus."
            )
        else:
            QMessageBox.information(
                self,
                "Allowed Origins perlu dikonfigurasi",
                "Restrict Allowed Origins belum dapat dipreview sebagai perubahan siap-Apply. "
                "Trusted origins harus dikonfigurasi terlebih dahulu."
            )
            return

        QMessageBox.information(
            self, title,
            f"Component: {component_name}\nControl: {control_name}\n\n"
            f"Current: {property_name} = {current}\n"
            f"Target : {property_name} = {target_text}\n\n"
            f"Potential impact:\n{impact}\n\n"
            f"Safety workflow:\n1. Backup lengkap {backup_scope}\n"
            "2. Apply perubahan tanpa mereset properti lain\n"
            "3. Live verification\n"
            "4. Rollback tersedia dari backup terakhir\n\n"
            "Belum ada konfigurasi yang diubah pada tahap preview."
        )
        self.append_log(f"Preview {control_name} ditampilkan.")

    def confirm_anonymous_access_impact(self):
        dialog = QDialog(self)
        dialog.setWindowTitle("Confirm Disable Anonymous Access")
        dialog.resize(760, 340)
        dialog.setMinimumWidth(720)
        layout = QVBoxLayout(dialog)
        warning = QLabel(
            "Disabling anonymous access can interrupt access to Portal pages, public web maps, "
            "dashboards, and web applications designed for unauthenticated users.\n\n"
            "This setting does not automatically change individual item sharing levels or "
            "underlying service permissions."
        )
        warning.setWordWrap(True)
        layout.addWidget(warning)
        check_apps = QCheckBox(
            "Saya sudah memeriksa web map dan aplikasi yang membutuhkan anonymous access."
        )
        check_impact = QCheckBox(
            "Saya memahami pengguna tanpa login dapat kehilangan akses ke Portal dan aplikasi."
        )
        # QCheckBox tidak memiliki setWordWrap() pada PySide6.
        # Teks dibuat cukup ringkas dan dialog diperlebar agar tetap terbaca.
        layout.addWidget(check_apps)
        layout.addWidget(check_impact)
        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel
        )
        ok_button = buttons.button(QDialogButtonBox.StandardButton.Ok)
        ok_button.setText("Apply")
        ok_button.setEnabled(False)
        def refresh_ok():
            ok_button.setEnabled(check_apps.isChecked() and check_impact.isChecked())
        check_apps.toggled.connect(refresh_ok)
        check_impact.toggled.connect(refresh_ok)
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)
        return dialog.exec() == QDialog.DialogCode.Accepted

    def choose_feature_service_xss_profile(self):
        raw = str(self.server_result.get("feature_service_xss_raw", "")).strip().lower()
        present = self.server_result.get("feature_service_xss_present", False)
        if present and raw not in ("input", "inputoutput"):
            QMessageBox.warning(
                self, "Nilai XSS tidak dikenal",
                "Pemilihan profil dihentikan karena nilai featureServiceXSSFilter tidak dikenali."
            )
            return None
        if present and raw == "inputoutput":
            QMessageBox.information(
                self, "Profil Advanced sudah aktif",
                "Current profile sudah Advanced (inputOutput). Tidak ada perubahan yang tersedia."
            )
            return None
        current_text = (
            "Input scanning (Basic)" if raw == "input"
            else "Not configured (effective default input)"
        )
        dialog = QDialog(self)
        dialog.setWindowTitle("Choose Feature Service XSS Profile")
        dialog.resize(700, 360)
        dialog.setMinimumWidth(680)
        layout = QVBoxLayout(dialog)
        layout.setSpacing(12)
        heading = QLabel(
            "Choose Feature Service XSS Profile\n\n"
            f"Current: {current_text}\n\n"
            "Choose one profile to apply:"
        )
        heading.setWordWrap(True)
        layout.addWidget(heading)

        # Gunakan checkbox square seperti dialog Portal Servlet, tetapi tetap exclusive.
        basic_check = QCheckBox("Basic — input")
        advanced_check = QCheckBox("Advanced — inputOutput")
        profile_group = QButtonGroup(dialog)
        profile_group.setExclusive(True)
        profile_group.addButton(basic_check)
        profile_group.addButton(advanced_check)
        layout.addWidget(basic_check)
        basic_help = QLabel("    Memindai edit yang masuk. Sesuai baseline Basic.")
        basic_help.setWordWrap(True)
        layout.addWidget(basic_help)
        layout.addWidget(advanced_check)
        advanced_help = QLabel(
            "    Memindai edit dan feature yang dikembalikan ke client. "
            "Dapat menambah overhead performa."
        )
        advanced_help.setWordWrap(True)
        layout.addWidget(advanced_help)

        selection_status = QLabel()
        layout.addWidget(selection_status)
        def update_selection_status():
            if advanced_check.isChecked():
                selection_status.setText("Selected: Advanced — inputOutput")
            elif basic_check.isChecked():
                selection_status.setText("Selected: Basic — input")
            else:
                selection_status.setText("Select a profile to continue.")
        basic_check.toggled.connect(update_selection_status)
        advanced_check.toggled.connect(update_selection_status)

        # Current Basic tidak boleh dipilih ulang; Advanced menjadi pilihan upgrade.
        if raw == "input":
            basic_check.setText("Basic — input (current)")
            basic_check.setEnabled(False)
            basic_check.setToolTip("Profil Basic sedang aktif.")
            advanced_check.setChecked(True)
        else:
            basic_check.setChecked(True)
        update_selection_status()

        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel
        )
        apply_button = buttons.button(QDialogButtonBox.StandardButton.Ok)
        apply_button.setText("Apply Selected Profile")
        def refresh_apply_button():
            apply_button.setEnabled(
                basic_check.isChecked() or advanced_check.isChecked()
            )
        basic_check.toggled.connect(refresh_apply_button)
        advanced_check.toggled.connect(refresh_apply_button)
        refresh_apply_button()
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return None
        return "inputOutput" if advanced_check.isChecked() else "input"
    def confirm_feature_service_xss_advanced(self):
        dialog = QDialog(self)
        dialog.setWindowTitle("Confirm Apply Advanced XSS Profile")
        dialog.resize(760, 360)
        dialog.setMinimumWidth(720)
        layout = QVBoxLayout(dialog)
        warning = QLabel(
            "Target: featureServiceXSSFilter=inputOutput. Feature Service baru akan "
            "memindai edit masuk dan feature yang dikembalikan kepada client. Pemindaian "
            "output dapat menambah overhead performa atau memengaruhi workflow yang "
            "menggunakan HTML pada atribut dan popup."
        )
        warning.setWordWrap(True)
        layout.addWidget(warning)
        check_scope = QCheckBox(
            "Saya memahami bahwa setting ini berlaku sebagai default untuk Feature Service baru."
        )
        check_impact = QCheckBox(
            "Saya sudah mempertimbangkan dampak performa dan kompatibilitas aplikasi client."
        )
        layout.addWidget(check_scope)
        layout.addWidget(check_impact)
        status_label = QLabel("Centang kedua pernyataan untuk melanjutkan.")
        layout.addWidget(status_label)
        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel
        )
        ok_button = buttons.button(QDialogButtonBox.StandardButton.Ok)
        ok_button.setText("Apply Advanced — Set inputOutput")
        ok_button.setEnabled(False)

        def refresh_ok():
            ready = check_scope.isChecked() and check_impact.isChecked()
            ok_button.setEnabled(ready)
            status_label.setText(
                "Siap menerapkan profil Advanced." if ready
                else "Centang kedua pernyataan untuk melanjutkan."
            )

        check_scope.toggled.connect(refresh_ok)
        check_impact.toggled.connect(refresh_ok)
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)
        return dialog.exec() == QDialog.DialogCode.Accepted

    def apply_feature_service_xss_profile(self):
        target = self.choose_feature_service_xss_profile()
        if not target:
            return
        if target == "inputOutput":
            if self.confirm_feature_service_xss_advanced():
                self.start_hardening_worker("apply_advanced", "feature_service_xss")
            return

        answer = QMessageBox.warning(
            self, "Apply Basic XSS Profile",
            "Target: featureServiceXSSFilter=input\n\n"
            "Profil Basic akan menjadi default untuk Feature Service baru. "
            "System Properties lain tidak akan diubah.\n\nLanjutkan?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if answer == QMessageBox.StandardButton.Yes:
            self.start_hardening_worker("apply", "feature_service_xss")

    def choose_portal_servlets(self):
        compatibility = (self.portal_result or {}).get(
            "portal_servlet_compatibility", "UNVALIDATED"
        )
        version = str((self.portal_result or {}).get("version", "Tidak tersedia"))
        if compatibility != "SUPPORTED":
            QMessageBox.information(
                self, "Portal Servlet Hardening tidak tersedia",
                f"Portal version: {version}\nCompatibility: {compatibility}\n\n"
                "Apply diblokir secara fail-safe. Portal 12.1 menolak property servlet; "
                "versi lain harus divalidasi terlebih dahulu."
            )
            return None
        states = (self.portal_result or {}).get("portal_servlet_properties", {})
        dialog = QDialog(self)
        dialog.setWindowTitle("Select Portal Servlets to Disable")
        dialog.resize(700, 380)
        layout = QVBoxLayout(dialog)
        heading = QLabel(
            f"Portal version: {version} (SUPPORTED)\n\n"
            "Pilih servlet yang akan dinonaktifkan. Servlet yang tidak dipilih dan "
            "seluruh System Properties lain tidak akan diubah."
        )
        heading.setWordWrap(True)
        layout.addWidget(heading)
        specs = (
            ("disableLegendServlet", "Disable Legend Servlet"),
            ("disablePrintServlet", "Disable Print Servlet"),
            ("disableWFSServlet", "Disable WFS Servlet"),
        )
        checks = {}
        for name, label in specs:
            entry = states.get(name, {})
            disabled = entry.get("present") and str(entry.get("value")).strip().lower() == "true"
            current = "currently disabled" if disabled else "currently enabled/not configured"
            check = QCheckBox(f"{label} ({current})")
            if disabled:
                check.setEnabled(False)
                check.setToolTip("Servlet ini sudah disabled secara eksplisit.")
            layout.addWidget(check)
            checks[name] = check
        status = QLabel("Pilih minimal satu servlet yang belum disabled.")
        layout.addWidget(status)
        buttons = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel
        )
        ok_button = buttons.button(QDialogButtonBox.StandardButton.Ok)
        ok_button.setText("Apply Selected Servlets")
        def refresh():
            selected = [name for name, check in checks.items() if check.isChecked()]
            ok_button.setEnabled(bool(selected))
            status.setText(
                "Selected: " + ", ".join(selected)
                if selected else "Pilih minimal satu servlet yang belum disabled."
            )
        for check in checks.values():
            check.toggled.connect(refresh)
        refresh()
        buttons.accepted.connect(dialog.accept)
        buttons.rejected.connect(dialog.reject)
        layout.addWidget(buttons)
        if dialog.exec() != QDialog.DialogCode.Accepted:
            return None
        return tuple(name for name, check in checks.items() if check.isChecked())
    def show_https_guidance(self):
        evidence = self.https_enforcement_evidence()
        findings = []
        if evidence["portal_policy"] != "PASS":
            findings.append({
                "title": f"PORTAL HTTPS POLICY — {evidence['portal_policy']}",
                "rows": [("Current evidence", f'allSSL={evidence["portal_all_ssl"]}'),
                         ("Arti hasil", "Portal belum terbukti mewajibkan HTTPS untuk seluruh komunikasi organisasi."),
                         ("Alasan penilaian", "Nilai allSSL belum PASS."),
                         ("Risiko / batas evidence", "Client dapat menggunakan HTTP jika policy tidak diwajibkan."),
                         ("Recommended action", "Aktifkan Allow access to the portal through HTTPS only; periksa integration URL yang masih HTTP; lalu Analyze ulang.")]
            })
        if evidence["server_policy"] != "PASS":
            findings.append({
                "title": f"ARCGIS SERVER HTTPS POLICY — {evidence['server_policy']}",
                "rows": [("Current evidence", evidence["server_protocol"]),
                         ("Evidence source", evidence["server_protocol_source"]),
                         ("Arti hasil", "ArcGIS Server belum terbukti menggunakan HTTPS Only."),
                         ("Alasan penilaian", "protocol atau httpEnabled/sslEnabled belum menunjukkan HTTPS Only."),
                         ("Risiko / batas evidence", "Service atau administrative traffic dapat tersedia melalui HTTP."),
                         ("Recommended action", "Ubah site protocol menjadi HTTPS Only dalam maintenance window; uji service dan federation; lalu Analyze ulang.")]
            })
        for item in evidence["records"]:
            if item.get("status") in ("PASS", "PASS WITH SCOPE NOTE"):
                continue
            probe = item.get("probe", {})
            http_result = str(probe.get("http_result", "UNKNOWN"))
            https_result = str(probe.get("https_result", "UNKNOWN"))
            meaning, reason, risk, action = self.https_result_detail(http_result, item.get("status"))
            platform = str(item.get("web_server", "Unknown"))
            if http_result == "REDIRECTS_TO_NON_HTTPS" and "tomcat" in platform.lower():
                action = ('Pastikan redirect target menggunakan https://; periksa HTTP Connector redirectPort="443" '
                          'dan transport-guarantee CONFIDENTIAL pada context Web Adaptor; periksa proxy/load balancer; '
                          'restart terencana; lalu Analyze ulang.')
            elif http_result == "PLAINTEXT_CONTENT_AVAILABLE" and "iis" in platform.lower():
                action = ("Periksa IIS HTTPS binding/certificate dan terapkan HTTP Redirect atau URL Rewrite pada "
                          "site/application path yang tepat; lalu Analyze ulang.")
            findings.append({
                "title": f"{item['component'].upper()} WEB ADAPTOR {item['number']} — {item.get('status')}",
                "rows": [("URL", item.get("url") or "Tidak tersedia"),
                         ("Platform", platform),
                         ("HTTP evidence", f'{http_result} / status {probe.get("http_status") or "-"}'),
                         ("HTTPS evidence", f'{https_result} / status {probe.get("https_status") or "-"}'),
                         ("Redirect target", probe.get("http_location") or "-"),
                         ("Arti hasil", meaning), ("Alasan penilaian", reason),
                         ("Risiko / batas evidence", risk), ("Recommended action", action)]
            })
        parts = [self.report_heading("HTTPS REMEDIATION GUIDANCE", 1)]
        if findings:
            parts.append(f'<p><strong>{len(findings)} mandatory finding</strong> memerlukan tindakan. Hanya finding yang belum PASS ditampilkan.</p>')
            for finding in findings:
                parts.append(self.report_heading(finding["title"], 2))
                parts.append(self.report_kv_table([
                    (label, self.report_escape(value)) for label, value in finding["rows"]
                ]))
        else:
            parts.append(self.report_heading("NO MANDATORY REMEDIATION FOUND", 2))
            parts.append(
                '<p>Portal policy dan Server policy sudah PASS. Seluruh discovered Web Adaptor '
                'memenuhi baseline; hasil PASS WITH SCOPE NOTE tetap memiliki batas cakupan jaringan.</p>'
            )
        hsts_setting = str(evidence.get("server_hsts_enabled", "")).strip().lower()
        hsts_observed = any(bool(item.get("probe", {}).get("hsts")) for item in evidence["records"])
        if hsts_setting != "true" or not hsts_observed:
            parts.append(self.report_heading("HSTS RECOMMENDATION (OPTIONAL)", 2))
            parts.append(
                '<p>HSTS atau <strong>HTTP Strict Transport Security</strong> meminta browser selalu menggunakan '
                'HTTPS untuk hostname tertentu dan membantu mencegah SSL-stripping.</p>'
                '<p>HSTS bersifat optional strengthening karena berlaku pada seluruh hostname, bukan hanya path '
                '<code>/portal</code> atau <code>/server</code>.</p>'
                '<p><strong>Sebelum mengaktifkan HSTS, pastikan:</strong></p>'
                '<ul><li>seluruh aplikasi pada hostname mendukung HTTPS;</li>'
                '<li>proxy, load balancer, WAF, dan alias menggunakan HTTPS;</li>'
                '<li>certificate valid dan proses renewal tersedia;</li>'
                '<li>tidak ada aplikasi legacy yang masih membutuhkan HTTP.</li></ul>'
                '<p>Tidak adanya HSTS tidak menggagalkan base HTTPS compliance, tetapi tetap dicatat sebagai '
                'rekomendasi keamanan tambahan.</p>'
            )
        parts.append(self.report_heading("SETELAH REMEDIATION", 2))
        parts.append('<p>Jalankan <strong>Analyze Hardening</strong> kembali untuk mengumpulkan fresh evidence. Guidance tidak menjalankan network probe ulang.</p>')
        self.show_readable_report(
            "HTTPS Remediation Guidance", ''.join(parts), width=960, height=740, rich=True
        )
        self.append_log("Targeted guidance Verify HTTPS Enforcement ditampilkan.")
    def apply_control(self, control):
        if control == "signed_ca_certificates":
            self.show_signed_ca_guidance()
            return
        if control == "https_enforcement":
            self.show_https_guidance()
            return
        if not self.control_supported(control):
            QMessageBox.information(
                self, "Apply belum tersedia",
                "Kontrol ini belum siap untuk Apply. Konfigurasi target diperlukan terlebih dahulu."
            )
            return
        labels = {
            "portal_directory": "Disable Portal Directory",
            "automatic_account": "Disable Automatic Enterprise Account Creation",
            "builtin_self_creation": "Disable Built-In Account Self-Creation",
            "public_profile_sharing": "Disable Public User Profile Sharing",
            "social_media_links": "Disable Show Social Media Links",
            "anonymous_access": "Disable Anonymous Access",
            "portal_servlets": "Disable Portal Servlets",
            "member_defaults_viewer": "Configure New Member Default Role as Viewer",
            "jsonp": "Disable JSONP Callback Functions",
            "services_directory": "Disable Services Directory",
            "standardized_queries": "Enable Standardized Queries",
            "feature_service_xss": "Verify Feature Service XSS Filter Default",
            "token_http_get": "Disable Token Acquisition via HTTP GET",
        }
        label = labels[control]
        if control == "anonymous_access":
            if self.confirm_anonymous_access_impact():
                self.start_hardening_worker("apply", control)
            return
        if control == "portal_servlets":
            selected = self.choose_portal_servlets()
            if selected:
                self.portal_servlet_selection = selected
                self.start_hardening_worker("apply", control)
            return
        if control == "feature_service_xss":
            self.apply_feature_service_xss_profile()
            return
        answer = QMessageBox.warning(
            self, f"Apply {label}",
            f"Aplikasi akan menjalankan {label}.\n\n"
            "Backup JSON dibuat sebelum perubahan dan hasil diverifikasi secara live.\n\n"
            "Lanjutkan?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if answer == QMessageBox.StandardButton.Yes:
            self.start_hardening_worker("apply", control)

    def rollback_control(self, control):
        backup = self.backup_for_control(control)
        if not backup:
            QMessageBox.warning(
                self, "Backup tidak tersedia",
                "Belum ada backup yang tersedia untuk kontrol ini."
            )
            return
        labels = {
            "portal_directory": "Portal Directory",
            "automatic_account": "Automatic Enterprise Account Creation",
            "builtin_self_creation": "Built-In Account Self-Creation",
            "public_profile_sharing": "Public User Profile Sharing",
            "social_media_links": "Show Social Media Links",
            "anonymous_access": "Anonymous Access",
            "portal_servlets": "Portal Servlets",
            "member_defaults_viewer": "New Member Default Role as Viewer",
            "jsonp": "JSONP",
            "services_directory": "Services Directory",
            "standardized_queries": "Standardized Queries",
            "feature_service_xss": "Feature Service XSS Filter Default",
            "token_http_get": "Token Acquisition via HTTP GET",
        }
        label = labels.get(control, control)
        answer = QMessageBox.warning(
            self, f"Rollback {label}",
            f"Konfigurasi {label} akan dikembalikan dari backup terakhir dan "
            "diverifikasi live.\n\n"
            f"Backup: {backup}\n\nLanjutkan?",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No,
            QMessageBox.StandardButton.No,
        )
        if answer == QMessageBox.StandardButton.Yes:
            self.start_hardening_worker("rollback", control)

    def start_hardening_worker(self, action, control):
        if self.hardening_thread is not None:
            return
        self.set_busy(True)
        if control == "portal_directory":
            backup = self.last_portal_directory_backup_file
        elif control == "automatic_account":
            backup = self.last_automatic_account_backup_file
        elif control == "builtin_self_creation":
            backup = self.last_builtin_self_creation_backup_file
        elif control == "public_profile_sharing":
            backup = self.last_public_profile_sharing_backup_file
        elif control == "social_media_links":
            backup = self.last_social_media_links_backup_file
        elif control == "anonymous_access":
            backup = self.last_anonymous_access_backup_file
        elif control == "portal_servlets":
            backup = self.last_portal_servlets_backup_file
        elif control == "member_defaults_viewer":
            backup = self.last_member_defaults_viewer_backup_file
        elif control == "jsonp":
            backup = self.last_jsonp_backup_file
        elif control == "standardized_queries":
            backup = self.last_standardized_queries_backup_file
        elif control == "feature_service_xss":
            backup = self.last_feature_service_xss_backup_file
        elif control == "token_http_get":
            backup = self.last_token_http_get_backup_file
        else:
            backup = self.last_services_directory_backup_file
        self.hardening_thread = QThread()
        if control in (
            "portal_directory", "automatic_account", "builtin_self_creation",
            "public_profile_sharing", "social_media_links", "anonymous_access",
            "portal_servlets", "member_defaults_viewer",
        ):
            self.hardening_worker = PortalDirectoryHardeningWorker(
                action=action,
                control=control,
                portal_admin_url=self.portal_admin_url.text(),
                portal_username=self.portal_username.text(),
                portal_password=self.portal_password.text(),
                verify_ssl=self.verify_ssl.isChecked(),
                portal_token=self.portal_token,
                portal_token_expires=self.portal_token_expires,
                backup_file=backup,
                selected_properties=self.portal_servlet_selection,
                portal_version=(self.portal_result or {}).get("version"),
            )
        else:
            self.hardening_worker = ServerDirectoryHardeningWorker(
                action=action,
                control=control,
                server_admin_url=self.server_admin_url.text(),
                server_username=self.server_username.text(),
                server_password=self.server_password.text(),
                verify_ssl=self.verify_ssl.isChecked(),
                server_token=self.server_token,
                server_token_expires=self.server_token_expires,
                backup_file=backup,
            )
        self.hardening_worker.moveToThread(self.hardening_thread)
        self.hardening_thread.started.connect(self.hardening_worker.run)
        self.hardening_worker.progress.connect(self.on_progress)
        self.hardening_worker.completed.connect(self.on_hardening_action_completed)
        self.hardening_worker.finished.connect(self.hardening_thread.quit)
        self.hardening_worker.finished.connect(self.hardening_worker.deleteLater)
        self.hardening_thread.finished.connect(self.hardening_thread.deleteLater)
        self.hardening_thread.finished.connect(self.on_hardening_worker_finished)
        self.hardening_thread.start()

    def update_verified_control_immediately(self, control, verified_value):
        """Update memory dan tabel segera dari hasil live verification worker."""
        normalized = str(verified_value).lower() == "true"
        if not self.server_result:
            self.server_result = {"success": True}

        if control == "portal_directory":
            if not self.portal_result:
                self.portal_result = {"success": True}
            self.portal_result["portal_directory_disabled"] = normalized
        elif control == "automatic_account":
            if not self.portal_result:
                self.portal_result = {"success": True}
            self.portal_result["automatic_account_creation"] = normalized
        elif control == "builtin_self_creation":
            if not self.portal_result:
                self.portal_result = {"success": True}
            verified_text = str(verified_value)
            present = not verified_text.startswith("Not configured")
            self.portal_result["builtin_self_creation_present"] = present
            self.portal_result["builtin_self_creation_disabled"] = (
                True if verified_text.startswith("Disabled (explicit)")
                else False if verified_text.startswith("Enabled (explicit)")
                else None
            )
        elif control == "public_profile_sharing":
            if not self.portal_result:
                self.portal_result = {"success": True}
            verified_text = str(verified_value)
            present = not verified_text.startswith("Unknown")
            self.portal_result["public_profile_setting_present"] = present
            self.portal_result["public_profile_update_disabled"] = (
                True if verified_text.startswith("Disabled (explicit)")
                else False if verified_text.startswith("Enabled (explicit)")
                else None
            )
        elif control == "anonymous_access":
            if not self.portal_result:
                self.portal_result = {"success": True}
            verified_text = str(verified_value)
            self.portal_result["access"] = (
                "private" if verified_text.startswith("Disabled (private)")
                else "public" if verified_text.startswith("Enabled (public)")
                else "Tidak tersedia"
            )
        elif control == "social_media_links":
            if not self.portal_result:
                self.portal_result = {"success": True}
            verified_text = str(verified_value)
            present = not verified_text.startswith("Not configured")
            self.portal_result["social_media_links_present"] = present
            self.portal_result["social_media_links_raw"] = (
                False if verified_text.startswith("Disabled (explicit)")
                else True if verified_text.startswith("Enabled (explicit)")
                else None
            )
            if not isinstance(self.portal_result.get("portal_properties"), dict):
                self.portal_result["portal_properties"] = {}
            if present and self.portal_result["social_media_links_raw"] is not None:
                self.portal_result["portal_properties"]["showSocialMediaLinks"] = (
                    self.portal_result["social_media_links_raw"]
                )
            else:
                self.portal_result["portal_properties"].pop("showSocialMediaLinks", None)
        elif control == "member_defaults_viewer":
            if not self.portal_result:
                self.portal_result = {"success": True}
            settings = verified_value if isinstance(verified_value, dict) else {}
            self.portal_result["user_default_settings"] = settings
        elif control == "portal_servlets":
            if not self.portal_result:
                self.portal_result = {"success": True}
            verified_properties = verified_value if isinstance(verified_value, dict) else None
            if verified_properties is not None:
                self.portal_result["system_properties"] = verified_properties
                self.portal_result["portal_servlet_properties"] = {
                    name: {
                        "present": name in verified_properties,
                        "value": verified_properties.get(name),
                    }
                    for name in (
                        "disableLegendServlet", "disablePrintServlet", "disableWFSServlet"
                    )
                }
        elif control == "jsonp":
            self.server_result["callback_functions_enabled"] = normalized
        elif control == "token_http_get":
            verified_text = str(verified_value)
            present = not verified_text.startswith("Not configured")
            self.server_result["allow_http_get_present"] = present
            self.server_result["allow_http_get_raw"] = (
                "false" if verified_text.startswith("Disabled (explicit)")
                else "true" if verified_text.startswith("Enabled (explicit)")
                else None
            )
        elif control == "feature_service_xss":
            verified_text = str(verified_value)
            present = not verified_text.startswith("Not configured")
            self.server_result["feature_service_xss_present"] = present
            self.server_result["feature_service_xss_raw"] = (
                "inputOutput" if verified_text.startswith("Input + output")
                else "input" if verified_text.startswith("Input scanning")
                else None
            )
        elif control == "standardized_queries":
            verified_text = str(verified_value)
            property_present = not verified_text.startswith("Not configured")
            self.server_result["standardized_queries_present"] = property_present
            self.server_result["standardized_queries_raw"] = (
                "true" if verified_text.startswith("Enabled (explicit)")
                else "false" if verified_text.startswith("Disabled (explicit)")
                else None
            )
        else:
            self.server_result["services_directory_enabled"] = normalized

        if control == "portal_servlets":
            self.render_controls()
            row = self.control_rows.get(control)
        else:
            row = self.control_rows.get(control)
        if row is not None and control != "portal_servlets":
            current_item = self.controls_table.item(row, 2)
            status_item = self.controls_table.item(row, 4)
            if current_item:
                if control in (
                    "standardized_queries", "feature_service_xss", "builtin_self_creation",
                    "token_http_get", "public_profile_sharing", "social_media_links", "anonymous_access",
                ):
                    full_value = str(verified_value)
                    current_item.setText(" ".join(full_value.split()))
                    current_item.setToolTip(full_value)
                else:
                    current_item.setText("true" if normalized else "false")
            if status_item:
                if control == "standardized_queries":
                    compliant = str(verified_value).startswith("Enabled (explicit)")
                elif control == "feature_service_xss":
                    compliant = (
                        str(verified_value).startswith("Input scanning")
                        or str(verified_value).startswith("Input + output")
                    )
                elif control in (
                    "builtin_self_creation", "token_http_get", "public_profile_sharing",
                    "social_media_links",
                ):
                    compliant = str(verified_value).startswith("Disabled (explicit)")
                elif control == "anonymous_access":
                    compliant = str(verified_value).startswith("Disabled (private)")
                else:
                    compliant = normalized if control == "portal_directory" else not normalized
                status = "COMPLIANT" if compliant else "NON-COMPLIANT"
                status_item.setData(Qt.ItemDataRole.UserRole, status)
                status_item.setText(self.wrap_table_words(status, 1))
                status_item.setToolTip(status)
                status_item.setForeground(self.semantic_status_color(status))

        self.update_connection_results()
        stamp = datetime.now().strftime("%d %b %Y %H:%M:%S")
        self.last_refreshed.setText(
            f"Data source: LIVE VERIFIED | Last verified: {stamp} | Component: OK"
        )
        self.activity.setText(
            "Activity: Perubahan berhasil diverifikasi. Menjadwalkan full live Analyze..."
        )
        self.refresh_action_buttons()

    def on_hardening_action_completed(self, result):
        self.portal_token = result.get("portal_token", self.portal_token)
        self.portal_token_expires = result.get(
            "portal_token_expires", self.portal_token_expires
        )
        self.server_token = result.get("server_token", self.server_token)
        self.server_token_expires = result.get(
            "server_token_expires", self.server_token_expires
        )
        action = result.get("action")
        control = result.get("control")
        label = (
            "Portal Directory" if control == "portal_directory"
            else "Automatic Enterprise Account Creation" if control == "automatic_account"
            else "Built-In Account Self-Creation" if control == "builtin_self_creation"
            else "Public User Profile Sharing" if control == "public_profile_sharing"
            else "Show Social Media Links" if control == "social_media_links"
            else "Anonymous Access" if control == "anonymous_access"
            else "Portal Servlets" if control == "portal_servlets"
            else "New Member Default Role as Viewer" if control == "member_defaults_viewer"
            else "JSONP" if control == "jsonp"
            else "Standardized Queries" if control == "standardized_queries"
            else "Feature Service XSS Filter Default" if control == "feature_service_xss"
            else "Token Acquisition via HTTP GET" if control == "token_http_get"
            else "Services Directory"
        )
        if not result.get("success"):
            QMessageBox.critical(
                self, f"{label} operation gagal",
                f"Action: {action}\n\n{result.get('error', 'Penyebab tidak tersedia.')}"
            )
            self.append_log(f"{label} {action} gagal: {result.get('error')}")
            return

        # Worker sudah membaca ulang endpoint live. Gunakan hasil verifikasi itu
        # untuk memperbarui tabel segera, tanpa menunggu full Analyze berikutnya.
        self.update_verified_control_immediately(
            control,
            result.get("verified_properties")
            if control == "portal_servlets"
            else result.get("verified_settings")
            if control == "member_defaults_viewer"
            else result.get("verified_value")
        )

        if action in ("apply", "apply_advanced"):
            if result.get("already_compliant"):
                QMessageBox.information(
                    self, f"{label} sudah compliant",
                    f"{label} sudah compliant. Tidak ada perubahan yang diperlukan."
                )
            else:
                backup = result.get("backup_file")
                if control == "portal_directory":
                    self.last_portal_directory_backup_file = backup
                elif control == "automatic_account":
                    self.last_automatic_account_backup_file = backup
                elif control == "builtin_self_creation":
                    self.last_builtin_self_creation_backup_file = backup
                elif control == "public_profile_sharing":
                    self.last_public_profile_sharing_backup_file = backup
                elif control == "social_media_links":
                    self.last_social_media_links_backup_file = backup
                elif control == "anonymous_access":
                    self.last_anonymous_access_backup_file = backup
                elif control == "portal_servlets":
                    self.last_portal_servlets_backup_file = backup
                elif control == "member_defaults_viewer":
                    self.last_member_defaults_viewer_backup_file = backup
                elif control == "jsonp":
                    self.last_jsonp_backup_file = backup
                elif control == "standardized_queries":
                    self.last_standardized_queries_backup_file = backup
                elif control == "feature_service_xss":
                    self.last_feature_service_xss_backup_file = backup
                elif control == "token_http_get":
                    self.last_token_http_get_backup_file = backup
                else:
                    self.last_services_directory_backup_file = backup
                profile_text = result.get("applied_profile")
                profile_line = f"Profile: {profile_text}\n" if profile_text else ""
                QMessageBox.information(
                    self, f"{label} hardening berhasil",
                    "Backup JSON berhasil dibuat.\nApply berhasil.\n"
                    f"{profile_line}Live verification: nilai aktual = "
                    f"{result.get('verified_value')}\n\nBackup: {backup}"
                )
                self.append_log(f"{label} Apply verified. Backup: {backup}")
                self.refresh_action_buttons()
        else:
            QMessageBox.information(
                self, f"Rollback {label} berhasil",
                "Konfigurasi berhasil dikembalikan dari backup dan diverifikasi live.\n\n"
                f"Nilai aktual = {result.get('verified_value')}"
            )
            self.append_log(f"{label} Rollback verified.")
            self.refresh_action_buttons()
        self.followup_live_analyze_pending = True

    def on_hardening_worker_finished(self):
        self.hardening_worker = None
        self.hardening_thread = None
        self.set_busy(False)
        if getattr(self, "followup_live_analyze_pending", False):
            self.followup_live_analyze_pending = False
            # Beri waktu singkat agar perubahan terpropagasi melalui Web Adaptor/proxy.
            QTimer.singleShot(750, self.run_followup_live_analyze)

    def run_followup_live_analyze(self):
        if self.thread is not None or self.hardening_thread is not None:
            QTimer.singleShot(500, self.run_followup_live_analyze)
            return
        self.append_log("Menjalankan full live Analyze setelah hardening...")
        self.start_request("all", "analyze")

    def invalidate_connection(self):
        if self.connected:
            self.connected = False
            self.portal_token = None
            self.portal_token_expires = 0
            self.server_token = None
            self.server_token_expires = 0
            self.tabs.setTabEnabled(1, False)
            self.controls_table.setRowCount(0)
            self.last_assessment_at = None
            self.export_assessment_button.setEnabled(False)
            self.last_refreshed.setText("Data source: target connection berubah")
            self.append_log("Target berubah; Hardening Controls dikunci kembali.")

    def clear_fields(self):
        for field in (
            self.portal_admin_url, self.portal_username, self.portal_password,
            self.server_admin_url, self.server_username, self.server_password,
        ):
            field.clear()
        self.portal_result = None
        self.server_result = None
        self.last_builtin_self_creation_backup_file = None
        self.last_member_defaults_viewer_backup_file = None
        self.last_public_profile_sharing_backup_file = None
        self.last_social_media_links_backup_file = None
        self.last_anonymous_access_backup_file = None
        self.last_jsonp_backup_file = None
        self.last_services_directory_backup_file = None
        self.last_standardized_queries_backup_file = None
        self.last_feature_service_xss_backup_file = None
        self.last_token_http_get_backup_file = None
        self.portal_token = None
        self.portal_token_expires = 0
        self.server_token = None
        self.server_token_expires = 0
        self.connected = False
        self.portal_status.setText("Portal status: Belum diuji.")
        self.server_status.setText("Server status: Belum diuji.")
        self.tabs.setTabEnabled(1, False)
        self.controls_table.setRowCount(0)
        self.control_rows = {}
        self.action_buttons = {}
        self.activity.setText("Activity: Ready")

    @staticmethod
    def sanitize_evidence(value):
        sensitive = {
            "token", "portal_token", "server_token", "password", "portal_password",
            "server_password", "cookie", "authorization", "sharedkey", "shared_key",
        }
        if isinstance(value, dict):
            clean = {}
            for key, item in value.items():
                normalized = str(key).lower().replace("-", "_")
                if normalized in sensitive or normalized.endswith("_token") or "password" in normalized:
                    continue
                clean[str(key)] = MainWindow.sanitize_evidence(item)
            return clean
        if isinstance(value, (list, tuple)):
            return [MainWindow.sanitize_evidence(item) for item in value]
        if isinstance(value, Path):
            return str(value)
        if isinstance(value, (str, int, float, bool)) or value is None:
            return value
        return str(value)
    def assessment_control_catalog(self):
        return {
            "https_enforcement": {
                "purpose": "Memastikan Portal, ArcGIS Server, dan seluruh registered Web Adaptor menggunakan HTTPS sesuai baseline.",
                "checks": "Portal allSSL, Server protocol/httpEnabled/sslEnabled, registered Web Adaptor, HTTP redirect/exposure, HTTPS reachability, dan HSTS sebagai optional strengthening.",
                "risk": "Traffic dapat tersedia tanpa enkripsi atau evidence jaringan belum mencakup jalur lain.",
                "action": "Ikuti targeted guidance untuk policy Portal/Server dan web tier, kemudian jalankan Analyze kembali.",
                "type": "Verify-only", "restart": "Tidak oleh tools; remediation manual mungkin memerlukan restart web tier atau service.",
            },
            "signed_ca_certificates": {
                "purpose": "Memastikan CA-signed certificate aktif pada web tier, Portal native HTTPS, dan setiap ArcGIS Server machine.",
                "checks": "Unique TLS listener; active alias; leaf, intermediate, root; issuer/subject; SAN; validity; key size; signature; fingerprint; dan native machine identity.",
                "risk": "Client dapat menolak koneksi atau tidak dapat memverifikasi identitas endpoint.",
                "action": "Import dan assign CA-signed certificate beserta chain pada layer yang belum memenuhi baseline, lalu Analyze kembali.",
                "type": "Verify-only", "restart": "Tidak oleh tools; remediation manual mungkin memerlukan restart atau propagation.",
            },
            "portal_directory": {"purpose":"Mengurangi paparan browsable Portal Administrator Directory.","checks":"Portal security configuration disableServicesDirectory.","risk":"Directory browsable membantu reconnaissance endpoint administratif.","action":"Set disableServicesDirectory=true jika belum compliant.","type":"Configurable","restart":"Tidak diharapkan."},
            "automatic_account": {"purpose":"Mencegah pembuatan enterprise account otomatis saat first sign-in.","checks":"Portal security configuration enableAutomaticAccountCreation.","risk":"Identitas dapat menjadi member tanpa proses provisioning yang disetujui.","action":"Set enableAutomaticAccountCreation=false.","type":"Configurable","restart":"Tidak diharapkan."},
            "member_defaults_viewer": {"purpose":"Menerapkan least privilege untuk member baru.","checks":"userDefaultSettings, Viewer userLicenseType, dan Viewer role yang di-resolve dari katalog live.","risk":"Member baru dapat menerima privilege lebih tinggi dari kebutuhan.","action":"Set default user type dan role ke Viewer.","type":"Configurable","restart":"Tidak."},
            "builtin_self_creation": {"purpose":"Mencegah pengguna membuat built-in Portal account secara mandiri.","checks":"Portal System Property disableSignup.","risk":"Account dapat dibuat di luar proses provisioning.","action":"Set disableSignup=true.","type":"Configurable","restart":"Portal restart atau propagation mungkin terjadi."},
            "public_profile_sharing": {"purpose":"Membatasi perubahan dan publikasi informasi profil oleh member.","checks":"Portal property updateUserProfileDisabled.","risk":"Informasi profil dapat mendukung reconnaissance dan social engineering.","action":"Set updateUserProfileDisabled=true.","type":"Configurable","restart":"Tidak diharapkan."},
            "social_media_links": {"purpose":"Menyembunyikan tautan social media pada item dan group.","checks":"portalProperties.showSocialMediaLinks.","risk":"Tautan eksternal menambah jalur keluar dan paparan informasi.","action":"Set showSocialMediaLinks=false.","type":"Configurable","restart":"Tidak diharapkan."},
            "anonymous_access": {"purpose":"Mewajibkan autentikasi untuk akses Portal.","checks":"Portal access dan anonymous evidence probe.","risk":"Portal, web map, dashboard, atau aplikasi dapat diakses tanpa login.","action":"Set access=private setelah impact review.","type":"Configurable","restart":"Tidak diharapkan; aplikasi public dapat terdampak."},
            "portal_servlets": {"purpose":"Menonaktifkan servlet Portal yang tidak diperlukan secara selektif.","checks":"disableLegendServlet, disablePrintServlet, disableWFSServlet dan kompatibilitas versi Portal.","risk":"Servlet yang tidak diperlukan menambah attack surface.","action":"Pilih servlet yang aman dinonaktifkan setelah dependency review.","type":"Selective configurable","restart":"Portal restart atau propagation diharapkan."},
            "services_directory": {"purpose":"Menonaktifkan browsable HTML ArcGIS Server Services Directory.","checks":"servicesDirEnabled.","risk":"Directory mempermudah enumeration service dan metadata.","action":"Set servicesDirEnabled=false.","type":"Configurable","restart":"Tidak diharapkan."},
            "jsonp": {"purpose":"Menonaktifkan legacy JSONP callback functions.","checks":"callbackFunctionsEnabled.","risk":"JSONP memperluas risiko cross-origin script execution pada aplikasi legacy.","action":"Set callbackFunctionsEnabled=false setelah compatibility review.","type":"Configurable","restart":"Tidak diharapkan."},
            "token_http_get": {"purpose":"Mencegah pengiriman credential melalui URL saat memperoleh token.","checks":"Token Manager properties.allowHttpGet dan positive POST test.","risk":"Credential pada query string dapat tersimpan di history dan log.","action":"Set allowHttpGet=false dan gunakan HTTP POST.","type":"Configurable","restart":"Tidak diharapkan."},
            "standardized_queries": {"purpose":"Memastikan standardized SQL query aktif secara eksplisit.","checks":"Server System Property standardizedQueries.","risk":"Non-standard query dapat meningkatkan risiko SQL injection atau query yang tidak terkontrol.","action":"Set standardizedQueries=true dan uji query aplikasi.","type":"Configurable","restart":"Tidak diharapkan."},
            "feature_service_xss": {"purpose":"Menetapkan XSS filtering default untuk Feature Service baru.","checks":"Server System Property featureServiceXSSFilter.","risk":"Input atau output berbahaya dapat mencapai client jika filtering tidak memadai.","action":"Pilih Basic input atau Advanced inputOutput dan lakukan regression test.","type":"Configurable profile","restart":"Tidak diharapkan."},
            "allowed_origins": {"purpose":"Membatasi origin yang boleh mengakses ArcGIS REST resources.","checks":"Services Directory allowedOrigins.","risk":"Wildcard origin memperluas permukaan cross-origin access.","action":"Tentukan daftar trusted origins sebelum penerapan.","type":"Guidance / requires target","restart":"Tergantung implementasi."},
        }
    def collect_assessment_controls(self):
        catalog = self.assessment_control_catalog()
        controls = []
        for control_id, row in sorted(self.control_rows.items(), key=lambda pair: pair[1]):
            values = []
            for column in range(6):
                item = self.controls_table.item(row, column)
                values.append(item.toolTip() if item and item.toolTip() else item.text() if item else "")
            status_item = self.controls_table.item(row, 4)
            status = (status_item.data(Qt.ItemDataRole.UserRole) if status_item else None) or values[4]
            detail = catalog.get(control_id, {})
            controls.append({
                "id": control_id, "component": values[0], "control": values[1],
                "current": values[2], "target": values[3], "status": str(status),
                "risk": values[5], "control_type": detail.get("type", "Assessment"),
                "purpose": detail.get("purpose", "Menilai konfigurasi keamanan terhadap baseline utility."),
                "checks": detail.get("checks", "Current configuration dan target baseline."),
                "risk_explanation": detail.get("risk", "Kondisi memerlukan review berdasarkan status assessment."),
                "recommended_action": detail.get("action", "Review evidence dan jalankan Analyze kembali setelah perubahan."),
                "restart": detail.get("restart", "Review sesuai komponen."),
                "backup": self.backup_for_control(control_id),
            })
        return controls
    @staticmethod
    def evidence_hash(value):
        payload = json.dumps(value, ensure_ascii=False, sort_keys=True, default=str).encode("utf-8")
        return hashlib.sha256(payload).hexdigest()
    @staticmethod
    def assessment_hostname(url):
        hostname = (urlparse(str(url or "")).hostname or "unknown-environment").lower()
        safe = "".join(ch if ch.isalnum() or ch == "-" else "_" for ch in hostname)
        while "__" in safe:
            safe = safe.replace("__", "_")
        return safe.strip("_")[:80] or "unknown-environment"

    @staticmethod
    def normalized_environment_url(value):
        parsed = urlparse(normalize_url(str(value or "")))
        host = (parsed.hostname or "").lower()
        port = f":{parsed.port}" if parsed.port else ""
        path = "/" + "/".join(part for part in parsed.path.lower().split("/") if part)
        return f"{parsed.scheme.lower()}://{host}{port}{path}".rstrip("/")

    def build_assessment_package(self, assessment_type="Baseline Assessment", baseline_id=None):
        timestamp = self.last_assessment_at or datetime.now().astimezone()
        portal_url = self.portal_admin_url.text().strip()
        server_url = self.server_admin_url.text().strip()
        assessment_id = (
            f"AEH-{self.assessment_hostname(portal_url)}-"
            f"{timestamp.strftime('%Y%m%d-%H%M%S')}"
        )
        controls = self.collect_assessment_controls()
        evidence = self.sanitize_evidence({
            "portal": self.portal_result or {}, "server": self.server_result or {},
        })
        return {
            "assessment_id": assessment_id,
            "assessment_type": assessment_type,
            "phase": "BASELINE" if assessment_type == "Baseline Assessment" else "CURRENT",
            "comparison_performed": False,
            "baseline_assessment_id": baseline_id,
            "assessed_at": timestamp.isoformat(),
            "operator": "Anggi Yulianto",
            "portal_admin_url": portal_url,
            "server_admin_url": server_url,
            "portal_organization_id": (self.portal_result or {}).get("id", "Tidak tersedia"),
            "environment_hostname": urlparse(portal_url).hostname or "Tidak tersedia",
            "portal_version": (self.portal_result or {}).get("version", "Tidak tersedia"),
            "server_version": (self.server_result or {}).get("version", "Tidak tersedia"),
            "controls": controls,
            "evidence": evidence,
            "evidence_sha256": self.evidence_hash(evidence),
            "sensitive_values_included": False,
        }

    @staticmethod
    def docx_set_cell_shading(cell, fill):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        tc_pr = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        tc_pr.append(shading)
    @staticmethod
    def docx_set_cell_text(cell, text, bold=False, color=None, size=9):
        cell.text = ""
        paragraph = cell.paragraphs[0]
        run = paragraph.add_run(str(text if text not in (None, "") else "-"))
        run.bold = bold
        run.font.size = __import__("docx").shared.Pt(size)
        if color:
            run.font.color.rgb = __import__("docx").shared.RGBColor.from_string(color)
        cell.vertical_alignment = __import__("docx").enum.table.WD_CELL_VERTICAL_ALIGNMENT.CENTER
    def export_assessment_docx(self, package, path):
        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.shared import Inches, Pt, RGBColor
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.7); section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7); section.right_margin = Inches(0.7)
        styles = document.styles
        styles["Normal"].font.name = "Aptos"; styles["Normal"].font.size = Pt(9.5)
        for style_name, size, color in (("Title",26,"17365D"),("Heading 1",17,"17365D"),("Heading 2",13,"2F5597"),("Heading 3",11,"44546A")):
            style=styles[style_name]; style.font.name="Aptos Display"; style.font.size=Pt(size); style.font.color.rgb=RGBColor.from_string(color)
        title=document.add_paragraph(); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
        run=title.add_run("ArcGIS Enterprise\nSecurity Hardening Assessment"); run.bold=True; run.font.size=Pt(25); run.font.color.rgb=RGBColor(23,54,93)
        subtitle=document.add_paragraph("Current Assessment Report"); subtitle.alignment=WD_ALIGN_PARAGRAPH.CENTER
        subtitle.runs[0].font.size=Pt(14); subtitle.runs[0].font.color.rgb=RGBColor(91,105,120)
        document.add_paragraph("")
        meta=document.add_table(rows=0, cols=2); meta.alignment=WD_TABLE_ALIGNMENT.CENTER
        metadata=[("Assessment ID",package["assessment_id"]),("Assessment date",package["assessed_at"]),("Prepared by",package["operator"]),("Portal Admin URL",package["portal_admin_url"]),("Server Admin URL",package["server_admin_url"]),("Portal version",package["portal_version"]),("Server version",package["server_version"]),("Configuration changes","None. Report reflects the latest successful Analyze." )]
        for label,value in metadata:
            cells=meta.add_row().cells; self.docx_set_cell_shading(cells[0],"D9EAF7"); self.docx_set_cell_text(cells[0],label,True,"17365D",9); self.docx_set_cell_text(cells[1],value,False,None,9)
        document.add_page_break()
        controls=package["controls"]
        counts={status:sum(1 for c in controls if c["status"]==status) for status in sorted({c["status"] for c in controls})}
        document.add_heading("Executive Summary", level=1)
        document.add_paragraph("Assessment ini mendokumentasikan current security posture berdasarkan live evidence dari ArcGIS Enterprise API, registered endpoints, network probes, dan certificate metadata. Tidak ada screenshot manual dan tidak ada perubahan konfigurasi saat export.")
        summary=document.add_table(rows=1,cols=2); summary.alignment=WD_TABLE_ALIGNMENT.LEFT
        self.docx_set_cell_text(summary.rows[0].cells[0],"Metric",True,"FFFFFF"); self.docx_set_cell_shading(summary.rows[0].cells[0],"17365D")
        self.docx_set_cell_text(summary.rows[0].cells[1],"Count",True,"FFFFFF"); self.docx_set_cell_shading(summary.rows[0].cells[1],"17365D")
        for label,value in [("Total controls",len(controls))]+list(counts.items()):
            cells=summary.add_row().cells; self.docx_set_cell_text(cells[0],label); self.docx_set_cell_text(cells[1],value)
        document.add_heading("Control Summary", level=1)
        table=document.add_table(rows=1,cols=7); table.alignment=WD_TABLE_ALIGNMENT.CENTER
        headers=["No.","Component","Control","Current","Target","Status","Risk"]
        for i,h in enumerate(headers): self.docx_set_cell_text(table.rows[0].cells[i],h,True,"FFFFFF",8); self.docx_set_cell_shading(table.rows[0].cells[i],"17365D")
        for idx,c in enumerate(controls,1):
            cells=table.add_row().cells
            values=[idx,c["component"],c["control"],c["current"],c["target"],c["status"],c["risk"]]
            for i,v in enumerate(values): self.docx_set_cell_text(cells[i],v,False,None,7.5)
            if c["status"] in ("NON-COMPLIANT","CRITICAL"): self.docx_set_cell_shading(cells[5],"F4CCCC")
            elif c["status"] in ("COMPLIANT","COMPLIANT WITH NOTE"): self.docx_set_cell_shading(cells[5],"D9EAD3")
            else: self.docx_set_cell_shading(cells[5],"FFF2CC")
        document.add_page_break()
        document.add_heading("Detailed Control Assessment", level=1)
        for idx,c in enumerate(controls,1):
            document.add_heading(f"{idx}. {c['control']}", level=2)
            info=document.add_table(rows=0,cols=2)
            for label,value in [("Component",c["component"]),("Control type",c["control_type"]),("Risk",c["risk"]),("Status",c["status"]),("Current",c["current"]),("Target",c["target"]),("Restart implication",c["restart"])]:
                cells=info.add_row().cells; self.docx_set_cell_shading(cells[0],"EAF2F8"); self.docx_set_cell_text(cells[0],label,True,"17365D",8.5); self.docx_set_cell_text(cells[1],value,False,None,8.5)
            for heading,text in [("Tujuan control",c["purpose"]),("Apa yang diperiksa",c["checks"]),("Alasan penilaian",f"Current value dibandingkan terhadap target baseline menghasilkan status {c['status']}.") ,("Risiko / batas evidence",c["risk_explanation"]),("Recommended action",c["recommended_action"]),("Evidence reference",f"assessment-summary.json | Control ID: {c['id']} | Evidence package SHA-256: {package['evidence_sha256']}")]:
                document.add_heading(heading, level=3); document.add_paragraph(text)
            document.add_paragraph("")
        document.add_heading("Evidence and Limitations", level=1)
        document.add_paragraph("Evidence package tidak menyimpan password, administrator token, cookie, authorization header, shared key, atau private key. Network evidence hanya mewakili jalur komputer operator. DNS, hosts file, firewall, IIS, Tomcat, WAF, dan load balancer tidak dibaca langsung kecuali melalui response yang terlihat dari endpoint.")
        footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.add_run(f"{package['assessment_id']} | ArcGIS Enterprise Hardening Utility")
        document.save(path)
    def export_assessment_xlsx(self, package, path):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
        from openpyxl.worksheet.table import Table, TableStyleInfo
        wb=Workbook(); ws=wb.active; ws.title="Assessment Summary"; ws.sheet_view.showGridLines=False
        dark="17365D"; blue="D9EAF7"; green="D9EAD3"; red="F4CCCC"; orange="FCE5CD"; gray="E7E6E6"
        ws["A1"]="ArcGIS Enterprise Security Hardening Assessment"; ws["A1"].font=Font(size=18,bold=True,color="FFFFFF"); ws["A1"].fill=PatternFill("solid",fgColor=dark); ws.merge_cells("A1:D1")
        meta=[("Assessment ID",package["assessment_id"]),("Assessment date",package["assessed_at"]),("Operator",package["operator"]),("Portal Admin URL",package["portal_admin_url"]),("Server Admin URL",package["server_admin_url"]),("Portal version",package["portal_version"]),("Server version",package["server_version"]),("Evidence SHA-256",package["evidence_sha256"])]
        for r,(label,value) in enumerate(meta,3): ws.cell(r,1,label); ws.cell(r,2,value); ws.cell(r,1).font=Font(bold=True,color=dark); ws.cell(r,1).fill=PatternFill("solid",fgColor=blue)
        controls=package["controls"]; start=13; ws.cell(start,1,"Status"); ws.cell(start,2,"Count")
        statuses=sorted({c["status"] for c in controls})
        for i,status in enumerate(statuses,start+1): ws.cell(i,1,status); ws.cell(i,2,f'=COUNTIF(\'Control Register\'!G:G,A{i})')
        ws.column_dimensions["A"].width=28; ws.column_dimensions["B"].width=80
        reg=wb.create_sheet("Control Register"); reg.sheet_view.showGridLines=False
        headers=["ID","Component","Control","Type","Risk","Current","Status","Target","Recommended Action","Restart","Owner","Due Date","Change Request","Maintenance Window","Closure Status","Notes"]
        reg.append(headers)
        for c in controls: reg.append([c["id"],c["component"],c["control"],c["control_type"],c["risk"],c["current"],c["status"],c["target"],c["recommended_action"],c["restart"],"","","","","Open" if c["status"] not in ("COMPLIANT",) else "Closed",""])
        findings=wb.create_sheet("Findings"); findings.sheet_view.showGridLines=False
        fheaders=["Finding ID","Control ID","Component","Severity","Finding","Risk / Limitation","Recommendation","Restart","Status"]
        findings.append(fheaders); count=1
        for c in controls:
            if c["status"] not in ("COMPLIANT",):
                findings.append([f"F-{count:03d}",c["id"],c["component"],c["risk"],f"{c['control']}: {c['status']} | Current: {c['current']}",c["risk_explanation"],c["recommended_action"],c["restart"],"Open"]); count+=1
        evidence_ws=wb.create_sheet("Evidence Index"); evidence_ws.sheet_view.showGridLines=False
        eheaders=["Evidence ID","Control ID","Phase","Source","Timestamp","Evidence File","SHA-256"]
        evidence_ws.append(eheaders)
        for i,c in enumerate(controls,1): evidence_ws.append([f"EV-{i:03d}",c["id"],"Current","Live Analyze",package["assessed_at"],"assessment-summary.json",package["evidence_sha256"]])
        cert=wb.create_sheet("Certificate Inventory"); cert.sheet_view.showGridLines=False
        cheaders=["Layer","Machine / Listener","Active Alias","Subject","Issuer","SAN","Valid Until","Days Remaining","Chain","Status"]
        cert.append(cheaders)
        signed=self.signed_ca_evidence()
        for item in signed.get("web_tier",[]): cert.append(["Web tier",f"{item.get('hostname')}:{item.get('port')}","N/A",item.get("subject_cn"),item.get("issuer_cn"),", ".join(item.get("sans",[])),item.get("not_after"),item.get("days_remaining"),"Trusted" if item.get("trusted") else "Unverified",item.get("status")])
        portal=signed.get("portal",{}); pc=portal.get("active_certificate") or {}; cert.append(["Portal",", ".join(m.get("machine_name","") for m in portal.get("machines",[])),portal.get("active_alias"),pc.get("subject"),pc.get("issuer"),", ".join(pc.get("sans",[])),pc.get("valid_until"),pc.get("days_remaining"),portal.get("chain",{}).get("status"),portal.get("status")])
        for m in signed.get("server",{}).get("machines",[]):
            mc=m.get("active_certificate") or {}; cert.append(["ArcGIS Server",m.get("machine_name"),m.get("active_alias"),mc.get("subject"),mc.get("issuer"),", ".join(mc.get("sans",[])),mc.get("valid_until"),mc.get("days_remaining"),m.get("chain",{}).get("status"),m.get("status")])
        for sheet in (reg,findings,evidence_ws,cert):
            sheet.freeze_panes="A2"; sheet.auto_filter.ref=sheet.dimensions
            for cell in sheet[1]: cell.font=Font(bold=True,color="FFFFFF"); cell.fill=PatternFill("solid",fgColor=dark); cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True)
            for row in sheet.iter_rows(min_row=2):
                for cell in row: cell.alignment=Alignment(vertical="top",wrap_text=True)
            for idx,col in enumerate(sheet.columns,1):
                max_len=min(55,max(len(str(cell.value or "")) for cell in col)+2); sheet.column_dimensions[__import__("openpyxl").utils.get_column_letter(idx)].width=max(12,max_len)
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    if str(cell.value) in ("NON-COMPLIANT","CRITICAL","FAIL"): cell.fill=PatternFill("solid",fgColor=red)
                    elif str(cell.value) in ("COMPLIANT","PASS"): cell.fill=PatternFill("solid",fgColor=green)
                    elif str(cell.value) in ("UNKNOWN","PARTIALLY VERIFIED","COMPLIANT WITH NOTE","PASS WITH NOTE"): cell.fill=PatternFill("solid",fgColor=orange)
        wb.calculation.fullCalcOnLoad=True; wb.calculation.forceFullCalc=True; wb.calculation.calcMode="auto"
        wb.save(path)
    @staticmethod
    def comparison_classification(before_status, after_status):
        before = str(before_status or "UNKNOWN").upper()
        after = str(after_status or "UNKNOWN").upper()
        good = {"COMPLIANT", "COMPLIANT WITH NOTE"}
        bad = {"NON-COMPLIANT", "CRITICAL", "NEEDS CONFIGURATION"}
        if before in bad and after in good:
            return "REMEDIATED", "CLOSED"
        if before == "PARTIALLY VERIFIED" and after in good:
            return "VERIFIED AND CLOSED", "CLOSED"
        if before in good and after in bad:
            return "REGRESSION", "OPEN"
        if before == "CRITICAL" and after == "NON-COMPLIANT":
            return "IMPROVED, STILL OPEN", "OPEN"
        if before in bad and after in bad:
            return "NO CHANGE, OPEN" if before == after else "CHANGED, STILL OPEN", "OPEN"
        if before in good and after in good:
            return "UNCHANGED, COMPLIANT" if before == after else "COMPLIANT CHANGE", "CLOSED"
        return "REVIEW REQUIRED", "REVIEW"

    def validate_baseline_environment(self, baseline):
        current_portal_id = str((self.portal_result or {}).get("id", "")).strip()
        baseline_portal_id = str(baseline.get("portal_organization_id", "")).strip()
        if (current_portal_id and baseline_portal_id and
                current_portal_id != "Tidak tersedia" and
                baseline_portal_id != "Tidak tersedia" and
                current_portal_id != baseline_portal_id):
            return False, "Portal organization ID berbeda. Comparison diblokir."
        pairs = (
            ("Portal Admin URL", baseline.get("portal_admin_url"), self.portal_admin_url.text()),
            ("Server Admin URL", baseline.get("server_admin_url"), self.server_admin_url.text()),
        )
        changed = [label for label, old, new in pairs
                   if self.normalized_environment_url(old) != self.normalized_environment_url(new)]
        if changed:
            return None, "Endpoint berubah: " + ", ".join(changed)
        return True, "Environment cocok."

    def build_comparison_package(self, baseline, current):
        before_by_id = {item.get("id"): item for item in baseline.get("controls", [])}
        comparisons = []
        for after in current.get("controls", []):
            before = before_by_id.get(after.get("id"), {})
            change, closure = self.comparison_classification(before.get("status"), after.get("status"))
            comparisons.append({
                "id": after.get("id"), "component": after.get("component"),
                "control": after.get("control"), "risk": after.get("risk"),
                "before_current": before.get("current", "Not available"),
                "before_status": before.get("status", "NOT AVAILABLE"),
                "after_current": after.get("current"), "after_status": after.get("status"),
                "change_classification": change, "closure_status": closure,
                "recommended_action": after.get("recommended_action"),
            })
        result = dict(current)
        result.update({
            "assessment_type": "Before/After Assessment",
            "phase": "AFTER", "comparison_performed": True,
            "baseline_assessment_id": baseline.get("assessment_id"),
            "baseline_assessed_at": baseline.get("assessed_at"),
            "comparisons": comparisons,
        })
        return result

    def export_comparison_docx(self, package, path):
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.shared import Inches, Pt
        doc = Document(); sec = doc.sections[0]
        sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65)
        sec.left_margin=Inches(.6); sec.right_margin=Inches(.6)
        title=doc.add_heading("ArcGIS Enterprise Hardening Before/After Report", 0)
        title.alignment=WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph(f"Baseline: {package.get('baseline_assessment_id')}\nAfter: {package.get('assessment_id')}")
        counts={}
        for row in package.get("comparisons", []): counts[row["change_classification"]]=counts.get(row["change_classification"],0)+1
        doc.add_heading("Executive Comparison Summary", level=1)
        for key in ("REMEDIATED","VERIFIED AND CLOSED","IMPROVED, STILL OPEN","NO CHANGE, OPEN","REGRESSION","UNCHANGED, COMPLIANT","REVIEW REQUIRED"):
            if counts.get(key): doc.add_paragraph(f"{key}: {counts[key]}")
        doc.add_heading("Control Comparison", level=1)
        table=doc.add_table(rows=1, cols=8); table.style="Table Grid"
        headers=("Component","Control","Before","Before status","After","After status","Change","Closure")
        for i,h in enumerate(headers): self.docx_set_cell_text(table.rows[0].cells[i],h,bold=True,size=8)
        for row in package.get("comparisons", []):
            cells=table.add_row().cells
            values=(row["component"],row["control"],row["before_current"],row["before_status"],row["after_current"],row["after_status"],row["change_classification"],row["closure_status"])
            for i,v in enumerate(values): self.docx_set_cell_text(cells[i],v,size=7)
        doc.add_heading("Evidence and Limitations", level=1)
        doc.add_paragraph("Before evidence berasal dari baseline folder yang dipilih. After evidence berasal dari Analyze Hardening terbaru. Tidak ada password, token, cookie, shared key, private key, atau password PFX/P12 yang disertakan.")
        doc.save(path)

    def export_comparison_xlsx(self, package, path):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb=Workbook(); ws=wb.active; ws.title="Assessment Summary"
        ws.append(["Field","Value"])
        for row in (("Assessment ID",package.get("assessment_id")),("Baseline Assessment ID",package.get("baseline_assessment_id")),("Assessment Type",package.get("assessment_type")),("After Date",package.get("assessed_at")),("Portal Admin URL",package.get("portal_admin_url")),("Server Admin URL",package.get("server_admin_url"))): ws.append(row)
        comp=wb.create_sheet("Control Comparison")
        headers=["Control ID","Component","Control","Risk","Before Current","Before Status","After Current","After Status","Change Classification","Closure Status","Recommended Action","Owner","Due Date","Change Request","Maintenance Window","Notes"]
        comp.append(headers)
        for r in package.get("comparisons", []): comp.append([r.get("id"),r.get("component"),r.get("control"),r.get("risk"),r.get("before_current"),r.get("before_status"),r.get("after_current"),r.get("after_status"),r.get("change_classification"),r.get("closure_status"),r.get("recommended_action"),"","","","",""])
        findings=wb.create_sheet("Findings"); findings.append(headers); [findings.append(list(row)) for row in comp.iter_rows(min_row=2, values_only=True) if row[9] != "CLOSED"]
        evidence=wb.create_sheet("Evidence Index"); evidence.append(["Assessment","File","SHA-256"]); evidence.append(["Before","evidence/before-evidence.json",package.get("baseline_evidence_sha256","")]); evidence.append(["After","evidence/after-evidence.json",package.get("evidence_sha256","")])
        cert=wb.create_sheet("Certificate Inventory"); cert.append(["Source","Evidence reference"]); cert.append(["After","evidence/after-evidence.json"])
        for sheet in wb.worksheets:
            sheet.freeze_panes="A2"; sheet.auto_filter.ref=sheet.dimensions
            for c in sheet[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="17365D"); c.alignment=Alignment(horizontal="center",wrap_text=True)
            for col in sheet.columns:
                letter=col[0].column_letter; sheet.column_dimensions[letter].width=min(55,max(12,max(len(str(c.value or "")) for c in col)+2))
            for row in sheet.iter_rows(min_row=2):
                for c in row: c.alignment=Alignment(vertical="top",wrap_text=True)
        wb.save(path)

    def write_assessment_package(self, package, output, baseline=None):
        output.mkdir(parents=True, exist_ok=False)
        evidence_dir=output/"evidence"; evidence_dir.mkdir(); logs_dir=output/"logs"; logs_dir.mkdir()
        safe=dict(package); safe.pop("evidence",None)
        if package.get("comparison_performed"):
            (evidence_dir/"before-evidence.json").write_text(json.dumps((baseline or {}).get("evidence",{}),indent=2,ensure_ascii=False,default=str),encoding="utf-8")
            (evidence_dir/"after-evidence.json").write_text(json.dumps(package["evidence"],indent=2,ensure_ascii=False,default=str),encoding="utf-8")
            safe["evidence_files"]=["evidence/before-evidence.json","evidence/after-evidence.json"]
            (output/"comparison.json").write_text(json.dumps(package.get("comparisons",[]),indent=2,ensure_ascii=False,default=str),encoding="utf-8")
            self.export_comparison_docx(package,output/"ArcGIS_Enterprise_Hardening_Before_After_Report.docx")
            self.export_comparison_xlsx(package,output/"ArcGIS_Enterprise_Hardening_Register.xlsx")
        else:
            (evidence_dir/"current-evidence.json").write_text(json.dumps(package["evidence"],indent=2,ensure_ascii=False,default=str),encoding="utf-8")
            safe["evidence_file"]="evidence/current-evidence.json"
            self.export_assessment_docx(package,output/"ArcGIS_Enterprise_Hardening_Report.docx")
            self.export_assessment_xlsx(package,output/"ArcGIS_Enterprise_Hardening_Register.xlsx")
        (output/"assessment-summary.json").write_text(json.dumps(safe,indent=2,ensure_ascii=False,default=str),encoding="utf-8")
        (logs_dir/"execution.log").write_text(self.log_label.toPlainText(),encoding="utf-8")

    def export_current_assessment(self):
        if not self.last_assessment_at or not self.portal_result or not self.server_result:
            QMessageBox.warning(self,"Belum ada assessment","Jalankan Analyze Hardening terlebih dahulu.")
            return
        selected=QFileDialog.getExistingDirectory(self,"Pilih folder induk atau folder baseline assessment")
        if not selected: return
        try:
            selected_path=Path(selected); summary_path=selected_path/"assessment-summary.json"
            if not summary_path.exists():
                package=self.build_assessment_package("Baseline Assessment")
                output=selected_path/package["assessment_id"]
                self.write_assessment_package(package,output)
                mode="Baseline Assessment"
            else:
                baseline=json.loads(summary_path.read_text(encoding="utf-8"))
                evidence_file=baseline.get("evidence_file")
                if evidence_file and (selected_path/evidence_file).exists():
                    baseline["evidence"]=json.loads((selected_path/evidence_file).read_text(encoding="utf-8"))
                match,message=self.validate_baseline_environment(baseline)
                if match is False: raise ValueError(message)
                if match is None:
                    answer=QMessageBox.question(self,"Endpoint environment berubah",message+"\n\nPortal organization ID masih cocok atau tidak tersedia. Lanjutkan comparison?",QMessageBox.StandardButton.Yes|QMessageBox.StandardButton.No,QMessageBox.StandardButton.No)
                    if answer != QMessageBox.StandardButton.Yes: return
                dialog=QMessageBox(self); dialog.setWindowTitle("Baseline assessment ditemukan")
                dialog.setText(f"Assessment ID: {baseline.get('assessment_id')}\nTanggal: {baseline.get('assessed_at')}\n\nPilih jenis report:")
                current_btn=dialog.addButton("Export Current Only",QMessageBox.ButtonRole.ActionRole)
                compare_btn=dialog.addButton("Create Before/After Report",QMessageBox.ButtonRole.AcceptRole)
                dialog.addButton(QMessageBox.StandardButton.Cancel); dialog.exec()
                clicked=dialog.clickedButton()
                if clicked not in (current_btn,compare_btn): return
                current=self.build_assessment_package("Current Assessment",baseline.get("assessment_id"))
                if clicked is current_btn:
                    output=selected_path/"current-assessments"/current["assessment_id"]
                    package=current; mode="Current Assessment"
                else:
                    package=self.build_comparison_package(baseline,current)
                    package["baseline_evidence_sha256"]=baseline.get("evidence_sha256","")
                    output=selected_path/"reassessments"/current["assessment_id"]
                    mode="Before/After Assessment"
                self.write_assessment_package(package,output,baseline)
            self.append_log(f"{mode} export selesai: {output}")
            QMessageBox.information(self,"Export Assessment berhasil",f"{mode} berhasil dibuat.\n\nFolder: {output}")
        except ModuleNotFoundError as error:
            QMessageBox.critical(self,"Dependency export belum tersedia",f"Module {error.name} belum tersedia.\n\nInstall dependencies:\npy -m pip install python-docx openpyxl")
        except Exception as error:
            QMessageBox.critical(self,"Export Assessment gagal",f"Export dihentikan tanpa mengubah konfigurasi ArcGIS.\n\n{type(error).__name__}: {error}")
            self.append_log(f"Assessment export gagal: {type(error).__name__}: {error}")

    def system_theme_name(self):
        # Use the native application palette as the source of truth for System mode.
        color = QApplication.instance().palette().color(
            QApplication.instance().palette().ColorRole.Window
        )
        return "Dark" if color.lightness() < 128 else "Light"

    def on_theme_changed(self, preference):
        self.apply_theme(preference, save=True)

    def apply_theme(self, preference, save=True):
        if preference not in ("System", "Dark", "Light"):
            preference = "System"
        self.theme_preference = preference
        effective = self.system_theme_name() if preference == "System" else preference
        self.current_effective_theme = effective
        if save:
            self.settings.setValue("appearance/theme", preference)
            self.settings.sync()
        self.setStyleSheet(
            self.dark_style_sheet() if effective == "Dark" else self.light_style_sheet()
        )
        self.refresh_semantic_colors()
        if getattr(self, "portal_result", None) or getattr(self, "server_result", None):
            self.update_connection_results()
        if hasattr(self, "status"):
            self.status.showMessage(
                f"Theme: {preference}" + (f" ({effective})" if preference == "System" else ""),
                2500,
            )

    def semantic_status_color(self, status):
        dark_colors = {
            "COMPLIANT": "#19c37d",
            "COMPLIANT WITH NOTE": "#19c37d",
            "PASS WITH NOTE": "#19c37d",
            "NON-COMPLIANT": "#ff5f56",
            "UNKNOWN": "#f4c542",
            "REVIEW": "#f4c542",
            "NEEDS CONFIGURATION": "#f4c542",
            "PARTIALLY VERIFIED": "#f4c542",
            "CRITICAL": "#ff2d20",
        }
        light_colors = {
            "COMPLIANT": "#2f6b3b",
            "COMPLIANT WITH NOTE": "#2f6b3b",
            "PASS WITH NOTE": "#2f6b3b",
            "NON-COMPLIANT": "#9f2d2d",
            "UNKNOWN": "#765400",
            "REVIEW": "#765400",
            "NEEDS CONFIGURATION": "#765400",
            "PARTIALLY VERIFIED": "#765400",
            "CRITICAL": "#7f1d1d",
        }
        palette = dark_colors if self.current_effective_theme == "Dark" else light_colors
        fallback = "#f4c542" if self.current_effective_theme == "Dark" else "#765400"
        return QColor(palette.get(str(status), fallback))

    def refresh_semantic_colors(self):
        if not hasattr(self, "controls_table"):
            return
        for row in range(self.controls_table.rowCount()):
            item = self.controls_table.item(row, 4)
            if not item:
                continue
            status = item.data(Qt.ItemDataRole.UserRole) or item.text().replace("\n", " ")
            item.setForeground(self.semantic_status_color(status))

    @staticmethod
    def dark_style_sheet():
        return """
            QMainWindow, QWidget { background:#1e1e1e; color:#f2f2f2; font-family:'Segoe UI'; font-size:13px; }
            QLabel#titleLabel { font-size:26px; font-weight:700; }
            QLabel#subtitleLabel, QLabel#sectionDescription { color:#b8b8b8; font-size:14px; }
            QLabel#sectionTitle { font-size:18px; font-weight:700; }
            QLabel#themeLabel { color:#cfcfcf; font-weight:600; }
            QCheckBox { background:transparent; }
            QFrame#card, QFrame#fixedFooter { background:#292929; border:1px solid #414141; border-radius:8px; }
            QLineEdit, QComboBox { background:#181818; color:#f2f2f2; border:1px solid #505050; border-radius:5px; padding:8px; }
            QComboBox::drop-down { border:none; width:24px; }
            QComboBox QAbstractItemView { background:#292929; color:#f2f2f2; selection-background-color:#0078d4; }
            QPushButton { background:#333; color:#f2f2f2; border:1px solid #565656; border-radius:5px; padding:9px 14px; font-weight:600; }
            QPushButton:hover { background:#414141; }
            QPushButton#primaryButton { background:#0078d4; color:white; border:none; }
            QPushButton:disabled { color:#777; background:#303030; }
            QLabel#resultLabel, QLabel#activityLabel { background:#202020; border:1px solid #444; border-radius:6px; padding:10px; }
            QTabWidget::pane { border:1px solid #414141; }
            QTabBar::tab { background:#292929; padding:10px 18px; margin-right:2px; }
            QTabBar::tab:selected { background:#0078d4; color:white; }
            QTableWidget { background:#202020; alternate-background-color:#242424; gridline-color:#444; border:1px solid #444; }
            QHeaderView::section { background:#303030; padding:8px; border:1px solid #444; font-weight:600; }
            QToolTip { background:#303030; color:#ffffff; border:1px solid #666; padding:5px; }
            QPlainTextEdit#executionLog { background:#171717; color:#f2f2f2; border:1px solid #444; border-radius:6px; padding:10px; font-family:'Consolas'; }
            QScrollBar:vertical, QScrollBar:horizontal { background:#222; }
            QScrollBar::handle:vertical, QScrollBar::handle:horizontal { background:#666; min-width:24px; min-height:24px; }
            QPushButton#rowPreviewButton, QPushButton#rowApplyButton, QPushButton#rowRollbackButton { min-width:54px; max-width:66px; padding:5px; font-size:11px; }
            QPushButton#rowApplyButton { min-width:48px; max-width:54px; background:#0078d4; color:white; border:none; }
            QPushButton#rowRollbackButton { min-width:62px; max-width:68px; background:#8a4b08; border:1px solid #b86b16; }
            QPushButton#rowRollbackButton:hover { background:#a45a0a; }
            QPushButton#rowPreviewButton:disabled, QPushButton#rowApplyButton:disabled, QPushButton#rowRollbackButton:disabled { color:#777; background:#303030; border:1px solid #444; }
        """

    @staticmethod
    def light_style_sheet():
        return """
            QMainWindow, QWidget { background:#f5f7fa; color:#000000; font-family:'Segoe UI'; font-size:13px; }
            QLabel { color:#000000; background:transparent; }
            QLabel#titleLabel { color:#000000; background:transparent; font-size:26px; font-weight:700; }
            QLabel#subtitleLabel, QLabel#sectionDescription { color:#000000; background:transparent; font-size:14px; }
            QLabel#sectionTitle { color:#000000; background:transparent; font-size:18px; font-weight:700; }
            QLabel#themeLabel { color:#000000; background:transparent; font-weight:600; }
            QCheckBox { color:#000000; background:transparent; padding:2px 0; }
            QCheckBox::indicator { background:#ffffff; border:1px solid #6b7280; border-radius:3px; width:16px; height:16px; }
            QCheckBox::indicator:checked { background:#0078d4; border-color:#0078d4; }
            QFormLayout QLabel { color:#000000; background:transparent; }
            QFrame#card, QFrame#fixedFooter { background:#ffffff; border:1px solid #cbd2d9; border-radius:8px; }
            QLineEdit, QComboBox { background:#ffffff; color:#000000; border:1px solid #aeb8c2; border-radius:5px; padding:8px; }
            QLineEdit:focus, QComboBox:focus { border:2px solid #0078d4; }
            QComboBox::drop-down { border:none; width:24px; }
            QComboBox QAbstractItemView { background:#ffffff; color:#000000; selection-background-color:#cfe8ff; selection-color:#000000; }
            QPushButton { background:#ffffff; color:#000000; border:1px solid #aeb8c2; border-radius:5px; padding:9px 14px; font-weight:600; }
            QPushButton:hover { background:#eef4fb; border-color:#7a9fc2; }
            QPushButton#primaryButton { background:#0078d4; color:white; border:none; }
            QPushButton:disabled { color:#8b949e; background:#edf0f3; border-color:#d5dbe1; }
            QLabel#resultLabel, QLabel#activityLabel { background:#ffffff; border:1px solid #cbd2d9; border-radius:6px; padding:10px; }
            QTabWidget::pane { background:#ffffff; border:1px solid #cbd2d9; }
            QTabBar::tab { background:#e9eef5; color:#000000; padding:10px 18px; margin-right:2px; border:1px solid #cbd2d9; }
            QTabBar::tab:selected { background:#0078d4; color:white; }
            QTableWidget { background:#ffffff; alternate-background-color:#f7f9fb; color:#000000; gridline-color:#d5dbe1; border:1px solid #cbd2d9; }
            QHeaderView::section { background:#e9eef5; color:#000000; padding:8px; border:1px solid #cbd2d9; font-weight:600; }
            QToolTip { background:#fffbe6; color:#000000; border:1px solid #b8a75b; padding:5px; }
            QPlainTextEdit#executionLog { background:#ffffff; color:#000000; border:1px solid #cbd2d9; border-radius:6px; padding:10px; font-family:'Consolas'; }
            QScrollBar:vertical, QScrollBar:horizontal { background:#eef1f4; }
            QScrollBar::handle:vertical, QScrollBar::handle:horizontal { background:#aab4be; min-width:24px; min-height:24px; }
            QPushButton#rowPreviewButton, QPushButton#rowApplyButton, QPushButton#rowRollbackButton { min-width:54px; max-width:66px; padding:5px; font-size:11px; }
            QPushButton#rowApplyButton { min-width:48px; max-width:54px; background:#0078d4; color:white; border:none; }
            QPushButton#rowRollbackButton { min-width:62px; max-width:68px; background:#fff1df; color:#7a3f00; border:1px solid #c77700; }
            QPushButton#rowRollbackButton:hover { background:#ffe2bd; }
            QPushButton#rowPreviewButton:disabled, QPushButton#rowApplyButton:disabled, QPushButton#rowRollbackButton:disabled { color:#8b949e; background:#edf0f3; border:1px solid #d5dbe1; }
        """


def main():
    app = QApplication(sys.argv)
    app_font = QFont("Segoe UI", 10)
    app.setFont(app_font)
    window = MainWindow()
    window.show()
    sys.exit(app.exec())


if __name__ == "__main__":
    main()
